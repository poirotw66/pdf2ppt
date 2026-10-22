"""P2 layout-quality metrics: grouping, list detection, title detection.

These metrics compare conversion outputs (``PageResult`` text boxes / PPTX
shapes) against optional ground-truth JSON fixtures. They are independent of
LibreOffice fidelity scoring so CI can exercise them without soffice.
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from .layout import bbox_area, intersection_area, iou
from .models import BBox, PageResult, TextBlock


@dataclass(slots=True)
class LayoutGroupingMetrics:
    """Paragraph/text-box grouping quality vs ground truth."""

    predicted_box_count: int
    ground_truth_paragraph_count: int
    box_count_ratio: float
    mean_paragraph_iou: float
    merge_error_count: int
    split_error_count: int
    unmatched_prediction_count: int
    unmatched_ground_truth_count: int
    list_precision: float | None = None
    list_recall: float | None = None
    list_f1: float | None = None
    title_detected: bool | None = None
    title_expected: bool | None = None
    title_correct: bool | None = None
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class GroundTruthParagraph:
    id: str
    bbox: BBox
    text: str = ""
    role: str = "body"  # title | body | list | caption | header | footer
    list_level: int | None = None


@dataclass(slots=True)
class LayoutGroundTruth:
    page_number: int = 1
    paragraphs: list[GroundTruthParagraph] = field(default_factory=list)
    expect_title: bool = False
    expect_list_item_count: int | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "expect_title": self.expect_title,
            "expect_list_item_count": self.expect_list_item_count,
            "paragraphs": [
                {
                    "id": item.id,
                    "bbox": list(item.bbox),
                    "text": item.text,
                    "role": item.role,
                    "list_level": item.list_level,
                }
                for item in self.paragraphs
            ],
        }


def load_layout_ground_truth(path: Path) -> LayoutGroundTruth:
    payload = json.loads(path.read_text(encoding="utf-8"))
    paragraphs = [
        GroundTruthParagraph(
            id=str(item.get("id") or f"gt_{index}"),
            bbox=_as_bbox(item["bbox"]),
            text=str(item.get("text") or ""),
            role=str(item.get("role") or "body"),
            list_level=item.get("list_level"),
        )
        for index, item in enumerate(payload.get("paragraphs") or [])
    ]
    return LayoutGroundTruth(
        page_number=int(payload.get("page_number") or 1),
        paragraphs=paragraphs,
        expect_title=bool(payload.get("expect_title")),
        expect_list_item_count=payload.get("expect_list_item_count"),
    )


def _as_bbox(raw: Any) -> BBox:
    values = [float(value) for value in list(raw)[:4]]
    return (values[0], values[1], values[2], values[3])


def evaluate_text_grouping(
    predicted_blocks: list[TextBlock],
    ground_truth: LayoutGroundTruth,
    *,
    match_iou_threshold: float = 0.35,
) -> LayoutGroupingMetrics:
    """Score text-box grouping against ground-truth paragraphs."""

    preds = list(predicted_blocks)
    gts = list(ground_truth.paragraphs)
    pred_count = len(preds)
    gt_count = len(gts)
    ratio = (pred_count / gt_count) if gt_count else (0.0 if pred_count == 0 else float("inf"))

    # Greedy IoU matching.
    pairs: list[tuple[float, int, int]] = []
    for pi, pred in enumerate(preds):
        for gi, gt in enumerate(gts):
            score = iou(pred.bbox, gt.bbox)
            if score >= match_iou_threshold:
                pairs.append((score, pi, gi))
    pairs.sort(reverse=True)
    matched_pred: set[int] = set()
    matched_gt: set[int] = set()
    matched_ious: list[float] = []
    for score, pi, gi in pairs:
        if pi in matched_pred or gi in matched_gt:
            continue
        matched_pred.add(pi)
        matched_gt.add(gi)
        matched_ious.append(score)

    merge_errors = 0
    split_errors = 0
    for pi, pred in enumerate(preds):
        covering = [
            gi
            for gi, gt in enumerate(gts)
            if intersection_area(pred.bbox, gt.bbox) / max(bbox_area(gt.bbox), 1.0) >= 0.5
        ]
        if len(covering) >= 2:
            merge_errors += 1
    for gi, gt in enumerate(gts):
        covering = [
            pi
            for pi, pred in enumerate(preds)
            if intersection_area(pred.bbox, gt.bbox) / max(bbox_area(pred.bbox), 1.0) >= 0.5
        ]
        if len(covering) >= 2:
            split_errors += 1

    list_precision, list_recall, list_f1 = _list_detection_scores(preds, gts)
    title_detected = any(
        block.block_role == "title" or block.use_title_placeholder or block.layout_label in {"title", "paragraph_title"}
        for block in preds
    )
    title_expected = ground_truth.expect_title or any(item.role == "title" for item in gts)
    title_correct: bool | None
    if not title_expected and not title_detected:
        title_correct = True
    elif title_expected:
        title_correct = title_detected
    else:
        title_correct = False

    return LayoutGroupingMetrics(
        predicted_box_count=pred_count,
        ground_truth_paragraph_count=gt_count,
        box_count_ratio=round(ratio, 6) if ratio != float("inf") else -1.0,
        mean_paragraph_iou=round(sum(matched_ious) / len(matched_ious), 6) if matched_ious else 0.0,
        merge_error_count=merge_errors,
        split_error_count=split_errors,
        unmatched_prediction_count=pred_count - len(matched_pred),
        unmatched_ground_truth_count=gt_count - len(matched_gt),
        list_precision=list_precision,
        list_recall=list_recall,
        list_f1=list_f1,
        title_detected=title_detected,
        title_expected=title_expected,
        title_correct=title_correct,
    )


def evaluate_page_layout(
    page_result: PageResult,
    ground_truth: LayoutGroundTruth,
) -> LayoutGroupingMetrics:
    return evaluate_text_grouping(page_result.text_blocks, ground_truth)


def _list_detection_scores(
    preds: list[TextBlock],
    gts: list[GroundTruthParagraph],
) -> tuple[float | None, float | None, float | None]:
    gt_list = [item for item in gts if item.role == "list"]
    pred_items = _expand_predicted_list_items(preds)
    if not gt_list and not pred_items:
        return None, None, None

    gt_items = len(gt_list)
    if gt_items == 0 and not pred_items:
        return None, None, None

    # Match predicted list *items* (often nested in one text box) to GT items.
    matched = 0
    used_gt: set[int] = set()
    for pred_bbox, _pred_text in pred_items:
        best_gi = None
        best_score = 0.0
        for gi, gt in enumerate(gt_list):
            if gi in used_gt:
                continue
            score = iou(pred_bbox, gt.bbox)
            # Also accept containment when a single list box covers many items.
            containment = intersection_area(pred_bbox, gt.bbox) / max(bbox_area(gt.bbox), 1.0)
            score = max(score, containment * 0.9)
            if score > best_score:
                best_score = score
                best_gi = gi
        if best_gi is not None and best_score >= 0.25:
            used_gt.add(best_gi)
            matched += 1

    pred_count = max(len(pred_items), 1) if pred_items else 0
    precision = matched / pred_count if pred_items else 0.0
    recall = matched / max(gt_items, 1) if gt_items else (1.0 if not pred_items else 0.0)
    if precision + recall <= 0:
        f1 = 0.0
    else:
        f1 = 2 * precision * recall / (precision + recall)
    return round(precision, 6), round(recall, 6), round(f1, 6)


def _expand_predicted_list_items(preds: list[TextBlock]) -> list[tuple[BBox, str]]:
    items: list[tuple[BBox, str]] = []
    for block in preds:
        paragraphs = block.paragraphs or []
        list_paras = [para for para in paragraphs if para.list_level is not None]
        if list_paras:
            # Distribute the parent bbox into horizontal bands for IoU matching.
            band_height = max(1.0, (block.bbox[3] - block.bbox[1]) / max(len(list_paras), 1))
            for index, para in enumerate(list_paras):
                y0 = block.bbox[1] + index * band_height
                y1 = y0 + band_height
                items.append(((block.bbox[0], y0, block.bbox[2], y1), para.text))
            continue
        if block.block_role == "list":
            lines = [line for line in block.text.splitlines() if line.strip()] or [block.text]
            band_height = max(1.0, (block.bbox[3] - block.bbox[1]) / max(len(lines), 1))
            for index, line in enumerate(lines):
                y0 = block.bbox[1] + index * band_height
                y1 = y0 + band_height
                items.append(((block.bbox[0], y0, block.bbox[2], y1), line))
    return items


def count_pptx_bullet_paragraphs(pptx_path: Path) -> int:
    """Count paragraphs that have a real bullet (buChar/buAutoNum), not literal glyphs."""

    from pptx import Presentation
    from pptx.oxml.ns import qn

    presentation = Presentation(str(pptx_path))
    count = 0
    bullet_tags = {qn("a:buChar"), qn("a:buAutoNum"), qn("a:buBlip")}
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                try:
                    p_pr = paragraph._p.pPr  # noqa: SLF001
                except Exception:  # noqa: BLE001
                    continue
                if p_pr is None:
                    continue
                if any(child.tag in bullet_tags for child in p_pr):
                    count += 1
    return count


def slide_has_title_placeholder_text(pptx_path: Path) -> bool:
    from pptx import Presentation

    presentation = Presentation(str(pptx_path))
    for slide in presentation.slides:
        title = slide.shapes.title
        if title is not None and (title.text or "").strip():
            return True
    return False


# ---------------------------------------------------------------------------
# Phase 3 layout-planner acceptance metrics
# ---------------------------------------------------------------------------


@dataclass(slots=True)
class PlannerAcceptanceMetrics:
    """Title / header-footer / reading-order scores vs planner ground truth."""

    title_accuracy: float | None
    header_footer_filter_rate: float | None
    reading_order_correctness: float | None
    title_detected: bool | None = None
    title_expected: bool | None = None
    expected_header_footer_count: int = 0
    filtered_header_footer_count: int = 0
    reading_order_pairs_checked: int = 0
    reading_order_pairs_correct: int = 0
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def load_planner_ground_truth(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Planner ground truth must be an object: {path}")
    return payload


def _normalize_match_text(text: str) -> str:
    return "".join(ch.lower() for ch in text if not ch.isspace())


def _text_match_score(predicted: str, ground_truth: str) -> float:
    pred = _normalize_match_text(predicted)
    gt = _normalize_match_text(ground_truth)
    if not gt:
        return 0.0
    if pred == gt:
        return 1.0
    if gt in pred or pred in gt:
        return 0.85
    # Cheap token overlap for short labels.
    pred_tokens = set(pred.split()) if " " in predicted else {pred}
    gt_tokens = set(gt.split()) if " " in ground_truth else {gt}
    if not pred_tokens or not gt_tokens:
        return 0.0
    inter = len(pred_tokens & gt_tokens)
    return inter / max(len(pred_tokens | gt_tokens), 1)


def evaluate_planner_acceptance(
    predicted_blocks: list[TextBlock],
    ground_truth: dict[str, Any],
    *,
    match_iou_threshold: float = 0.15,
) -> PlannerAcceptanceMetrics:
    """Score Phase-3 acceptance criteria against ``*.planner.json`` labels.

    Matching prefers normalized text containment over strict IoU so synthetic
    GT bboxes that only roughly cover real blocks still produce reading-order
    and header/footer scores.
    """

    paragraphs = list(ground_truth.get("paragraphs") or [])
    expect_title = bool(ground_truth.get("expect_title")) or any(
        str(item.get("role") or "") == "title" for item in paragraphs
    )
    title_detected = any(
        block.block_role == "title"
        or block.use_title_placeholder
        or block.layout_label in {"title", "paragraph_title"}
        for block in predicted_blocks
    )
    if not expect_title and not title_detected:
        title_accuracy: float | None = 1.0
    elif expect_title:
        title_accuracy = 1.0 if title_detected else 0.0
    else:
        title_accuracy = 0.0

    if "expect_header_footer_ids" in ground_truth:
        hf_ids = [str(item) for item in (ground_truth.get("expect_header_footer_ids") or [])]
    else:
        hf_ids = [
            str(item.get("id"))
            for item in paragraphs
            if str(item.get("role") or "") in {"header", "footer", "page_number"}
        ]
    hf_gts = [item for item in paragraphs if str(item.get("id")) in set(hf_ids)]
    filtered = 0
    for gt in hf_gts:
        gt_text = str(gt.get("text") or "")
        gt_bbox = _as_bbox(gt["bbox"])
        still_present = False
        for block in predicted_blocks:
            text_score = _text_match_score(block.text, gt_text) if gt_text else 0.0
            geom_score = max(
                iou(block.bbox, gt_bbox),
                intersection_area(block.bbox, gt_bbox) / max(bbox_area(gt_bbox), 1.0),
            )
            if text_score >= 0.85 or (text_score >= 0.5 and geom_score >= 0.15) or geom_score >= 0.5:
                still_present = True
                break
        if not still_present:
            filtered += 1
    filter_rate = (filtered / len(hf_gts)) if hf_gts else None

    matched: list[tuple[int, TextBlock]] = []
    used_pred: set[int] = set()
    for gt_index, gt in enumerate(paragraphs):
        if str(gt.get("role") or "") in {"header", "footer", "page_number"}:
            continue
        gt_text = str(gt.get("text") or "")
        gt_bbox = _as_bbox(gt["bbox"])
        best_pi = None
        best_score = 0.0
        for pi, pred in enumerate(predicted_blocks):
            if pi in used_pred:
                continue
            text_score = _text_match_score(pred.text, gt_text) if gt_text else 0.0
            geom_score = iou(pred.bbox, gt_bbox)
            score = max(text_score, geom_score * 0.9)
            if score > best_score:
                best_score = score
                best_pi = pi
        if best_pi is not None and best_score >= max(0.35, match_iou_threshold):
            used_pred.add(best_pi)
            matched.append((gt_index, predicted_blocks[best_pi]))

    pairs_checked = 0
    pairs_correct = 0
    for left in range(len(matched)):
        for right in range(left + 1, len(matched)):
            gt_i, pred_i = matched[left]
            gt_j, pred_j = matched[right]
            pairs_checked += 1
            if (gt_i < gt_j) == (pred_i.reading_order <= pred_j.reading_order):
                pairs_correct += 1
    order_score = (pairs_correct / pairs_checked) if pairs_checked else None

    return PlannerAcceptanceMetrics(
        title_accuracy=title_accuracy,
        header_footer_filter_rate=None if filter_rate is None else round(filter_rate, 6),
        reading_order_correctness=None if order_score is None else round(order_score, 6),
        title_detected=title_detected,
        title_expected=expect_title,
        expected_header_footer_count=len(hf_gts),
        filtered_header_footer_count=filtered,
        reading_order_pairs_checked=pairs_checked,
        reading_order_pairs_correct=pairs_correct,
    )


def evaluate_planner_acceptance_from_pptx(
    pptx_path: Path,
    ground_truth: dict[str, Any],
    *,
    slide_index: int = 0,
) -> PlannerAcceptanceMetrics:
    """Score acceptance from the actual PPTX (not IR-only / GT-oracle).

    Title: title placeholder text present when expected.
    Header/footer: GT HF strings absent from editable PPTX text.
    Reading order: consecutive non-HF GT strings appear in PPTX text order.
    """

    from pptx import Presentation

    from .refine import extract_pptx_slide_text

    presentation = Presentation(str(pptx_path))
    if slide_index < 0 or slide_index >= len(presentation.slides):
        return PlannerAcceptanceMetrics(
            title_accuracy=None,
            header_footer_filter_rate=None,
            reading_order_correctness=None,
            notes=["pptx slide missing"],
        )

    slide = presentation.slides[slide_index]
    title_shape = slide.shapes.title
    title_text = (title_shape.text or "").strip() if title_shape is not None else ""
    slide_text = extract_pptx_slide_text(pptx_path, slide_index=slide_index)
    normalized_slide = _normalize_match_text(slide_text)

    paragraphs = list(ground_truth.get("paragraphs") or [])
    expect_title = bool(ground_truth.get("expect_title")) or any(
        str(item.get("role") or "") == "title" for item in paragraphs
    )
    title_detected = bool(title_text)
    if expect_title:
        # Prefer placeholder; also accept GT title text present as any shape.
        gt_title = next(
            (str(item.get("text") or "") for item in paragraphs if str(item.get("role") or "") == "title"),
            "",
        )
        if title_detected:
            title_accuracy: float | None = 1.0
        elif gt_title and _normalize_match_text(gt_title) in normalized_slide:
            title_accuracy = 0.5
            title_detected = True
        else:
            title_accuracy = 0.0
    else:
        title_accuracy = 1.0 if not title_detected else 0.0

    if "expect_header_footer_ids" in ground_truth:
        hf_ids = [str(item) for item in (ground_truth.get("expect_header_footer_ids") or [])]
    else:
        hf_ids = [
            str(item.get("id"))
            for item in paragraphs
            if str(item.get("role") or "") in {"header", "footer", "page_number"}
        ]
    hf_gts = [item for item in paragraphs if str(item.get("id")) in set(hf_ids)]
    filtered = 0
    for gt in hf_gts:
        gt_text = str(gt.get("text") or "")
        if not gt_text:
            filtered += 1
            continue
        if _normalize_match_text(gt_text) not in normalized_slide:
            filtered += 1
    filter_rate = (filtered / len(hf_gts)) if hf_gts else None

    # Reading order from PPTX concatenated text positions of non-HF GT strings.
    body_gts = [
        item
        for item in paragraphs
        if str(item.get("role") or "") not in {"header", "footer", "page_number"}
    ]
    positions: list[tuple[int, int]] = []
    for gt_index, gt in enumerate(body_gts):
        needle = _normalize_match_text(str(gt.get("text") or ""))
        if not needle:
            continue
        pos = normalized_slide.find(needle)
        if pos >= 0:
            positions.append((gt_index, pos))
    pairs_checked = 0
    pairs_correct = 0
    for left in range(len(positions)):
        for right in range(left + 1, len(positions)):
            gi, pi = positions[left]
            gj, pj = positions[right]
            pairs_checked += 1
            if (gi < gj) == (pi <= pj):
                pairs_correct += 1
    order_score = (pairs_correct / pairs_checked) if pairs_checked else None

    return PlannerAcceptanceMetrics(
        title_accuracy=title_accuracy,
        header_footer_filter_rate=None if filter_rate is None else round(filter_rate, 6),
        reading_order_correctness=None if order_score is None else round(order_score, 6),
        title_detected=title_detected,
        title_expected=expect_title,
        expected_header_footer_count=len(hf_gts),
        filtered_header_footer_count=filtered,
        reading_order_pairs_checked=pairs_checked,
        reading_order_pairs_correct=pairs_correct,
        notes=["measured_from_pptx"],
    )
