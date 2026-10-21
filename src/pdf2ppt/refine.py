"""Deterministic render-compare-refine loop (roadmap P4).

After PPTX is built, optionally re-render each slide (LibreOffice → PDF → PNG),
compare against the source PDF page, diagnose localized problems with rule-based
checks (no LLM), apply bounded fixes to SlideIR / ``PageResult`` shapes, and
keep a change only when page SSIM **strictly improves** and editability / text
content do not regress.

Opt-in via ``ConversionOptions.refine`` / CLI ``--refine``. When LibreOffice
(``soffice``) is unavailable **or fails**, the loop skips with a warning,
restores the last accepted PPTX, and never fails the conversion.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import logging
import math
import tempfile
import zipfile
from collections.abc import Callable
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path
from time import perf_counter
from typing import Any

import cv2
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation

from .models import BBox, ImagePlacement, PageResult, TextBlock, TextParagraph
from .render_compare import (
    DEFAULT_COMPARE_DPI,
    RenderCompareError,
    compute_pixel_error,
    compute_ssim,
    extract_pdf_page_text,
    find_libreoffice_executable,
    inspect_slide_editability,
    render_pdf_page_image,
    render_pptx_pages,
    resize_to_match,
    text_token_f1,
)
from .text_style import (
    choose_measurement_font,
    classify_text_script,
    measure_text_dimensions,
)

logger = logging.getLogger(__name__)

DEFAULT_REFINE_ITERS = 2
DEFAULT_MAX_FIXES_PER_SLIDE = 8
# Bench on PR #14 fixtures: ~97% of attempts were rejected; 8 caps waste while
# still covering the rare accepted path (median accepts per improved slide ≤ 2).
DEFAULT_MAX_ATTEMPTS_PER_SLIDE = 8
# Cap real (XML-changing) attempts per element so one title cannot burn the slide budget.
DEFAULT_MAX_ATTEMPTS_PER_ELEMENT = 2
DEFAULT_TIME_BUDGET_SECONDS = 90.0
DEFAULT_FONT_SHRINK_STEP_PT = 0.5
# Measurement floor (layout heuristics); shrink floor is stricter (see below).
DEFAULT_MIN_FONT_SIZE_PT = 6.0
# Never shrink below max(absolute floor, source_size * ratio). 0.95 ≈ at most one
# 0.5pt step on typical body sizes (14pt → floor 13.3pt), preserving OCR legibility.
DEFAULT_SHRINK_MIN_FONT_SIZE_PT = 8.0
DEFAULT_MIN_FONT_RATIO = 0.95
DEFAULT_LIBREOFFICE_TIMEOUT_CAP_SECONDS = 180
# Floor for soffice trial timeouts. Real single-slide converts are ~1.5–3s; the
# previous 5s floor left short budgets with leftover time but zero trials.
# Effective minimum is max(floor, slack × measured per-page render).
DEFAULT_MIN_SOFFICE_TIMEOUT_SECONDS = 2.5
DEFAULT_SOFFICE_TIMEOUT_SLACK = 1.5
# Baseline cost model: soffice render + per-page PDF/diagnose work.
# Render: ~2.5s fixed + ~0.12s/page (PR #16 / #17 real soffice). Diagnosis
# (source PDF render + diagnose_slide) is ~0.45s/page on real decks and must be
# counted or short budgets spend the whole budget diagnosing.
DEFAULT_BASELINE_FIXED_SECONDS = 2.5
DEFAULT_BASELINE_PER_PAGE_SECONDS = 0.12
DEFAULT_DIAGNOSE_PER_PAGE_SECONDS = 0.45
# When a full-deck baseline+diagnose would leave too little room for trials,
# cap that phase to this fraction of the remaining wall budget.
DEFAULT_BASELINE_BUDGET_FRACTION = 0.40
# Weights for continuous page severity. Fidelity gap must dominate issue-count
# terms (real decks: raw issue sums correlated ~0.96 with score when weight=1).
SEVERITY_FIDELITY_WEIGHT = 40.0
SEVERITY_ISSUE_SUM_WEIGHT = 0.25
SEVERITY_ISSUE_COUNT_WEIGHT = 0.15
SEVERITY_ISSUE_AREA_WEIGHT = 0.25
SEVERITY_ISSUE_COUNT_NORM = 20.0  # log1p(n) / log1p(this) → ~1 at n=20
# SSIM is insensitive to flat recolors; a modest gap still catches structural
# misses. Local MAE catches mis-coloured / missing patches that SSIM misses.
DEFAULT_LOCAL_SSIM_GAP = 0.02
DEFAULT_LOCAL_MAE_MIN = 12.0
DEFAULT_LOCAL_MAE_PAGE_RATIO = 2.5
DEFAULT_OVERFLOW_RATIO = 1.02
DEFAULT_LINE_DRIFT_PT = 3.0
DEFAULT_BBOX_EDGE_MARGIN_PT = 0.5
# Strict improvement required (neutral / flat accepts are rejected).
METRIC_IMPROVE_EPSILON = 1e-4
EDITABILITY_EPSILON = 1e-6
TEXT_F1_EPSILON = 1e-6


@dataclass(slots=True)
class RefineIssue:
    """One localized problem found on a rendered slide."""

    kind: str
    element_id: str
    severity: float
    details: dict[str, Any] = field(default_factory=dict)

    @property
    def has_render_evidence(self) -> bool:
        return bool(self.details.get("render_evidence"))

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class RefineFixAttempt:
    """One rule-based fix tried during refinement."""

    issue_kind: str
    element_id: str
    action: str
    accepted: bool
    metric_before: float
    metric_after: float | None
    details: dict[str, Any] = field(default_factory=dict)
    xml_changed: bool | None = None
    noop: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class SlideRefineReport:
    """Per-slide refine summary written into the JSON report."""

    page_number: int
    iterations: int
    metric_before: float | None
    metric_after: float | None
    issues: list[RefineIssue] = field(default_factory=list)
    fixes: list[RefineFixAttempt] = field(default_factory=list)
    skipped_reason: str | None = None
    attempt_count: int = 0
    soffice_call_count: int = 0
    editability_before: float | None = None
    editability_after: float | None = None
    text_f1_before: float | None = None
    text_f1_after: float | None = None
    # Budget / ordering telemetry (deterministic for a given wall budget).
    severity: float | None = None
    # Wall seconds actually spent on this page (first-pass / reallocation / total).
    slot_seconds: float | None = None
    time_used_seconds: float | None = None
    borrowed_seconds: float | None = None
    # Sum of allowances granted (may exceed used; informational only).
    granted_seconds: float | None = None
    accepted_fix_count: int = 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "iterations": self.iterations,
            "metric_before": self.metric_before,
            "metric_after": self.metric_after,
            "issues": [issue.to_dict() for issue in self.issues],
            "fixes": [fix.to_dict() for fix in self.fixes],
            "skipped_reason": self.skipped_reason,
            "attempt_count": self.attempt_count,
            "soffice_call_count": self.soffice_call_count,
            "editability_before": self.editability_before,
            "editability_after": self.editability_after,
            "text_f1_before": self.text_f1_before,
            "text_f1_after": self.text_f1_after,
            "severity": self.severity,
            "slot_seconds": self.slot_seconds,
            "time_used_seconds": self.time_used_seconds,
            "borrowed_seconds": self.borrowed_seconds,
            "granted_seconds": self.granted_seconds,
            "accepted_fix_count": self.accepted_fix_count,
        }


@dataclass(slots=True)
class DeckRefineReport:
    """Aggregate refine report for one conversion."""

    enabled: bool
    skipped_reason: str | None
    iterations: int
    dpi: float
    pages: list[SlideRefineReport] = field(default_factory=list)
    runtime_seconds: float | None = None
    soffice_call_count: int = 0
    attempt_count: int = 0
    # Baseline / budget telemetry.
    baseline_mode: str | None = None  # "full" | "partial" | "skipped"
    baseline_page_count: int | None = None
    pages_skipped: list[dict[str, Any]] = field(default_factory=list)
    reallocation_passes: int = 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "enabled": self.enabled,
            "skipped_reason": self.skipped_reason,
            "iterations": self.iterations,
            "dpi": self.dpi,
            "runtime_seconds": self.runtime_seconds,
            "soffice_call_count": self.soffice_call_count,
            "attempt_count": self.attempt_count,
            "baseline_mode": self.baseline_mode,
            "baseline_page_count": self.baseline_page_count,
            "pages_skipped": list(self.pages_skipped),
            "reallocation_passes": self.reallocation_passes,
            "pages": [page.to_dict() for page in self.pages],
        }


# ---------------------------------------------------------------------------
# Geometry helpers (page points ↔ raster pixels)
# ---------------------------------------------------------------------------


def _page_to_pixel_bbox(
    bbox: BBox,
    *,
    page_width_pt: float,
    page_height_pt: float,
    image_size: tuple[int, int],
    pad_px: int = 0,
) -> tuple[int, int, int, int]:
    width_px, height_px = image_size
    scale_x = width_px / max(page_width_pt, 1.0)
    scale_y = height_px / max(page_height_pt, 1.0)
    x0 = int(math.floor(bbox[0] * scale_x)) - pad_px
    y0 = int(math.floor(bbox[1] * scale_y)) - pad_px
    x1 = int(math.ceil(bbox[2] * scale_x)) + pad_px
    y1 = int(math.ceil(bbox[3] * scale_y)) + pad_px
    return (
        max(0, x0),
        max(0, y0),
        min(width_px, x1),
        min(height_px, y1),
    )


def _union_bboxes(boxes: list[BBox]) -> BBox | None:
    if not boxes:
        return None
    return (
        min(box[0] for box in boxes),
        min(box[1] for box in boxes),
        max(box[2] for box in boxes),
        max(box[3] for box in boxes),
    )


def _bbox_width(bbox: BBox) -> float:
    return max(0.0, bbox[2] - bbox[0])


def _bbox_height(bbox: BBox) -> float:
    return max(0.0, bbox[3] - bbox[1])


def _bboxes_intersect(a: BBox, b: BBox) -> bool:
    return a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]


def _clamp_bbox_to_page(bbox: BBox, page_width_pt: float, page_height_pt: float) -> BBox:
    return (
        max(0.0, min(bbox[0], page_width_pt)),
        max(0.0, min(bbox[1], page_height_pt)),
        max(0.0, min(bbox[2], page_width_pt)),
        max(0.0, min(bbox[3], page_height_pt)),
    )


def _source_line_count(block: TextBlock) -> int:
    if block.line_bboxes:
        return max(1, len(block.line_bboxes))
    text = block.text or ""
    lines = [line for line in text.splitlines() if line.strip()]
    return max(1, len(lines) if lines else (1 if text.strip() else 0))


def _ink_mask(gray: np.ndarray) -> np.ndarray:
    threshold = float(np.median(gray)) - 18.0
    ink = gray < threshold
    if not np.any(ink):
        ink = gray > float(np.median(gray)) + 18.0
    return ink


def _ink_row_centers(image: Image.Image, bbox_px: tuple[int, int, int, int]) -> list[float]:
    """Y centers (in page-image pixels) of horizontal ink bands inside a crop."""

    x0, y0, x1, y1 = bbox_px
    if x1 <= x0 or y1 <= y0:
        return []
    crop = np.array(image.convert("RGB"))[y0:y1, x0:x1]
    if crop.size == 0:
        return []
    gray = cv2.cvtColor(crop, cv2.COLOR_RGB2GRAY)
    ink = _ink_mask(gray)
    if not np.any(ink):
        return []
    row_density = ink.mean(axis=1)
    active = row_density > max(0.02, float(row_density.mean()) * 0.5)
    centers: list[float] = []
    start: int | None = None
    for index, is_active in enumerate(active.tolist()):
        if is_active and start is None:
            start = index
        elif not is_active and start is not None:
            centers.append(y0 + (start + index - 1) / 2.0)
            start = None
    if start is not None:
        centers.append(y0 + (start + len(active) - 1) / 2.0)
    return centers


def _ink_extent_px(image: Image.Image, bbox_px: tuple[int, int, int, int]) -> tuple[int, int, int, int] | None:
    """Bounding box of ink pixels inside ``bbox_px``, or ``None`` if empty."""

    x0, y0, x1, y1 = bbox_px
    if x1 <= x0 or y1 <= y0:
        return None
    crop = np.array(image.convert("RGB"))[y0:y1, x0:x1]
    if crop.size == 0:
        return None
    gray = cv2.cvtColor(crop, cv2.COLOR_RGB2GRAY)
    ink = _ink_mask(gray)
    if not np.any(ink):
        return None
    ys, xs = np.where(ink)
    return (
        int(x0 + xs.min()),
        int(y0 + ys.min()),
        int(x0 + xs.max() + 1),
        int(y0 + ys.max() + 1),
    )


def region_ssim(
    reference: Image.Image,
    candidate: Image.Image,
    bbox_px: tuple[int, int, int, int],
) -> float | None:
    """SSIM on a pixel crop; ``None`` when the crop is empty or too small."""

    x0, y0, x1, y1 = bbox_px
    if x1 - x0 < 8 or y1 - y0 < 8:
        return None
    matched = resize_to_match(reference, candidate)
    ref_crop = reference.crop((x0, y0, x1, y1))
    hyp_crop = matched.crop((x0, y0, x1, y1))
    return compute_ssim(ref_crop, hyp_crop)


def region_mae(
    reference: Image.Image,
    candidate: Image.Image,
    bbox_px: tuple[int, int, int, int],
) -> float | None:
    """Mean absolute RGB error on a pixel crop; ``None`` when too small."""

    x0, y0, x1, y1 = bbox_px
    if x1 - x0 < 4 or y1 - y0 < 4:
        return None
    matched = resize_to_match(reference, candidate)
    ref_crop = reference.crop((x0, y0, x1, y1))
    hyp_crop = matched.crop((x0, y0, x1, y1))
    return float(compute_pixel_error(ref_crop, hyp_crop)["pixel_mae"])


# ---------------------------------------------------------------------------
# PPTX helpers (XML digest, text extract, editability, single-slide write)
# ---------------------------------------------------------------------------


def pptx_content_digest(pptx_path: Path) -> str:
    """Stable hash of slide XML + media so no-op writes are detectable."""

    digest = hashlib.sha256()
    with zipfile.ZipFile(pptx_path) as archive:
        names = sorted(
            name for name in archive.namelist() if name.startswith("ppt/slides/") or name.startswith("ppt/media/")
        )
        for name in names:
            digest.update(name.encode("utf-8"))
            digest.update(archive.read(name))
    return digest.hexdigest()


def extract_pptx_slide_text(pptx_path: Path, *, slide_index: int = 0) -> str:
    """Concatenate editable text-frame text from one slide (no OCR)."""

    presentation = Presentation(str(pptx_path))
    if slide_index < 0 or slide_index >= len(presentation.slides):
        return ""
    slide = presentation.slides[slide_index]
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = (shape.text_frame.text or "").strip()
            if text:
                parts.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    cell_text = (cell.text or "").strip()
                    if cell_text:
                        parts.append(cell_text)
    return "\n".join(parts)


def slide_editability_score(pptx_path: Path, *, slide_index: int = 0) -> float:
    presentation = Presentation(str(pptx_path))
    if slide_index < 0 or slide_index >= len(presentation.slides):
        return 0.0
    metrics = inspect_slide_editability(
        presentation.slides[slide_index],
        slide_width=int(presentation.slide_width or 1),
        slide_height=int(presentation.slide_height or 1),
    )
    return float(metrics.score)


def slide_text_f1_vs_pdf(pptx_path: Path, pdf_path: Path, *, page_index: int) -> float:
    """Bag-of-tokens F1: PDF page text vs editable PPTX text (no OCR)."""

    reference = extract_pdf_page_text(pdf_path, page_index)
    presentation = Presentation(str(pptx_path))
    # Single-slide trial PPTX always has one slide at index 0.
    slide_index = 0 if len(presentation.slides) == 1 else page_index
    candidate = extract_pptx_slide_text(pptx_path, slide_index=slide_index)
    return float(text_token_f1(reference, candidate)["text_f1"])


# ---------------------------------------------------------------------------
# Diagnosis (pure / deterministic; unit-tested without LibreOffice)
# ---------------------------------------------------------------------------


def diagnose_text_overflow(
    block: TextBlock,
    *,
    page_width_pt: float,
    page_height_pt: float,
    candidate: Image.Image | None = None,
    overflow_ratio: float = DEFAULT_OVERFLOW_RATIO,
) -> RefineIssue | None:
    """Flag overflow from render ink spill, or (lower priority) measurement."""

    x0, y0, x1, y1 = block.bbox
    past_edge = (
        x0 < -DEFAULT_BBOX_EDGE_MARGIN_PT
        or y0 < -DEFAULT_BBOX_EDGE_MARGIN_PT
        or x1 > page_width_pt + DEFAULT_BBOX_EDGE_MARGIN_PT
        or y1 > page_height_pt + DEFAULT_BBOX_EDGE_MARGIN_PT
    )

    render_spill = False
    spill_px = 0
    if candidate is not None:
        # Look slightly outside the box for ink that belongs to this block.
        box_px = _page_to_pixel_bbox(
            block.bbox,
            page_width_pt=page_width_pt,
            page_height_pt=page_height_pt,
            image_size=candidate.size,
            pad_px=0,
        )
        search_px = _page_to_pixel_bbox(
            block.bbox,
            page_width_pt=page_width_pt,
            page_height_pt=page_height_pt,
            image_size=candidate.size,
            pad_px=8,
        )
        extent = _ink_extent_px(candidate, search_px)
        if extent is not None:
            spill_left = max(0, box_px[0] - extent[0])
            spill_top = max(0, box_px[1] - extent[1])
            spill_right = max(0, extent[2] - box_px[2])
            spill_bottom = max(0, extent[3] - box_px[3])
            spill_px = spill_left + spill_top + spill_right + spill_bottom
            # Also treat ink past the slide edge as overflow.
            edge_spill = (
                extent[0] <= 1
                or extent[1] <= 1
                or extent[2] >= candidate.size[0] - 1
                or extent[3] >= candidate.size[1] - 1
            )
            render_spill = spill_px >= 3 or (edge_spill and past_edge)

    text = (block.text or "").strip()
    measured_overflow = False
    measured_w = 0.0
    measured_h = 0.0
    if text:
        script = classify_text_script(text)
        font_path = choose_measurement_font(script)
        font_size = max(DEFAULT_MIN_FONT_SIZE_PT, float(block.font_size or 12.0))
        if font_path is not None:
            measured_w, measured_h = measure_text_dimensions(text, int(round(font_size)), font_path)
            box_w = max(1.0, _bbox_width(block.bbox))
            box_h = max(1.0, _bbox_height(block.bbox))
            measured_overflow = measured_w > box_w * overflow_ratio or measured_h > box_h * overflow_ratio

    if not past_edge and not render_spill and not measured_overflow:
        return None

    # Prefer render evidence; measurement-only is demoted so it does not eat budget.
    render_evidence = render_spill or past_edge
    if not render_evidence and measured_overflow:
        severity = 0.15
    else:
        severity = 0.5 if past_edge else 0.0
        if render_spill:
            severity = max(severity, min(1.0, spill_px / 20.0))
        if measured_overflow and render_evidence:
            box_w = max(1.0, _bbox_width(block.bbox))
            box_h = max(1.0, _bbox_height(block.bbox))
            severity = max(
                severity,
                min(1.0, max(measured_w / box_w, measured_h / box_h, 1.0) - 1.0),
            )

    return RefineIssue(
        kind="text_overflow",
        element_id=block.id,
        severity=round(severity, 4),
        details={
            "past_slide_edge": past_edge,
            "render_spill": render_spill,
            "spill_px": spill_px,
            "measured_width_pt": round(measured_w, 3),
            "measured_height_pt": round(measured_h, 3),
            "bbox": list(block.bbox),
            "font_size": block.font_size,
            "render_evidence": render_evidence,
            "element_kind": "text",
        },
    )


def diagnose_line_drift(
    block: TextBlock,
    *,
    rendered_line_ys_pt: list[float],
    drift_pt: float = DEFAULT_LINE_DRIFT_PT,
) -> RefineIssue | None:
    """Flag when rendered line y-centers drift from PDF ``line_bboxes``."""

    if not block.line_bboxes or not rendered_line_ys_pt:
        return None
    expected = sorted((box[1] + box[3]) / 2.0 for box in block.line_bboxes)
    observed = sorted(rendered_line_ys_pt)
    pair_count = min(len(expected), len(observed))
    if pair_count == 0:
        return None
    drifts = [abs(expected[index] - observed[index]) for index in range(pair_count)]
    max_drift = max(drifts)
    if max_drift <= drift_pt:
        return None
    return RefineIssue(
        kind="line_drift",
        element_id=block.id,
        severity=round(min(1.0, max_drift / max(drift_pt * 4.0, 1.0)), 4),
        details={
            "max_drift_pt": round(max_drift, 3),
            "expected_ys_pt": [round(value, 3) for value in expected],
            "observed_ys_pt": [round(value, 3) for value in observed],
            "render_evidence": True,
            "element_kind": "text",
        },
    )


def diagnose_extra_wrap(
    block: TextBlock,
    *,
    candidate: Image.Image | None = None,
    page_width_pt: float | None = None,
    page_height_pt: float | None = None,
) -> RefineIssue | None:
    """Flag extra wrapping only when render ink shows more lines than the source."""

    source_lines = _source_line_count(block)
    if source_lines <= 0:
        return None
    if candidate is None or page_width_pt is None or page_height_pt is None:
        return None

    bbox_px = _page_to_pixel_bbox(
        block.bbox,
        page_width_pt=page_width_pt,
        page_height_pt=page_height_pt,
        image_size=candidate.size,
        pad_px=4,
    )
    ink_rows = _ink_row_centers(candidate, bbox_px)
    rendered_lines = len(ink_rows)
    if rendered_lines <= source_lines:
        return None
    return RefineIssue(
        kind="extra_wrap",
        element_id=block.id,
        severity=round(min(1.0, (rendered_lines - source_lines) / max(source_lines, 1)), 4),
        details={
            "source_line_count": source_lines,
            "estimated_line_count": rendered_lines,
            "font_size": block.font_size,
            "bbox": list(block.bbox),
            "render_evidence": True,
            "element_kind": "text",
        },
    )


def diagnose_low_local_ssim(
    *,
    element_id: str,
    element_kind: str,
    local_ssim: float | None,
    page_ssim: float,
    gap: float = DEFAULT_LOCAL_SSIM_GAP,
    local_mae: float | None = None,
    page_mae: float = 0.0,
    mae_min: float = DEFAULT_LOCAL_MAE_MIN,
    mae_page_ratio: float = DEFAULT_LOCAL_MAE_PAGE_RATIO,
) -> RefineIssue | None:
    """Flag an element region that is visually far worse than the page average."""

    ssim_bad = local_ssim is not None and local_ssim < page_ssim - gap
    mae_bad = local_mae is not None and local_mae >= mae_min and local_mae >= max(mae_min, page_mae * mae_page_ratio)
    if not ssim_bad and not mae_bad:
        return None
    severity = 0.0
    if ssim_bad and local_ssim is not None:
        severity = max(severity, min(1.0, (page_ssim - gap) - local_ssim))
    if mae_bad and local_mae is not None:
        severity = max(severity, min(1.0, local_mae / 64.0))
    return RefineIssue(
        kind="low_local_ssim",
        element_id=element_id,
        severity=round(severity, 4),
        details={
            "local_ssim": None if local_ssim is None else round(local_ssim, 6),
            "page_ssim": round(page_ssim, 6),
            "gap": gap,
            "local_mae": None if local_mae is None else round(local_mae, 4),
            "page_mae": round(page_mae, 4),
            "ssim_bad": ssim_bad,
            "mae_bad": mae_bad,
            "render_evidence": True,
            "element_kind": element_kind,
        },
    )


def _issue_sort_key(issue: RefineIssue) -> tuple[int, float, str, str]:
    # Render-evidence issues first, then higher severity.
    return (0 if issue.has_render_evidence else 1, -issue.severity, issue.kind, issue.element_id)


def diagnose_slide(
    reference: Image.Image,
    candidate: Image.Image,
    page_result: PageResult,
    *,
    page_ssim: float | None = None,
    local_ssim_gap: float = DEFAULT_LOCAL_SSIM_GAP,
) -> tuple[float, list[RefineIssue]]:
    """Run all deterministic per-slide checks. Returns ``(page_ssim, issues)``."""

    matched = resize_to_match(reference, candidate)
    ssim = page_ssim if page_ssim is not None else compute_ssim(reference, matched)
    page_mae = float(compute_pixel_error(reference, matched)["pixel_mae"])
    issues: list[RefineIssue] = []

    def _check_region(element_id: str, element_kind: str, bbox_px: tuple[int, int, int, int]) -> None:
        local = region_ssim(reference, matched, bbox_px)
        local_mae = region_mae(reference, matched, bbox_px)
        low = diagnose_low_local_ssim(
            element_id=element_id,
            element_kind=element_kind,
            local_ssim=local,
            page_ssim=ssim,
            gap=local_ssim_gap,
            local_mae=local_mae,
            page_mae=page_mae,
        )
        if low is not None:
            issues.append(low)

    for block in page_result.text_blocks:
        overflow = diagnose_text_overflow(
            block,
            page_width_pt=page_result.width_pt,
            page_height_pt=page_result.height_pt,
            candidate=matched,
        )
        if overflow is not None:
            issues.append(overflow)
        wrap = diagnose_extra_wrap(
            block,
            candidate=matched,
            page_width_pt=page_result.width_pt,
            page_height_pt=page_result.height_pt,
        )
        if wrap is not None:
            issues.append(wrap)

        bbox_px = _page_to_pixel_bbox(
            block.bbox,
            page_width_pt=page_result.width_pt,
            page_height_pt=page_result.height_pt,
            image_size=reference.size,
            pad_px=4,
        )
        if block.line_bboxes:
            ink_ys_px = _ink_row_centers(matched, bbox_px)
            scale_y = reference.size[1] / max(page_result.height_pt, 1.0)
            ink_ys_pt = [y / scale_y for y in ink_ys_px]
            drift = diagnose_line_drift(block, rendered_line_ys_pt=ink_ys_pt)
            if drift is not None:
                issues.append(drift)

        _check_region(block.id, "text", bbox_px)

    for shape in page_result.shape_elements:
        bbox_px = _page_to_pixel_bbox(
            shape.bbox,
            page_width_pt=page_result.width_pt,
            page_height_pt=page_result.height_pt,
            image_size=reference.size,
            pad_px=2,
        )
        _check_region(shape.id, "shape", bbox_px)

    for index, image_element in enumerate(page_result.image_elements):
        element_id = f"image_{index}"
        bbox_px = _page_to_pixel_bbox(
            image_element.bbox,
            page_width_pt=page_result.width_pt,
            page_height_pt=page_result.height_pt,
            image_size=reference.size,
            pad_px=2,
        )
        _check_region(element_id, "image", bbox_px)

    issues.sort(key=_issue_sort_key)
    return ssim, issues


# ---------------------------------------------------------------------------
# Rule-based fixes
# ---------------------------------------------------------------------------


def _find_text_block(page: PageResult, element_id: str) -> TextBlock | None:
    for block in page.text_blocks:
        if block.id == element_id:
            return block
    return None


def _replace_text_block(page: PageResult, updated: TextBlock) -> None:
    page.text_blocks = [updated if block.id == updated.id else block for block in page.text_blocks]


def shrink_font_floor(
    source_font_size: float,
    *,
    min_font_size_pt: float = DEFAULT_SHRINK_MIN_FONT_SIZE_PT,
    min_font_ratio: float = DEFAULT_MIN_FONT_RATIO,
) -> float:
    """Lower bound for ``shrink_font``: ``max(absolute_min, source * ratio)``."""

    source = max(0.0, float(source_font_size))
    absolute = max(0.0, float(min_font_size_pt))
    ratio = min(1.0, max(0.0, float(min_font_ratio)))
    return max(absolute, source * ratio)


def apply_shrink_font(
    page: PageResult,
    element_id: str,
    *,
    step_pt: float = DEFAULT_FONT_SHRINK_STEP_PT,
    min_font_size_pt: float = DEFAULT_SHRINK_MIN_FONT_SIZE_PT,
    min_font_ratio: float = DEFAULT_MIN_FONT_RATIO,
    source_font_size: float | None = None,
    source_paragraph_sizes: list[float] | None = None,
) -> dict[str, Any] | None:
    """Shrink ``block.font_size`` and every paragraph ``font_size`` (writer prefers paras).

    Never shrinks below ``max(min_font_size_pt, source * min_font_ratio)``. ``source_*``
    should be the pre-refine sizes so repeated shrinks cannot dig under the ratio floor.
    """

    block = _find_text_block(page, element_id)
    if block is None:
        return None
    current = float(block.font_size or 12.0)
    source = float(source_font_size) if source_font_size is not None else current
    floor = shrink_font_floor(
        source,
        min_font_size_pt=min_font_size_pt,
        min_font_ratio=min_font_ratio,
    )
    new_size = max(floor, round(current - step_pt, 2))
    paragraphs = list(block.paragraphs or [])
    para_changed = False
    new_paragraphs: list[TextParagraph] | None = None
    if paragraphs:
        new_paragraphs = []
        for index, para in enumerate(paragraphs):
            para_size = float(para.font_size) if para.font_size is not None else current
            if source_paragraph_sizes is not None and index < len(source_paragraph_sizes):
                para_source = float(source_paragraph_sizes[index])
            else:
                para_source = para_size
            para_floor = shrink_font_floor(
                para_source,
                min_font_size_pt=min_font_size_pt,
                min_font_ratio=min_font_ratio,
            )
            para_new = max(para_floor, round(para_size - step_pt, 2))
            if para_new < para_size - 1e-9:
                para_changed = True
            new_paragraphs.append(replace(para, font_size=para_new))
    if new_size >= current - 1e-9 and not para_changed:
        return None
    updated = replace(
        block,
        font_size=new_size if new_size < current - 1e-9 else block.font_size,
        paragraphs=new_paragraphs if new_paragraphs is not None else block.paragraphs,
    )
    _replace_text_block(page, updated)
    return {
        "font_size_before": current,
        "font_size_after": updated.font_size,
        "font_size_floor": round(floor, 3),
        "source_font_size": round(source, 3),
        "paragraph_sizes_scaled": para_changed,
    }


def apply_match_line_bboxes(page: PageResult, element_id: str) -> dict[str, Any] | None:
    """Widen / reposition a text box to the union of its source ``line_bboxes``."""

    block = _find_text_block(page, element_id)
    if block is None or not block.line_bboxes:
        return None
    union = _union_bboxes(list(block.line_bboxes))
    if union is None:
        return None
    pad = max(0.5, (union[3] - union[1]) * 0.05)
    new_bbox = _clamp_bbox_to_page(
        (union[0] - pad, union[1] - pad, union[2] + pad, union[3] + pad),
        page.width_pt,
        page.height_pt,
    )
    if new_bbox == block.bbox:
        return None
    before = list(block.bbox)
    _replace_text_block(page, replace(block, bbox=new_bbox))
    return {"bbox_before": before, "bbox_after": list(new_bbox)}


def apply_widen_for_wrap(page: PageResult, element_id: str) -> dict[str, Any] | None:
    """Widen a text box so the measured single-line width fits (no extra wrap)."""

    block = _find_text_block(page, element_id)
    if block is None:
        return None
    text = (block.text or "").strip()
    if not text:
        return None
    script = classify_text_script(text)
    font_path = choose_measurement_font(script)
    if font_path is None:
        return None
    font_size = max(DEFAULT_MIN_FONT_SIZE_PT, float(block.font_size or 12.0))
    source_lines = max(1, _source_line_count(block))
    lines = [line for line in text.splitlines() if line.strip()] or [text.replace("\n", " ")]
    max_line_w = 0.0
    for line in lines:
        line_w, _line_h = measure_text_dimensions(line, int(round(font_size)), font_path)
        max_line_w = max(max_line_w, line_w)
    needed_width = max_line_w * 1.02
    current_width = _bbox_width(block.bbox)
    if needed_width <= current_width:
        return None
    x0, y0, x1, y1 = block.bbox
    extra = needed_width - current_width
    new_x1 = min(page.width_pt, x1 + extra)
    remaining = needed_width - (new_x1 - x0)
    new_x0 = max(0.0, x0 - max(0.0, remaining))
    line_h = max(font_size * 1.2, _bbox_height(block.bbox) / source_lines)
    needed_height = line_h * source_lines
    new_y1 = min(page.height_pt, max(y1, y0 + needed_height))
    new_bbox = (new_x0, y0, new_x1, new_y1)
    if new_bbox == block.bbox:
        return None
    before = list(block.bbox)
    # Pin font size so a larger box cannot grow glyphs via fit_text / fallback sizing.
    _replace_text_block(page, replace(block, bbox=new_bbox, pin_font_size=True))
    return {
        "bbox_before": before,
        "bbox_after": list(new_bbox),
        "needed_width_pt": round(needed_width, 3),
        "pin_font_size": True,
    }


def _erase_overlapping_text_from_crop(
    crop: Image.Image,
    *,
    crop_bbox_px: tuple[int, int, int, int],
    page: PageResult,
    image_size: tuple[int, int],
) -> Image.Image:
    """Paint over glyphs of editable text that intersect a non-text raster crop."""

    cx0, cy0, cx1, cy1 = crop_bbox_px
    crop_box_pt = (
        cx0 * page.width_pt / max(image_size[0], 1),
        cy0 * page.height_pt / max(image_size[1], 1),
        cx1 * page.width_pt / max(image_size[0], 1),
        cy1 * page.height_pt / max(image_size[1], 1),
    )
    result = crop.convert("RGB").copy()
    draw = ImageDraw.Draw(result)
    arr = np.array(result)
    # Fill color: median of crop (approximates local background).
    fill = tuple(int(v) for v in np.median(arr.reshape(-1, 3), axis=0))
    for block in page.text_blocks:
        if not _bboxes_intersect(block.bbox, crop_box_pt):
            continue
        tb = _page_to_pixel_bbox(
            block.bbox,
            page_width_pt=page.width_pt,
            page_height_pt=page.height_pt,
            image_size=image_size,
            pad_px=1,
        )
        # Convert to crop-local coordinates.
        lx0 = max(0, tb[0] - cx0)
        ly0 = max(0, tb[1] - cy0)
        lx1 = min(result.size[0], tb[2] - cx0)
        ly1 = min(result.size[1], tb[3] - cy0)
        if lx1 > lx0 and ly1 > ly0:
            draw.rectangle((lx0, ly0, lx1 - 1, ly1 - 1), fill=fill)
    return result


def apply_raster_fallback(
    page: PageResult,
    element_id: str,
    *,
    reference_page_image: Image.Image,
) -> dict[str, Any] | None:
    """Last resort for **non-text** regions: paste a PDF crop with text glyphs erased.

    Never applied to text elements (would ghost under editable titles).
    """

    if _find_text_block(page, element_id) is not None:
        return None

    bbox: BBox | None = None
    kind = "unknown"
    for shape in page.shape_elements:
        if shape.id == element_id:
            bbox = shape.bbox
            kind = "shape"
            break
    if bbox is None and element_id.startswith("image_"):
        try:
            index = int(element_id.split("_", 1)[1])
        except (IndexError, ValueError):
            index = -1
        if 0 <= index < len(page.image_elements):
            bbox = page.image_elements[index].bbox
            kind = "image"

    if bbox is None:
        return None

    bbox_px = _page_to_pixel_bbox(
        bbox,
        page_width_pt=page.width_pt,
        page_height_pt=page.height_pt,
        image_size=reference_page_image.size,
        pad_px=1,
    )
    x0, y0, x1, y1 = bbox_px
    if x1 <= x0 or y1 <= y0:
        return None
    crop = reference_page_image.crop((x0, y0, x1, y1)).convert("RGB")
    crop = _erase_overlapping_text_from_crop(
        crop,
        crop_bbox_px=bbox_px,
        page=page,
        image_size=reference_page_image.size,
    )
    buffer = io.BytesIO()
    crop.save(buffer, format="PNG")
    png_bytes = buffer.getvalue()

    placement = ImagePlacement(bbox=bbox, png_bytes=png_bytes, z_order=-10)
    page.image_elements = list(page.image_elements) + [placement]

    removed_shape = False
    if kind == "shape":
        before_count = len(page.shape_elements)
        page.shape_elements = [shape for shape in page.shape_elements if shape.id != element_id]
        removed_shape = len(page.shape_elements) < before_count

    return {
        "kind": kind,
        "bbox": list(bbox),
        "crop_px": [x0, y0, x1, y1],
        "removed_shape": removed_shape,
        "bytes": len(png_bytes),
        "text_glyphs_erased": True,
    }


def propose_fixes_for_issue(issue: RefineIssue) -> list[str]:
    """Ordered fix actions for an issue kind (deterministic priority)."""

    element_kind = str(issue.details.get("element_kind") or "unknown")
    if issue.kind in {"text_overflow", "extra_wrap"}:
        # Never raster-fallback text.
        return ["shrink_font", "widen_bbox", "match_line_bboxes"]
    if issue.kind == "line_drift":
        return ["match_line_bboxes", "widen_bbox", "shrink_font"]
    if issue.kind == "low_local_ssim":
        if element_kind == "text":
            return ["match_line_bboxes", "widen_bbox", "shrink_font"]
        # Non-text visual defects: try raster first (the real defect class).
        return ["raster_fallback"]
    if element_kind == "text":
        return ["shrink_font", "widen_bbox", "match_line_bboxes"]
    return ["raster_fallback"]


def apply_fix_action(
    page: PageResult,
    *,
    element_id: str,
    action: str,
    reference_page_image: Image.Image | None,
    min_font_size_pt: float = DEFAULT_SHRINK_MIN_FONT_SIZE_PT,
    min_font_ratio: float = DEFAULT_MIN_FONT_RATIO,
    source_font_size: float | None = None,
    source_paragraph_sizes: list[float] | None = None,
) -> dict[str, Any] | None:
    if action == "shrink_font":
        return apply_shrink_font(
            page,
            element_id,
            min_font_size_pt=min_font_size_pt,
            min_font_ratio=min_font_ratio,
            source_font_size=source_font_size,
            source_paragraph_sizes=source_paragraph_sizes,
        )
    if action == "widen_bbox":
        return apply_widen_for_wrap(page, element_id)
    if action == "match_line_bboxes":
        return apply_match_line_bboxes(page, element_id)
    if action == "raster_fallback":
        if reference_page_image is None:
            return None
        return apply_raster_fallback(page, element_id, reference_page_image=reference_page_image)
    return None


# ---------------------------------------------------------------------------
# Accept / reject loop
# ---------------------------------------------------------------------------


RenderPagesFn = Callable[..., list[Image.Image]]
RenderPdfPageFn = Callable[..., Image.Image]
WritePptxFn = Callable[..., Path]


def _deepcopy_pages(pages: list[PageResult]) -> list[PageResult]:
    return copy.deepcopy(pages)


def min_soffice_timeout_seconds(
    last_render_seconds: float | None = None,
    *,
    last_render_page_count: int = 1,
) -> float:
    """Minimum soffice allowance before issuing a trial (avoids false 1s timeouts).

    Uses ``max(floor, slack × per-page render)`` so a multi-page baseline does
    not inflate the trial floor to the whole-deck duration. Floor is ~2.5s
    (single-slide LibreOffice converts are typically 1.5–3s).
    """

    floor = float(DEFAULT_MIN_SOFFICE_TIMEOUT_SECONDS)
    if last_render_seconds is not None and last_render_seconds > 0:
        per_page = float(last_render_seconds) / float(max(1, int(last_render_page_count)))
        return max(floor, float(DEFAULT_SOFFICE_TIMEOUT_SLACK) * per_page)
    return floor


def _soffice_timeout_seconds(remaining_seconds: float, *, minimum: float) -> int:
    """Clamp remaining budget to a soffice timeout under the 180s cap and above ``minimum``."""

    capped = min(float(DEFAULT_LIBREOFFICE_TIMEOUT_CAP_SECONDS), max(float(remaining_seconds), float(minimum)))
    return max(1, int(math.ceil(capped)))


def _issue_area_fraction(issue: RefineIssue, page_area_pt: float) -> float:
    """Fraction of page area covered by an issue bbox (0 when unknown)."""

    bbox = issue.details.get("bbox")
    if not isinstance(bbox, (list, tuple)) or len(bbox) != 4 or page_area_pt <= 0:
        return 0.0
    width = max(0.0, float(bbox[2]) - float(bbox[0]))
    height = max(0.0, float(bbox[3]) - float(bbox[1]))
    return min(1.0, (width * height) / page_area_pt)


def page_refine_severity(
    *,
    page_ssim: float | None,
    issues: list[RefineIssue],
    page_width_pt: float = 1.0,
    page_height_pt: float = 1.0,
) -> float:
    """Continuous discriminating page severity (not clamped to 1.0).

    Fidelity gap ``(1 - SSIM)`` is scaled by :data:`SEVERITY_FIDELITY_WEIGHT` so
    it dominates normalised issue terms (mean severity, soft count, area).
    """

    fidelity_gap = 0.0 if page_ssim is None else max(0.0, 1.0 - float(page_ssim))
    if not issues:
        return round(SEVERITY_FIDELITY_WEIGHT * fidelity_gap, 6)

    page_area = max(1.0, float(page_width_pt) * float(page_height_pt))
    issue_count = len(issues)
    mean_sev = sum(float(issue.severity) for issue in issues) / float(issue_count)
    count_term = math.log1p(float(issue_count)) / math.log1p(float(SEVERITY_ISSUE_COUNT_NORM))
    area_term = min(1.0, sum(_issue_area_fraction(issue, page_area) for issue in issues))
    score = (
        SEVERITY_FIDELITY_WEIGHT * fidelity_gap
        + SEVERITY_ISSUE_SUM_WEIGHT * mean_sev
        + SEVERITY_ISSUE_COUNT_WEIGHT * count_term
        + SEVERITY_ISSUE_AREA_WEIGHT * area_term
    )
    return round(score, 6)


def page_issue_severity(issues: list[RefineIssue]) -> float:
    """Aggregate diagnosed severity for a page (continuous; 0 if clean).

    Prefer :func:`page_refine_severity` when baseline SSIM is available so
    fidelity gap discriminates pages whose per-issue severities all saturate.
    """

    return page_refine_severity(page_ssim=None, issues=issues)


def page_prerank_score(page: PageResult) -> float:
    """Cheap conversion-time risk score used before baseline render.

    Favours pages with more shapes/tables/images, lower pipeline quality, or
    non-element backgrounds so short budgets still baseline the riskiest slides
    first (deterministic: higher score wins, page index breaks ties).
    """

    quality = page.quality_score
    quality_risk = (
        (1.0 - float(quality.text_confidence))
        + (1.0 - float(quality.layout_overlap_score))
        + (1.0 - float(quality.style_recovery_score))
        + (1.0 - float(quality.editable_ratio))
    )
    bg_risk = 0.0 if page.background_mode == "elements" else 0.5
    fallback_risk = 1.0 if page.fallback_reason else 0.0
    large_text = sum(1 for block in page.text_blocks if float(block.font_size or 0.0) >= 24.0)
    return (
        0.3 * len(page.text_blocks)
        + 0.5 * len(page.shape_elements)
        + 0.8 * len(page.table_elements)
        + 0.4 * len(page.image_elements)
        + 1.5 * quality_risk
        + bg_risk
        + fallback_risk
        + 0.4 * large_text
    )


def estimate_baseline_render_seconds(
    page_count: int,
    *,
    last_render_seconds: float | None = None,
    last_page_count: int | None = None,
) -> float:
    """Estimate soffice wall time for a baseline render of ``page_count`` slides."""

    count = max(0, int(page_count))
    if (
        last_render_seconds is not None
        and last_render_seconds > 0
        and last_page_count is not None
        and last_page_count > 0
    ):
        scaled = float(last_render_seconds) * (float(count) / float(last_page_count))
        return max(float(DEFAULT_MIN_SOFFICE_TIMEOUT_SECONDS), scaled)
    return float(DEFAULT_BASELINE_FIXED_SECONDS) + float(DEFAULT_BASELINE_PER_PAGE_SECONDS) * float(count)


def estimate_diagnose_seconds(
    page_count: int,
    *,
    diagnose_per_page_seconds: float | None = None,
) -> float:
    """Estimate wall time to PDF-render + diagnose ``page_count`` baseline slides."""

    per_page = (
        float(DEFAULT_DIAGNOSE_PER_PAGE_SECONDS)
        if diagnose_per_page_seconds is None
        else max(0.0, float(diagnose_per_page_seconds))
    )
    return per_page * float(max(0, int(page_count)))


def estimate_baseline_phase_seconds(
    page_count: int,
    *,
    last_render_seconds: float | None = None,
    last_page_count: int | None = None,
    diagnose_per_page_seconds: float | None = None,
) -> float:
    """Estimate render + diagnosis cost for a baseline covering ``page_count`` slides."""

    return estimate_baseline_render_seconds(
        page_count,
        last_render_seconds=last_render_seconds,
        last_page_count=last_page_count,
    ) + estimate_diagnose_seconds(
        page_count,
        diagnose_per_page_seconds=diagnose_per_page_seconds,
    )


def select_baseline_page_indices(
    pages: list[PageResult],
    *,
    remaining_budget: float,
    minimum_trial_seconds: float,
    baseline_fraction: float = DEFAULT_BASELINE_BUDGET_FRACTION,
    last_render_seconds: float | None = None,
    last_page_count: int | None = None,
    diagnose_per_page_seconds: float | None = None,
) -> tuple[list[int], str]:
    """Choose which pages to include in the baseline render under a wall budget.

    Returns ``(indices_sorted_by_page, baseline_mode)`` where mode is
    ``\"full\"``, ``\"partial\"``, or ``\"skipped\"``.

    Estimates include **both** soffice render and per-page diagnosis. When a
    full-deck phase would leave too little room for trials, only the highest
    :func:`page_prerank_score` pages are baselined (still written/rendered as
    one contiguous PPTX so pixels match a per-page render of those slides).
    Unselected pages are reported as ``time_budget`` skips.
    """

    count = len(pages)
    if count == 0:
        return [], "skipped"

    remaining = max(0.0, float(remaining_budget))
    minimum = max(0.0, float(minimum_trial_seconds))
    if remaining + 1e-9 < minimum:
        return [], "skipped"

    def _estimate(n: int) -> float:
        return estimate_baseline_phase_seconds(
            n,
            last_render_seconds=last_render_seconds,
            last_page_count=last_page_count,
            diagnose_per_page_seconds=diagnose_per_page_seconds,
        )

    full_est = _estimate(count)
    # Prefer a full deck whenever it leaves room for at least one trial call.
    if full_est + minimum <= remaining + 1e-9:
        return list(range(count)), "full"

    # Otherwise only spend up to a fraction of the budget on baseline+diagnose,
    # and still try to leave one trial slot.
    fraction_cap = remaining * max(0.0, min(1.0, float(baseline_fraction)))
    baseline_cap = min(remaining, max(fraction_cap, remaining - minimum))

    preranked = sorted(
        range(count),
        key=lambda index: (-page_prerank_score(pages[index]), index),
    )
    best_k = 0
    for k in range(1, count + 1):
        est = _estimate(k)
        if est + minimum <= remaining + 1e-9 and est <= baseline_cap + 1e-9:
            best_k = k
        elif best_k == 0 and est <= remaining - minimum + 1e-9:
            best_k = k
            break
        else:
            if best_k > 0:
                break
            if est > remaining + 1e-9:
                break

    if best_k <= 0:
        # Last resort: one page if render+diagnose still leaves a trial crumb,
        # or if we can at least diagnose one page before the trial floor.
        if _estimate(1) + minimum <= remaining + 1e-9 or _estimate(1) <= remaining + 1e-9:
            best_k = 1
        else:
            return [], "skipped"

    selected = sorted(preranked[:best_k])
    mode = "full" if best_k >= count else "partial"
    return selected, mode


def order_pages_for_refine(page_severities: list[tuple[int, float]]) -> list[int]:
    """Order page indices for refine trials.

    Highest diagnosed severity first; page index ascending as a deterministic
    tie-break so early slides win when severities are equal (users care more
    about the start of the deck). Clean pages (severity 0) still appear so
    reports stay complete, but callers typically skip trials when there are
    no issues.
    """

    return [index for index, _severity in sorted(page_severities, key=lambda item: (-item[1], item[0]))]


def page_time_allowance_seconds(
    remaining_seconds: float,
    *,
    pages_left: int,
    minimum_seconds: float,
    cap_at_fair_share: bool = False,
) -> float:
    """Wall-clock slice for the next page under a shared budget.

    Uses an equal share of remaining time among pages still to visit. When
    ``cap_at_fair_share`` is true (first pass), never inflate the slot above
    the equal share — leftover time is left for the reallocation pass instead
    of dumping onto the last page. When false, a slot may be raised to
    ``minimum_seconds`` so a single trial can still start.
    """

    if pages_left <= 0:
        return 0.0
    remaining = max(0.0, float(remaining_seconds))
    minimum = max(0.0, float(minimum_seconds))
    if remaining + 1e-9 < minimum:
        return 0.0
    fair_share = remaining / float(pages_left)
    if cap_at_fair_share:
        return min(remaining, fair_share)
    if fair_share + 1e-9 < minimum:
        return min(remaining, minimum)
    return fair_share


def _source_font_snapshot(page: PageResult) -> dict[str, tuple[float, list[float]]]:
    """Capture pre-refine font sizes so shrink floors stay relative to the source."""

    snapshot: dict[str, tuple[float, list[float]]] = {}
    for block in page.text_blocks:
        block_size = float(block.font_size or 12.0)
        para_sizes = [
            float(para.font_size) if para.font_size is not None else block_size for para in (block.paragraphs or [])
        ]
        snapshot[block.id] = (block_size, para_sizes)
    return snapshot


@dataclass(slots=True)
class _PageTrialState:
    """Mutable per-page trial state shared across first pass and reallocation."""

    page_index: int
    slide_report: SlideRefineReport
    current_metric: float
    current_edit: float
    current_text_f1: float
    fixes_applied: int
    attempts: int
    improving_iters: int
    element_attempts: dict[str, int]
    last_good_digest: str
    source_fonts: dict[str, tuple[float, list[float]]]
    single_path: Path
    has_untried_candidates: bool
    tried_keys: set[tuple[str, str, str]]
    # Actual wall time used (first-pass slot / reallocation / total).
    slot_seconds: float
    time_used_seconds: float
    borrowed_seconds: float
    granted_seconds: float
    page_soffice_before: int
    last_slot_unused: float = 0.0


def refine_pages(
    pages: list[PageResult],
    *,
    pdf_path: Path,
    output_pptx: Path,
    slide_width_pt: float,
    slide_height_pt: float,
    iterations: int = DEFAULT_REFINE_ITERS,
    max_fixes_per_slide: int = DEFAULT_MAX_FIXES_PER_SLIDE,
    max_attempts_per_slide: int = DEFAULT_MAX_ATTEMPTS_PER_SLIDE,
    max_attempts_per_element: int = DEFAULT_MAX_ATTEMPTS_PER_ELEMENT,
    time_budget_seconds: float = DEFAULT_TIME_BUDGET_SECONDS,
    min_font_size_pt: float = DEFAULT_SHRINK_MIN_FONT_SIZE_PT,
    min_font_ratio: float = DEFAULT_MIN_FONT_RATIO,
    dpi: float = DEFAULT_COMPARE_DPI,
    write_pptx: WritePptxFn | None = None,
    render_pptx: RenderPagesFn | None = None,
    render_pdf_page: RenderPdfPageFn | None = None,
    work_dir: Path | None = None,
) -> tuple[list[PageResult], DeckRefineReport]:
    """Run the render-compare-refine loop.

    Returns the (possibly improved) page list and a deck-level report. When
    LibreOffice is missing or fails, restores the last accepted PPTX and
    returns a skipped report instead of raising.

    Budget and ordering (deterministic):

    * **Baseline**: soffice render + per-page diagnosis (both costed). When a
      full-deck phase would leave too little room for trials, only the highest
      conversion-time risk pages are baselined (``baseline_mode=partial``);
      diagnosis also stops when the trial floor would be breached. Others are
      ``pages_skipped`` with reason ``time_budget``. Subset slides stay
      pixel-identical to a per-page render. A started render may overrun by at
      most one call's duration.
    * **Page order**: continuous severity ``40×(1-SSIM) + normalised issues``,
      then ascending page index. Clean pages skip trials.
    * **Per-page share**: first-pass slots capped at equal fair share; trial
      floor is ``max(~2.5s, 1.5× measured per-page render)``.
    * **Reallocation**: round-robin second pass (one attempt per page, severity
      order) while budget and untried candidates remain.
    * **Attempts**: XML no-ops do not count; real attempts are capped per
      element, and the per-element counter **resets after an accepted fix**.
    """

    from .pipeline import render_pages_to_pptx

    started = perf_counter()
    iters = max(0, int(iterations))
    soffice_calls = 0
    total_attempts = 0
    deck_skipped: str | None = None
    last_render_seconds: float | None = None
    last_render_page_count: int = 1
    diagnose_per_page_seconds: float | None = None
    baseline_mode: str | None = None
    baseline_page_count: int | None = None
    pages_skipped: list[dict[str, Any]] = []
    reallocation_passes = 0

    def _elapsed() -> float:
        return perf_counter() - started

    def _remaining_budget() -> float:
        return max(0.0, float(time_budget_seconds) - _elapsed())

    def _trial_minimum() -> float:
        return min_soffice_timeout_seconds(
            last_render_seconds,
            last_render_page_count=last_render_page_count,
        )

    def _report(
        *,
        working_pages: list[PageResult],
        skipped: str | None,
        page_reports: list[SlideRefineReport] | None = None,
    ) -> tuple[list[PageResult], DeckRefineReport]:
        return working_pages, DeckRefineReport(
            enabled=True,
            skipped_reason=skipped,
            iterations=iters,
            dpi=dpi,
            pages=list(page_reports or []),
            runtime_seconds=round(_elapsed(), 3),
            soffice_call_count=soffice_calls,
            attempt_count=total_attempts,
            baseline_mode=baseline_mode,
            baseline_page_count=baseline_page_count,
            pages_skipped=list(pages_skipped),
            reallocation_passes=reallocation_passes,
        )

    if iters <= 0:
        return _report(working_pages=pages, skipped="refine_iters<=0")

    if find_libreoffice_executable() is None and render_pptx is None:
        logger.warning(
            "Refine skipped: LibreOffice (soffice) is not available. Install libreoffice-impress to enable --refine."
        )
        return _report(working_pages=pages, skipped="libreoffice_unavailable")

    write_fn = write_pptx or render_pages_to_pptx
    render_pptx_fn = render_pptx or render_pptx_pages
    render_pdf_fn = render_pdf_page or render_pdf_page_image

    working = _deepcopy_pages(pages)
    owns_work = work_dir is None
    root = Path(tempfile.mkdtemp(prefix="pdf2ppt-refine-")) if owns_work else Path(work_dir or ".")
    root.mkdir(parents=True, exist_ok=True)

    def _safe_render(
        pptx_path: Path,
        label: str,
        *,
        timeout_seconds: int,
        page_count: int = 1,
    ) -> list[Image.Image]:
        nonlocal soffice_calls, last_render_seconds, last_render_page_count
        soffice_calls += 1
        render_started = perf_counter()
        try:
            return render_pptx_fn(
                pptx_path,
                dpi=dpi,
                work_dir=root / label,
                timeout_seconds=timeout_seconds,
            )
        finally:
            last_render_seconds = max(0.0, perf_counter() - render_started)
            last_render_page_count = max(1, int(page_count))

    def _write_full() -> None:
        write_fn(
            working,
            output_pptx,
            slide_width_pt=slide_width_pt,
            slide_height_pt=slide_height_pt,
        )

    def _write_single(page_index: int, path: Path) -> None:
        write_fn(
            [working[page_index]],
            path,
            slide_width_pt=slide_width_pt,
            slide_height_pt=slide_height_pt,
        )

    def _write_subset(page_indices: list[int], path: Path) -> None:
        write_fn(
            [working[index] for index in page_indices],
            path,
            slide_width_pt=slide_width_pt,
            slide_height_pt=slide_height_pt,
        )

    def _trial_timeout_for(available: float) -> int | None:
        """Return a soffice timeout, or ``None`` when the remaining share is too small."""

        minimum = _trial_minimum()
        if available + 1e-9 < minimum:
            return None
        return _soffice_timeout_seconds(available, minimum=minimum)

    try:
        # Always materialise the full deck first (final artifact / restore point).
        _write_full()
        remaining = _remaining_budget()
        minimum = _trial_minimum()
        baseline_indices, baseline_mode = select_baseline_page_indices(
            working,
            remaining_budget=remaining,
            minimum_trial_seconds=minimum,
            diagnose_per_page_seconds=diagnose_per_page_seconds,
        )
        baseline_page_count = len(baseline_indices)
        baseline_set = set(baseline_indices)

        if not baseline_indices:
            for page in working:
                pages_skipped.append(
                    {
                        "page_number": page.page_number,
                        "reason": "time_budget",
                        "detail": "baseline_not_scheduled",
                    }
                )
            empty_reports = [
                SlideRefineReport(
                    page_number=page.page_number,
                    iterations=0,
                    metric_before=None,
                    metric_after=None,
                    skipped_reason="time_budget",
                    severity=round(page_prerank_score(page), 6),
                )
                for page in working
            ]
            return _report(working_pages=working, skipped=None, page_reports=empty_reports)

        for page_index, page in enumerate(working):
            if page_index in baseline_set:
                continue
            pages_skipped.append(
                {
                    "page_number": page.page_number,
                    "reason": "time_budget",
                    "detail": "baseline_not_scheduled",
                }
            )

        if baseline_mode == "full":
            baseline_path = output_pptx
        else:
            baseline_path = root / "baseline-partial.pptx"
            _write_subset(baseline_indices, baseline_path)

        remaining = _remaining_budget()
        minimum = _trial_minimum()
        if remaining + 1e-9 < minimum:
            empty_reports = [
                SlideRefineReport(
                    page_number=page.page_number,
                    iterations=0,
                    metric_before=None,
                    metric_after=None,
                    skipped_reason="time_budget",
                    severity=round(page_prerank_score(page), 6),
                )
                for page in working
            ]
            for page_index, page in enumerate(working):
                if page_index in baseline_set:
                    pages_skipped.append(
                        {
                            "page_number": page.page_number,
                            "reason": "time_budget",
                            "detail": "baseline_not_started",
                        }
                    )
            baseline_mode = "skipped"
            return _report(working_pages=working, skipped=None, page_reports=empty_reports)

        baseline_timeout = _soffice_timeout_seconds(remaining, minimum=minimum)
        try:
            baseline_images = _safe_render(
                baseline_path,
                "baseline-deck" if baseline_mode == "full" else "baseline-partial",
                timeout_seconds=baseline_timeout,
                page_count=len(baseline_indices),
            )
        except RenderCompareError as error:
            logger.warning("Refine skipped: LibreOffice render failed (%s)", error)
            _write_full()
            return _report(working_pages=pages, skipped="render_failed")

        if len(baseline_images) < len(baseline_indices):
            logger.warning(
                "Refine skipped: baseline render returned %s slides, expected %s",
                len(baseline_images),
                len(baseline_indices),
            )
            _write_full()
            return _report(working_pages=pages, skipped="render_failed")

        references: list[Image.Image | None] = [None] * len(working)
        candidates: list[Image.Image | None] = [None] * len(working)
        metrics_before: list[float | None] = [None] * len(working)
        issues_by_page: list[list[RefineIssue]] = [[] for _ in working]
        severities: list[tuple[int, float]] = []
        diagnose_failures: dict[int, str] = {}
        reports_by_index: dict[int, SlideRefineReport] = {}
        diagnosed_indices: set[int] = set()

        for offset, page_index in enumerate(baseline_indices):
            candidates[page_index] = baseline_images[offset]

        # Non-baselined pages are already skipped.
        for page_index, page in enumerate(working):
            if page_index not in baseline_set:
                reports_by_index[page_index] = SlideRefineReport(
                    page_number=page.page_number,
                    iterations=0,
                    metric_before=None,
                    metric_after=None,
                    skipped_reason="time_budget",
                    severity=round(page_prerank_score(page), 6),
                )
                severities.append((page_index, 0.0))

        diagnose_samples: list[float] = []
        for offset, page_index in enumerate(baseline_indices):
            page = working[page_index]
            # Leave room for at least one trial; stop diagnosing early on short budgets.
            if _remaining_budget() + 1e-9 < _trial_minimum():
                for rest_index in baseline_indices[offset:]:
                    if rest_index in reports_by_index:
                        continue
                    reports_by_index[rest_index] = SlideRefineReport(
                        page_number=working[rest_index].page_number,
                        iterations=0,
                        metric_before=None,
                        metric_after=None,
                        skipped_reason="time_budget",
                        severity=round(page_prerank_score(working[rest_index]), 6),
                    )
                    severities.append((rest_index, 0.0))
                    pages_skipped.append(
                        {
                            "page_number": working[rest_index].page_number,
                            "reason": "time_budget",
                            "detail": "diagnose_deadline",
                        }
                    )
                break

            diagnose_started = perf_counter()
            try:
                reference = render_pdf_fn(pdf_path, page_index, dpi=dpi)
            except Exception as error:  # noqa: BLE001
                logger.warning("Refine page %s: PDF render failed (%s)", page.page_number, error)
                diagnose_failures[page_index] = "render_failed"
                severities.append((page_index, 0.0))
                diagnose_samples.append(max(0.0, perf_counter() - diagnose_started))
                continue
            references[page_index] = reference
            candidate = candidates[page_index]
            assert candidate is not None
            metric_before, issues = diagnose_slide(reference, candidate, page)
            metrics_before[page_index] = metric_before
            issues_by_page[page_index] = list(issues)
            diagnosed_indices.add(page_index)
            severity = page_refine_severity(
                page_ssim=metric_before,
                issues=issues,
                page_width_pt=page.width_pt,
                page_height_pt=page.height_pt,
            )
            severities.append((page_index, severity))
            diagnose_samples.append(max(0.0, perf_counter() - diagnose_started))
            if diagnose_samples:
                diagnose_per_page_seconds = sum(diagnose_samples) / float(len(diagnose_samples))

        visit_order = order_pages_for_refine(severities)
        work_order = [
            index
            for index in visit_order
            if index in diagnosed_indices and issues_by_page[index] and index not in diagnose_failures
        ]

        for fail_index, reason in diagnose_failures.items():
            reports_by_index[fail_index] = SlideRefineReport(
                page_number=working[fail_index].page_number,
                iterations=0,
                metric_before=None,
                metric_after=None,
                skipped_reason=reason,
                severity=0.0,
            )

        # Clean baselined pages: record diagnosis, skip trials (no soffice).
        for page_index, page in enumerate(working):
            if page_index in reports_by_index:
                continue
            if issues_by_page[page_index]:
                continue
            metric = metrics_before[page_index]
            severity = next((value for index, value in severities if index == page_index), 0.0)
            reports_by_index[page_index] = SlideRefineReport(
                page_number=page.page_number,
                iterations=0,
                metric_before=metric,
                metric_after=metric,
                issues=[],
                severity=severity,
                slot_seconds=0.0,
                time_used_seconds=0.0,
                borrowed_seconds=0.0,
                accepted_fix_count=0,
            )

        # Mutable trial state retained across first pass + reallocation pass.
        trial_states: dict[int, _PageTrialState] = {}
        severity_by_index = {index: value for index, value in severities}

        def _mark_remaining_time_budget(indices: list[int]) -> None:
            for later_index in indices:
                if later_index in reports_by_index and later_index not in trial_states:
                    continue
                if later_index in trial_states:
                    state = trial_states[later_index]
                    if state.slide_report.skipped_reason is None and state.attempts == 0:
                        state.slide_report.skipped_reason = "time_budget"
                    reports_by_index[later_index] = state.slide_report
                    continue
                if later_index in reports_by_index:
                    continue
                reports_by_index[later_index] = SlideRefineReport(
                    page_number=working[later_index].page_number,
                    iterations=0,
                    metric_before=metrics_before[later_index],
                    metric_after=metrics_before[later_index],
                    issues=list(issues_by_page[later_index]),
                    skipped_reason="time_budget",
                    severity=severity_by_index.get(later_index),
                )
                pages_skipped.append(
                    {
                        "page_number": working[later_index].page_number,
                        "reason": "time_budget",
                        "detail": "trial_not_reached",
                    }
                )

        def _init_trial_state(page_index: int) -> _PageTrialState | None:
            page = working[page_index]
            reference_opt = references[page_index]
            metric_opt = metrics_before[page_index]
            candidate = candidates[page_index]
            if reference_opt is None or metric_opt is None or candidate is None:
                reports_by_index[page_index] = SlideRefineReport(
                    page_number=page.page_number,
                    iterations=0,
                    metric_before=None,
                    metric_after=None,
                    issues=list(issues_by_page[page_index]),
                    skipped_reason="render_failed",
                    severity=severity_by_index.get(page_index),
                )
                return None

            single_path = root / f"page-{page_index}.pptx"
            _write_single(page_index, single_path)
            try:
                edit_before = slide_editability_score(single_path, slide_index=0)
                text_f1_before = slide_text_f1_vs_pdf(single_path, pdf_path, page_index=page_index)
            except Exception:  # noqa: BLE001
                edit_before = 0.0
                text_f1_before = 0.0

            slide_report = SlideRefineReport(
                page_number=page.page_number,
                iterations=0,
                metric_before=metric_opt,
                metric_after=metric_opt,
                issues=list(issues_by_page[page_index]),
                editability_before=edit_before,
                editability_after=edit_before,
                text_f1_before=text_f1_before,
                text_f1_after=text_f1_before,
                severity=severity_by_index.get(page_index),
                slot_seconds=0.0,
                time_used_seconds=0.0,
                borrowed_seconds=0.0,
                granted_seconds=0.0,
                accepted_fix_count=0,
            )
            return _PageTrialState(
                page_index=page_index,
                slide_report=slide_report,
                current_metric=metric_opt,
                current_edit=edit_before,
                current_text_f1=text_f1_before,
                fixes_applied=0,
                attempts=0,
                improving_iters=0,
                element_attempts={},
                last_good_digest=pptx_content_digest(single_path),
                source_fonts=_source_font_snapshot(page),
                single_path=single_path,
                has_untried_candidates=True,
                tried_keys=set(),
                slot_seconds=0.0,
                time_used_seconds=0.0,
                borrowed_seconds=0.0,
                granted_seconds=0.0,
                page_soffice_before=soffice_calls,
                last_slot_unused=0.0,
            )

        def _run_page_trials(
            state: _PageTrialState,
            *,
            allowance: float,
            borrowed: float,
            max_attempts_this_slot: int | None = None,
            is_reallocation: bool = False,
        ) -> None:
            """Run fix trials for one page under a wall-clock allowance."""

            nonlocal total_attempts, deck_skipped

            page_index = state.page_index
            page = working[page_index]
            reference = references[page_index]
            candidate = candidates[page_index]
            assert reference is not None and candidate is not None

            state.granted_seconds += float(allowance)
            state.slide_report.granted_seconds = round(state.granted_seconds, 3)

            page_started = perf_counter()
            page_deadline = page_started + max(0.0, float(allowance))
            attempts_at_slot_start = state.attempts
            attempts_this_slot = 0
            slide_report = state.slide_report
            # Clear a prior first-pass time_budget skip if we got reallocated time.
            if borrowed > 0 and slide_report.skipped_reason == "time_budget":
                slide_report.skipped_reason = None

            def _trial_available_seconds() -> float:
                """Seconds available for the next soffice trial in this slot."""

                remaining_now = _remaining_budget()
                minimum_now = _trial_minimum()
                until_deadline = page_deadline - perf_counter()
                if remaining_now + 1e-9 < minimum_now:
                    return 0.0
                if until_deadline + 1e-9 >= minimum_now:
                    return min(until_deadline, remaining_now)
                # First trial in this slot may use global remainder when the
                # fair-share crumb was burned by setup, but never beyond the
                # remaining wall budget.
                if state.attempts == attempts_at_slot_start:
                    return min(remaining_now, max(minimum_now, until_deadline))
                return 0.0

            made_any_progress = False
            while state.improving_iters < iters and state.fixes_applied < max_fixes_per_slide:
                if state.attempts >= max_attempts_per_slide:
                    state.has_untried_candidates = False
                    break
                if max_attempts_this_slot is not None and attempts_this_slot >= max_attempts_this_slot:
                    break
                available_now = _trial_available_seconds()
                if available_now + 1e-9 < _trial_minimum():
                    slide_report.skipped_reason = slide_report.skipped_reason or "time_budget"
                    break

                _, issues = diagnose_slide(reference, candidate, page, page_ssim=state.current_metric)
                if not issues:
                    state.has_untried_candidates = False
                    break

                made_progress = False
                found_candidate = False
                for issue in issues:
                    if state.attempts >= max_attempts_per_slide or state.fixes_applied >= max_fixes_per_slide:
                        break
                    if state.element_attempts.get(issue.element_id, 0) >= max_attempts_per_element:
                        continue
                    for action in propose_fixes_for_issue(issue):
                        key = (issue.kind, issue.element_id, action)
                        if key in state.tried_keys:
                            continue
                        if state.attempts >= max_attempts_per_slide or state.fixes_applied >= max_fixes_per_slide:
                            break
                        if state.element_attempts.get(issue.element_id, 0) >= max_attempts_per_element:
                            break
                        found_candidate = True
                        available = _trial_available_seconds()
                        trial_timeout = _trial_timeout_for(available)
                        if trial_timeout is None:
                            slide_report.skipped_reason = slide_report.skipped_reason or "time_budget"
                            state.has_untried_candidates = True
                            break

                        source_size: float | None = None
                        source_paras: list[float] | None = None
                        if issue.element_id in state.source_fonts:
                            source_size, source_paras = state.source_fonts[issue.element_id]

                        snapshot = _deepcopy_pages(working)
                        details = apply_fix_action(
                            working[page_index],
                            element_id=issue.element_id,
                            action=action,
                            reference_page_image=reference,
                            min_font_size_pt=min_font_size_pt,
                            min_font_ratio=min_font_ratio,
                            source_font_size=source_size,
                            source_paragraph_sizes=source_paras,
                        )
                        if details is None:
                            state.tried_keys.add(key)
                            continue

                        _write_single(page_index, state.single_path)
                        trial_digest = pptx_content_digest(state.single_path)
                        xml_changed = trial_digest != state.last_good_digest
                        if not xml_changed:
                            working[:] = snapshot
                            page = working[page_index]
                            _write_single(page_index, state.single_path)
                            state.tried_keys.add(key)
                            slide_report.fixes.append(
                                RefineFixAttempt(
                                    issue_kind=issue.kind,
                                    element_id=issue.element_id,
                                    action=action,
                                    accepted=False,
                                    metric_before=state.current_metric,
                                    metric_after=state.current_metric,
                                    details={**details, "reject_reason": "xml_unchanged"},
                                    xml_changed=False,
                                    noop=True,
                                )
                            )
                            continue

                        state.attempts += 1
                        attempts_this_slot += 1
                        total_attempts += 1
                        state.element_attempts[issue.element_id] = state.element_attempts.get(issue.element_id, 0) + 1
                        state.tried_keys.add(key)

                        try:
                            trial_rendered = _safe_render(
                                state.single_path,
                                f"trial-p{page_index}-a{state.attempts}-{action}",
                                timeout_seconds=trial_timeout,
                                page_count=1,
                            )
                        except RenderCompareError as error:
                            logger.warning(
                                "Refine render failed mid-loop (%s); restoring last accepted state",
                                error,
                            )
                            working[:] = snapshot
                            page = working[page_index]
                            _write_full()
                            slide_report.skipped_reason = "render_failed"
                            deck_skipped = "render_failed"
                            slide_report.fixes.append(
                                RefineFixAttempt(
                                    issue_kind=issue.kind,
                                    element_id=issue.element_id,
                                    action=action,
                                    accepted=False,
                                    metric_before=state.current_metric,
                                    metric_after=None,
                                    details={
                                        **details,
                                        "reject_reason": "render_failed",
                                        "soffice_timeout_seconds": trial_timeout,
                                    },
                                    xml_changed=True,
                                    noop=False,
                                )
                            )
                            state.has_untried_candidates = False
                            state.improving_iters = iters
                            break

                        if not trial_rendered:
                            working[:] = snapshot
                            page = working[page_index]
                            _write_single(page_index, state.single_path)
                            slide_report.fixes.append(
                                RefineFixAttempt(
                                    issue_kind=issue.kind,
                                    element_id=issue.element_id,
                                    action=action,
                                    accepted=False,
                                    metric_before=state.current_metric,
                                    metric_after=None,
                                    details={**details, "reject_reason": "missing_rendered_slide"},
                                    xml_changed=True,
                                    noop=False,
                                )
                            )
                            continue

                        trial_candidate = trial_rendered[0]
                        trial_metric = compute_ssim(reference, resize_to_match(reference, trial_candidate))
                        try:
                            trial_edit = slide_editability_score(state.single_path, slide_index=0)
                            trial_text_f1 = slide_text_f1_vs_pdf(state.single_path, pdf_path, page_index=page_index)
                        except Exception:  # noqa: BLE001
                            trial_edit = state.current_edit
                            trial_text_f1 = state.current_text_f1

                        ssim_improved = trial_metric > state.current_metric + METRIC_IMPROVE_EPSILON
                        edit_ok = trial_edit + EDITABILITY_EPSILON >= state.current_edit
                        text_ok = trial_text_f1 + TEXT_F1_EPSILON >= state.current_text_f1
                        accepted = ssim_improved and edit_ok and text_ok

                        reject_reason = None
                        if not ssim_improved:
                            reject_reason = "ssim_not_improved"
                        elif not edit_ok:
                            reject_reason = "editability_regressed"
                        elif not text_ok:
                            reject_reason = "text_f1_regressed"

                        slide_report.fixes.append(
                            RefineFixAttempt(
                                issue_kind=issue.kind,
                                element_id=issue.element_id,
                                action=action,
                                accepted=accepted,
                                metric_before=state.current_metric,
                                metric_after=trial_metric,
                                details={
                                    **details,
                                    **({"reject_reason": reject_reason} if reject_reason is not None else {}),
                                    "editability_before": state.current_edit,
                                    "editability_after": trial_edit,
                                    "text_f1_before": state.current_text_f1,
                                    "text_f1_after": trial_text_f1,
                                },
                                xml_changed=True,
                                noop=False,
                            )
                        )
                        if accepted:
                            state.current_metric = trial_metric
                            state.current_edit = trial_edit
                            state.current_text_f1 = trial_text_f1
                            candidate = trial_candidate
                            candidates[page_index] = trial_candidate
                            page = working[page_index]
                            state.last_good_digest = trial_digest
                            state.fixes_applied += 1
                            state.improving_iters += 1
                            made_progress = True
                            made_any_progress = True
                            state.element_attempts[issue.element_id] = 0
                            break

                        working[:] = snapshot
                        page = working[page_index]
                        _write_single(page_index, state.single_path)
                        if max_attempts_this_slot is not None and attempts_this_slot >= max_attempts_this_slot:
                            break
                    else:
                        continue
                    break

                if slide_report.skipped_reason in {"render_failed", "time_budget"}:
                    break
                if max_attempts_this_slot is not None and attempts_this_slot >= max_attempts_this_slot:
                    break
                if not made_progress:
                    if not found_candidate:
                        state.has_untried_candidates = False
                    break

            slot_used = max(0.0, perf_counter() - page_started)
            state.time_used_seconds += slot_used
            state.last_slot_unused = max(0.0, float(allowance) - slot_used)
            if is_reallocation:
                state.borrowed_seconds += slot_used
            else:
                state.slot_seconds += slot_used
            slide_report.iterations = state.improving_iters
            slide_report.metric_after = state.current_metric
            slide_report.editability_after = state.current_edit
            slide_report.text_f1_after = state.current_text_f1
            slide_report.attempt_count = state.attempts
            slide_report.soffice_call_count = soffice_calls - state.page_soffice_before
            slide_report.slot_seconds = round(state.slot_seconds, 3)
            slide_report.borrowed_seconds = round(state.borrowed_seconds, 3)
            slide_report.time_used_seconds = round(state.time_used_seconds, 3)
            slide_report.granted_seconds = round(state.granted_seconds, 3)
            slide_report.accepted_fix_count = sum(1 for fix in slide_report.fixes if fix.accepted)
            if (
                not state.has_untried_candidates
                and slide_report.skipped_reason == "time_budget"
                and (state.attempts > 0 or made_any_progress)
            ):
                # Candidates exhausted under the slot; keep telemetry but prefer
                # not to look like a hard skip when work completed normally.
                slide_report.skipped_reason = None
            reports_by_index[page_index] = slide_report

        # ----- First pass: severity order, fair-share capped slots -----
        # Snapshot remaining once so later pages cannot inherit all leftover time
        # (priority inversion). Unused crumbs go to the reallocation pass.
        first_pass_budget = _remaining_budget()
        first_pass_pages = len(work_order)
        first_pass_fair = first_pass_budget / float(first_pass_pages) if first_pass_pages > 0 else 0.0

        for order_pos, page_index in enumerate(work_order):
            remaining = _remaining_budget()
            minimum_now = _trial_minimum()
            if remaining + 1e-9 < minimum_now:
                _mark_remaining_time_budget(work_order[order_pos:])
                break

            allowance = min(first_pass_fair, remaining)
            # Cap at the initial fair share; do not inflate to dump leftovers on
            # the last page. If the crumb is below the trial floor, defer to
            # reallocation instead of skipping the rest of the queue.
            if allowance + 1e-9 < minimum_now:
                continue

            state = _init_trial_state(page_index)
            if state is None:
                continue

            remaining_after_setup = _remaining_budget()
            minimum_after_setup = _trial_minimum()
            if remaining_after_setup + 1e-9 < minimum_after_setup:
                state.slide_report.skipped_reason = "time_budget"
                reports_by_index[page_index] = state.slide_report
                trial_states[page_index] = state
                continue

            allowance = min(allowance, first_pass_fair, remaining_after_setup)
            if allowance + 1e-9 < minimum_after_setup:
                # Setup burned the fair crumb; leave the page for reallocation.
                trial_states[page_index] = state
                reports_by_index[page_index] = state.slide_report
                continue

            trial_states[page_index] = state
            _run_page_trials(state, allowance=allowance, borrowed=0.0, is_reallocation=False)

            if deck_skipped == "render_failed":
                for later_index in work_order[order_pos + 1 :]:
                    if later_index in reports_by_index:
                        continue
                    reports_by_index[later_index] = SlideRefineReport(
                        page_number=working[later_index].page_number,
                        iterations=0,
                        metric_before=metrics_before[later_index],
                        metric_after=metrics_before[later_index],
                        issues=list(issues_by_page[later_index]),
                        skipped_reason="render_failed",
                        severity=severity_by_index.get(later_index),
                    )
                break

        # ----- Second pass: round-robin one attempt per page by severity -----
        if deck_skipped != "render_failed":
            while True:
                remaining = _remaining_budget()
                minimum_now = _trial_minimum()
                if remaining + 1e-9 < minimum_now:
                    break

                candidates_left: list[int] = []
                for index, state in trial_states.items():
                    if (
                        state.has_untried_candidates
                        and state.attempts < max_attempts_per_slide
                        and state.fixes_applied < max_fixes_per_slide
                        and state.improving_iters < iters
                        and state.slide_report.skipped_reason != "render_failed"
                    ):
                        candidates_left.append(index)
                for index in work_order:
                    if index in trial_states or index in diagnose_failures:
                        continue
                    if not issues_by_page[index]:
                        continue
                    existing = reports_by_index.get(index)
                    if existing is not None and existing.skipped_reason == "render_failed":
                        continue
                    candidates_left.append(index)

                seen: set[int] = set()
                ordered_left: list[int] = []
                for index in order_pages_for_refine(
                    [(idx, severity_by_index.get(idx, 0.0)) for idx in candidates_left]
                ):
                    if index in seen:
                        continue
                    seen.add(index)
                    ordered_left.append(index)

                if not ordered_left:
                    break

                reallocation_passes += 1
                progressed = False
                for page_index in ordered_left:
                    remaining = _remaining_budget()
                    minimum_now = _trial_minimum()
                    if remaining + 1e-9 < minimum_now:
                        break
                    allowance = min(remaining, minimum_now)
                    if page_index not in trial_states:
                        state = _init_trial_state(page_index)
                        if state is None:
                            continue
                        trial_states[page_index] = state
                    else:
                        state = trial_states[page_index]
                    before_attempts = state.attempts
                    _run_page_trials(
                        state,
                        allowance=allowance,
                        borrowed=allowance,
                        max_attempts_this_slot=1,
                        is_reallocation=True,
                    )
                    if state.attempts > before_attempts:
                        progressed = True
                    if deck_skipped == "render_failed":
                        break
                if deck_skipped == "render_failed":
                    break
                if not progressed:
                    break

        # Emit reports in original page order for stable JSON.
        page_reports = [
            reports_by_index.get(
                index,
                SlideRefineReport(
                    page_number=working[index].page_number,
                    iterations=0,
                    metric_before=metrics_before[index],
                    metric_after=metrics_before[index],
                    issues=list(issues_by_page[index]),
                    skipped_reason="time_budget",
                    severity=severity_by_index.get(index),
                ),
            )
            for index in range(len(working))
        ]
        for report in page_reports:
            if report.skipped_reason == "time_budget":
                already = any(entry.get("page_number") == report.page_number for entry in pages_skipped)
                if not already:
                    pages_skipped.append(
                        {
                            "page_number": report.page_number,
                            "reason": "time_budget",
                            "detail": "trial_or_baseline",
                        }
                    )

        _write_full()
        return _report(working_pages=working, skipped=deck_skipped, page_reports=page_reports)
    except RenderCompareError as error:
        logger.warning("Refine skipped: LibreOffice render failed (%s)", error)
        try:
            _write_full()
        except Exception:  # noqa: BLE001
            logger.warning("Could not restore PPTX after refine render failure", exc_info=True)
        return _report(working_pages=pages, skipped="render_failed")
    finally:
        if owns_work:
            import shutil

            shutil.rmtree(root, ignore_errors=True)


def write_refine_report(report: DeckRefineReport, path: Path) -> Path:
    """Write the per-slide refine JSON report."""

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return output


__all__ = [
    "DEFAULT_MAX_ATTEMPTS_PER_ELEMENT",
    "DEFAULT_MAX_ATTEMPTS_PER_SLIDE",
    "DEFAULT_MAX_FIXES_PER_SLIDE",
    "DEFAULT_MIN_FONT_RATIO",
    "DEFAULT_MIN_SOFFICE_TIMEOUT_SECONDS",
    "DEFAULT_REFINE_ITERS",
    "DEFAULT_SHRINK_MIN_FONT_SIZE_PT",
    "DEFAULT_TIME_BUDGET_SECONDS",
    "METRIC_IMPROVE_EPSILON",
    "DeckRefineReport",
    "RefineFixAttempt",
    "RefineIssue",
    "SlideRefineReport",
    "apply_fix_action",
    "apply_match_line_bboxes",
    "apply_raster_fallback",
    "apply_shrink_font",
    "apply_widen_for_wrap",
    "diagnose_extra_wrap",
    "diagnose_line_drift",
    "diagnose_low_local_ssim",
    "diagnose_slide",
    "diagnose_text_overflow",
    "DEFAULT_DIAGNOSE_PER_PAGE_SECONDS",
    "SEVERITY_FIDELITY_WEIGHT",
    "estimate_baseline_phase_seconds",
    "estimate_baseline_render_seconds",
    "estimate_diagnose_seconds",
    "min_soffice_timeout_seconds",
    "order_pages_for_refine",
    "page_issue_severity",
    "page_prerank_score",
    "page_refine_severity",
    "page_time_allowance_seconds",
    "pptx_content_digest",
    "propose_fixes_for_issue",
    "refine_pages",
    "region_ssim",
    "select_baseline_page_indices",
    "shrink_font_floor",
    "write_refine_report",
]
