"""P2 layout analysis tests — heuristic path + mocked PP-StructureV3."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from PIL import Image
from pptx import Presentation

from pdf2ppt.core import ConversionOptions
from pdf2ppt.layout import (
    LayoutEngine,
    LayoutPageResult,
    LayoutRegion,
    analyze_layout_heuristic,
    apply_layout_to_page,
    normalize_layout_label,
    parse_layout_regions_from_dict,
    resolve_layout_model_dir,
    strip_bullet_prefix,
)
from pdf2ppt.layout_metrics import (
    count_pptx_bullet_paragraphs,
    evaluate_text_grouping,
    load_layout_ground_truth,
    slide_has_title_placeholder_text,
)
from pdf2ppt.models import TextBlock, TextParagraph
from pdf2ppt.ppt_render import add_text_block, render_page_to_slide

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _block(
    block_id: str,
    text: str,
    bbox: tuple[float, float, float, float],
    *,
    role: str = "body",
    font_size: float = 14.0,
) -> TextBlock:
    return TextBlock(
        id=block_id,
        source="native",
        bbox=bbox,
        text=text,
        confidence=0.99,
        font_size=font_size,
        block_role=role,
    )


class NormalizeLabelTests(unittest.TestCase):
    def test_maps_ppstructure_labels(self) -> None:
        self.assertEqual(normalize_layout_label("doc_title"), "title")
        self.assertEqual(normalize_layout_label("figure_title"), "caption")
        self.assertEqual(normalize_layout_label("page_number"), "page_number")
        self.assertEqual(normalize_layout_label("unknown_xyz"), "other")


class HeuristicLayoutTests(unittest.TestCase):
    def test_keeps_flowchart_labels_separate(self) -> None:
        blocks = [
            _block("t", "Onboarding Flow", (40, 24, 600, 70), role="title", font_size=24),
            _block("a", "Collect Info", (128, 268, 312, 312), font_size=16),
            _block("b", "Verify", (528, 268, 752, 312), font_size=16),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        texts = {block.text for block in applied.text_blocks}
        self.assertIn("Collect Info", texts)
        self.assertIn("Verify", texts)
        merged = [block for block in applied.text_blocks if "Collect" in block.text and "Verify" in block.text]
        self.assertEqual(merged, [])

    def test_groups_stacked_list_lines(self) -> None:
        blocks = [
            _block("t", "Sprint Goals", (48, 36, 400, 80), role="title", font_size=28),
            _block("l1", "• Ship layout analysis", (72, 130, 700, 160)),
            _block("l2", "◦ Detect multi-level bullets", (96, 170, 700, 200)),
            _block("l3", "• Keep heuristic fallback", (72, 250, 700, 280)),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            skip_header_footer=True,
        )
        list_blocks = [block for block in applied.text_blocks if block.block_role == "list"]
        self.assertTrue(list_blocks)
        paragraphs = list_blocks[0].paragraphs or []
        self.assertGreaterEqual(len(paragraphs), 2)
        self.assertTrue(any(para.list_level is not None for para in paragraphs))
        # Literal bullet glyphs stripped from paragraph text.
        self.assertFalse(any(para.text.startswith(("•", "◦", "-", "*")) for para in paragraphs))

    def test_skips_ocr_footer_and_page_number(self) -> None:
        blocks = [
            _block("body", "Hello", (60, 120, 400, 160)),
            TextBlock(
                id="foot",
                source="ocr",
                bbox=(48, 510, 300, 535),
                text="Confidential",
                confidence=0.9,
                font_size=10,
                image_bbox=(48, 510, 300, 535),
            ),
            TextBlock(
                id="num",
                source="ocr",
                bbox=(880, 510, 940, 535),
                text="3",
                confidence=0.9,
                font_size=10,
                image_bbox=(880, 510, 940, 535),
            ),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            skip_header_footer=True,
        )
        texts = [block.text for block in applied.text_blocks]
        self.assertIn("Hello", texts)
        self.assertNotIn("Confidential", texts)
        self.assertNotIn("3", texts)

    def test_keeps_native_footer_visible_on_digital(self) -> None:
        blocks = [
            _block("body", "Hello", (60, 120, 400, 160)),
            _block("foot", "Confidential", (48, 510, 300, 535), font_size=10),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            skip_header_footer=True,
        )
        texts = [block.text for block in applied.text_blocks]
        self.assertIn("Confidential", texts)

    def test_chart_axis_labels_not_footer(self) -> None:
        # Axis text deep in the footer band would match y0>=0.93*H, but a chart
        # shape within 40pt above must block the footer label.
        blocks = [
            _block("t", "Monthly Active Users", (48, 36, 500, 80), role="title", font_size=24),
            _block("axis", "Jan  Feb  Mar  Apr  May  Jun  Jul", (120, 505, 800, 530), font_size=11),
        ]
        shapes = [(140.0, 200.0, 200.0, 470.0), (220.0, 150.0, 280.0, 470.0)]
        regions = analyze_layout_heuristic(
            blocks,
            page_width=960,
            page_height=540,
            shape_bboxes=shapes,
        )
        labels = {region.label for region in regions}
        self.assertNotIn("footer", labels)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            skip_header_footer=True,
        )
        self.assertTrue(any("Jan" in block.text for block in applied.text_blocks))

    def test_shallow_bottom_labels_not_footer_without_tiny_type(self) -> None:
        # geometric_chart axis at ~0.89*H must never be footer.
        blocks = [
            _block("axis", "Jan  Feb  Mar  Apr  May  Jun  Jul", (120, 480, 800, 500), font_size=14),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        self.assertEqual(regions[0].label, "text")

    def test_strip_bullet_prefix_levels(self) -> None:
        cleaned, level, style, prefix = strip_bullet_prefix("• Ship it", left=72, min_left=72)
        self.assertEqual(cleaned, "Ship it")
        self.assertEqual(level, 0)
        self.assertEqual(style, "bullet")
        self.assertTrue(prefix.startswith("•"))
        cleaned, level, style, prefix = strip_bullet_prefix("1. First", left=96, min_left=72)
        self.assertEqual(cleaned, "First")
        self.assertEqual(level, 1)
        self.assertEqual(style, "arabicPeriod")
        self.assertTrue(prefix.startswith("1."))
        cleaned, level, style, prefix = strip_bullet_prefix("a. Nested", left=96, min_left=72)
        self.assertEqual(cleaned, "Nested")
        self.assertEqual(style, "alphaLcPeriod")
        # Modest nested indent must map to level 1, not jump to 2 via rounding.
        cleaned, level, style, prefix = strip_bullet_prefix("a. Nested", left=99, min_left=72)
        self.assertEqual(level, 1)
        self.assertEqual(style, "alphaLcPeriod")


class PpstructureMockTests(unittest.TestCase):
    def test_ppstructure_parses_mocked_boxes(self) -> None:
        image = Image.new("RGB", (960, 540), color=(255, 255, 255))
        fake_boxes = [
            {"cls_id": 0, "label": "doc_title", "score": 0.99, "coordinate": [48, 36, 500, 90]},
            {"cls_id": 2, "label": "text", "score": 0.95, "coordinate": [60, 140, 450, 400]},
            {"cls_id": 2, "label": "text", "score": 0.94, "coordinate": [510, 140, 900, 400]},
            {"cls_id": 4, "label": "figure", "score": 0.91, "coordinate": [180, 110, 780, 380]},
            {"cls_id": 6, "label": "figure_title", "score": 0.88, "coordinate": [180, 395, 780, 440]},
            {"cls_id": 5, "label": "footer", "score": 0.8, "coordinate": [40, 500, 200, 530]},
        ]

        class FakeResult:
            def __init__(self) -> None:
                self.json = {"res": {"boxes": fake_boxes}}

        class FakeDetector:
            def predict(self, *_args, **_kwargs):
                return [FakeResult()]

        engine = LayoutEngine(engine="ppstructure")
        with patch.object(LayoutEngine, "_get_detector", return_value=FakeDetector()):
            result = engine.analyze(image, page_width=960, page_height=540, page_number=1)

        self.assertEqual(result.engine, "ppstructure")
        labels = {region.label for region in result.regions}
        self.assertIn("title", labels)
        self.assertIn("figure", labels)
        self.assertIn("caption", labels)
        self.assertIn("footer", labels)
        # Image coords map 1:1 when image size matches page points.
        title = next(region for region in result.regions if region.label == "title")
        self.assertAlmostEqual(title.bbox[0], 48, places=0)

    def test_ppstructure_falls_back_when_detector_fails(self) -> None:
        blocks = [_block("a", "Hello", (40, 40, 200, 80))]
        engine = LayoutEngine(engine="ppstructure")
        with patch.object(LayoutEngine, "_get_detector", side_effect=RuntimeError("no paddle")):
            result = engine.analyze(
                None,
                page_width=960,
                page_height=540,
                text_blocks=blocks,
            )
        self.assertEqual(result.engine, "heuristic")
        self.assertIsNotNone(result.fallback_reason)

    def test_parse_layout_regions_from_dict(self) -> None:
        regions = parse_layout_regions_from_dict(
            [{"label": "list", "bbox": [10, 20, 100, 200], "score": 0.9}],
            page_number=2,
        )
        self.assertEqual(len(regions), 1)
        self.assertEqual(regions[0].label, "list")
        self.assertEqual(regions[0].reading_order, 1)


class ApplyLayoutFigureTableTests(unittest.TestCase):
    def test_crops_figure_regions(self) -> None:
        import numpy as np

        array = np.full((540, 960, 3), (30, 90, 160), dtype=np.uint8)
        array[110:380, 180:780] = (200, 180, 40)
        image = Image.fromarray(array, mode="RGB")
        regions = [
            LayoutRegion(
                id="fig1",
                label="figure",
                bbox=(180, 110, 780, 380),
                confidence=0.95,
                reading_order=1,
                engine="ppstructure",
            )
        ]
        applied = apply_layout_to_page(
            [],
            LayoutPageResult(regions=regions, engine="ppstructure"),
            page_width=960,
            page_height=540,
            page_image=image,
        )
        self.assertEqual(len(applied.figure_images), 1)
        self.assertTrue(applied.figure_images[0].png_bytes.startswith(b"\x89PNG"))

    def test_defers_text_inside_native_table(self) -> None:
        from pdf2ppt.models import TableCell, TableElement

        table = TableElement(
            id="t1",
            bbox=(100, 100, 500, 300),
            row_count=1,
            col_count=1,
            cells=[TableCell(row=0, col=0, text="Cell", text_block_ids=["n1"])],
            column_widths=[400],
            row_heights=[200],
        )
        blocks = [_block("n1", "Cell", (120, 120, 400, 180))]
        regions = [
            LayoutRegion(
                id="table",
                label="table",
                bbox=(100, 100, 500, 300),
                confidence=0.9,
                reading_order=1,
            )
        ]
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            table_elements=[table],
        )
        self.assertEqual(applied.text_blocks, [])
        self.assertIn("n1", applied.skipped_block_ids)


class BulletRenderTests(unittest.TestCase):
    def test_set_paragraph_bullet_emits_bu_char(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        block = TextBlock(
            id="list1",
            source="native",
            bbox=(72, 120, 600, 280),
            text="One\nTwo",
            confidence=1.0,
            font_size=16,
            block_role="list",
            paragraphs=[
                TextParagraph(text="One", list_level=0),
                TextParagraph(text="Two", list_level=1),
            ],
        )
        add_text_block(slide, block, scale_x=1.0, scale_y=1.0)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bullets.pptx"
            presentation.save(path)
            self.assertGreaterEqual(count_pptx_bullet_paragraphs(path), 2)

    def test_title_placeholder_mapping(self) -> None:
        from pdf2ppt.models import PageResult, QualityScore

        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        page = PageResult(
            page_number=1,
            page_kind="digital",
            background_mode="elements",
            width_pt=960,
            height_pt=540,
            text_blocks=[
                TextBlock(
                    id="title",
                    source="native",
                    bbox=(60, 40, 700, 100),
                    text="Platform Overview",
                    confidence=1.0,
                    font_size=28,
                    block_role="title",
                    use_title_placeholder=True,
                    reading_order=1,
                ),
                TextBlock(
                    id="body",
                    source="native",
                    bbox=(60, 140, 400, 300),
                    text="Body copy",
                    confidence=1.0,
                    font_size=14,
                    reading_order=2,
                ),
            ],
            quality_score=QualityScore(1, 1, 1, 1),
            fallback_reason=None,
        )
        render_page_to_slide(
            presentation,
            blank,
            page,
            slide_width_pt=960,
            slide_height_pt=540,
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "title.pptx"
            presentation.save(path)
            self.assertTrue(slide_has_title_placeholder_text(path))


class LayoutMetricsFixtureTests(unittest.TestCase):
    def test_flowchart_grouping_metrics_no_merge(self) -> None:
        gt = load_layout_ground_truth(FIXTURES / "flowchart.layout.json")
        blocks = [
            _block("t", "Onboarding Flow", (40, 24, 600, 70), role="title", font_size=24),
            _block("s", "Start", (128, 138, 312, 182), font_size=16),
            _block("c", "Collect Info", (128, 268, 312, 312), font_size=16),
            _block("v", "Verify", (528, 268, 752, 312), font_size=16),
            _block("d", "Done", (328, 418, 552, 462), font_size=16),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        metrics = evaluate_text_grouping(applied.text_blocks, gt)
        self.assertEqual(metrics.merge_error_count, 0)
        self.assertTrue(metrics.title_correct)
        self.assertGreaterEqual(metrics.mean_paragraph_iou, 0.35)

    def test_bullet_list_fixture_metrics(self) -> None:
        gt = load_layout_ground_truth(FIXTURES / "bullet_list.layout.json")
        blocks = [
            _block("t", "Sprint Goals", (48, 36, 800, 90), role="title", font_size=28),
            _block("a", "• Ship layout analysis for scanned decks", (72, 130, 880, 164)),
            _block("b", "◦ Detect multi-level bullets", (96, 170, 880, 204)),
            _block("c", "◦ Map titles to placeholders", (96, 205, 880, 239)),
            _block("d", "• Keep heuristic fallback for CI", (72, 250, 880, 284)),
            _block("e", "◦ No model download required", (96, 290, 880, 324)),
            _block("f", "• Report grouping metrics in the PR", (72, 340, 880, 374)),
            TextBlock(
                id="foot",
                source="ocr",
                bbox=(48, 510, 400, 535),
                text="Confidential",
                confidence=0.9,
                font_size=10,
            ),
            TextBlock(
                id="num",
                source="ocr",
                bbox=(880, 510, 940, 535),
                text="3",
                confidence=0.9,
                font_size=10,
            ),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
            skip_header_footer=True,
        )
        metrics = evaluate_text_grouping(applied.text_blocks, gt)
        self.assertTrue(metrics.title_correct)
        self.assertIsNotNone(metrics.list_f1)
        assert metrics.list_f1 is not None
        self.assertGreaterEqual(metrics.list_f1, 0.5)
        # OCR footer/page number stay in the raster background (skipped as text).
        joined = " ".join(block.text for block in applied.text_blocks)
        self.assertNotIn("Confidential", joined)


class PipelineLayoutOptionTests(unittest.TestCase):
    def test_conversion_options_default_heuristic(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            options = ConversionOptions(
                input_path=Path(tmp) / "in.pdf",
                output_path=Path(tmp) / "out.pptx",
                report_path=Path(tmp) / "out.json",
            )
        self.assertEqual(options.layout_engine, "heuristic")
        self.assertTrue(options.layout_skip_header_footer)

    def test_analyze_page_records_layout_engine(self) -> None:
        from pdf2ppt.pipeline import analyze_page

        pdf_path = FIXTURES / "flowchart.pdf"
        import fitz

        options = ConversionOptions(
            input_path=pdf_path,
            output_path=pdf_path.with_suffix(".pptx"),
            report_path=pdf_path.with_suffix(".json"),
            vector_mode="off",
            layout_engine="heuristic",
        )
        ocr = MagicMock()
        with fitz.open(pdf_path) as document:
            page = document[0]
            result = analyze_page(page, options, ocr)
        self.assertEqual(result.layout_engine, "heuristic")
        self.assertTrue(result.layout_regions)
        texts = [block.text for block in result.text_blocks]
        self.assertTrue(any("Collect Info" in text for text in texts))
        self.assertTrue(any("Verify" in text for text in texts))
        self.assertFalse(any("Collect Info" in text and "Verify" in text for text in texts))


class ModelDirResolutionTests(unittest.TestCase):
    def test_resolve_layout_model_dir_none_when_missing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            self.assertIsNone(resolve_layout_model_dir(Path(tmp) / "model"))
            self.assertIsNone(resolve_layout_model_dir(Path(tmp)))

    def test_resolve_layout_model_dir_named_folder(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            named = Path(tmp) / "PP-DocLayout_plus-L"
            named.mkdir()
            (named / "inference.yml").write_text("model: ok\n", encoding="utf-8")
            self.assertEqual(resolve_layout_model_dir(tmp), str(named))

    def test_get_detector_omits_model_dir_when_absent(self) -> None:
        captured: dict[str, object] = {}

        class FakeDetector:
            def __init__(self, **kwargs):
                captured.update(kwargs)

            def predict(self, *_args, **_kwargs):
                return []

        with tempfile.TemporaryDirectory() as tmp:
            engine = LayoutEngine(engine="ppstructure", model_root=Path(tmp) / "missing-model")
            with patch("paddleocr.LayoutDetection", FakeDetector):
                # Clear process cache so this root is exercised.
                from pdf2ppt import layout as layout_mod

                layout_mod._DETECTOR_CACHE.clear()
                engine._detector = None
                engine._get_detector()
        self.assertNotIn("model_dir", captured)
        self.assertEqual(captured.get("model_name"), "PP-DocLayout_plus-L")


class RealModelOptInTests(unittest.TestCase):
    """Real PP-StructureV3 download is opt-in via PDF2PPT_RUN_LAYOUT_MODEL=1."""

    def test_real_model_skipped_by_default(self) -> None:
        import os

        if os.environ.get("PDF2PPT_RUN_LAYOUT_MODEL") != "1":
            self.skipTest("Set PDF2PPT_RUN_LAYOUT_MODEL=1 to run the real layout model")
        from pdf2ppt.layout import is_ppstructure_available

        if not is_ppstructure_available():
            self.skipTest("LayoutDetection not importable")
        # Default model_root is ./model (OCR). Without DocLayout weights there,
        # detector must still start via PaddleX's default cache — not fall back.
        image = Image.new("RGB", (640, 480), color=(255, 255, 255))
        engine = LayoutEngine(engine="ppstructure", model_root=Path("model"))
        result = engine.analyze(image, page_width=640, page_height=480)
        self.assertEqual(result.engine, "ppstructure")


class TitleZOrderTests(unittest.TestCase):
    def test_title_placeholder_above_background_picture(self) -> None:
        from pptx.oxml.ns import qn

        from pdf2ppt.models import PageResult, QualityScore
        from pdf2ppt.ppt_render import render_page_to_slide

        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        # Minimal PNG background.
        import io

        buffer = io.BytesIO()
        Image.new("RGB", (16, 16), color=(40, 80, 120)).save(buffer, format="PNG")
        page = PageResult(
            page_number=1,
            page_kind="digital",
            background_mode="overlay",
            width_pt=960,
            height_pt=540,
            background_image_bytes=buffer.getvalue(),
            text_blocks=[
                TextBlock(
                    id="title",
                    source="native",
                    bbox=(60, 40, 700, 100),
                    text="Platform Overview",
                    confidence=1.0,
                    font_size=28,
                    block_role="title",
                    use_title_placeholder=True,
                    reading_order=1,
                    bold=False,
                ),
            ],
            quality_score=QualityScore(1, 1, 1, 1),
            fallback_reason=None,
        )
        render_page_to_slide(presentation, blank, page, slide_width_pt=960, slide_height_pt=540)
        slide = presentation.slides[0]
        tree = slide.shapes._spTree  # noqa: SLF001
        children = list(tree)
        title_el = slide.shapes.title._element  # noqa: SLF001
        pic_els = [child for child in children if child.tag == qn("p:pic")]
        self.assertTrue(pic_els)
        self.assertGreater(children.index(title_el), children.index(pic_els[0]))

    def test_page_background_shape_dropped_when_raster_present(self) -> None:
        """figure_caption: full-page white fill must not cover the overlay picture."""

        from pptx.oxml.ns import qn

        from pdf2ppt.models import PageResult, QualityScore, ShapeElement
        from pdf2ppt.ppt_render import render_page_to_slide

        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        import io

        buffer = io.BytesIO()
        Image.new("RGB", (16, 16), color=(200, 180, 40)).save(buffer, format="PNG")
        page = PageResult(
            page_number=1,
            page_kind="digital",
            background_mode="overlay",
            width_pt=960,
            height_pt=540,
            background_image_bytes=buffer.getvalue(),
            shape_elements=[
                ShapeElement(
                    id="bg_rect",
                    kind="rect",
                    bbox=(0, 0, 960, 540),
                    fill_color="#FFFFFF",
                    is_page_background=True,
                    z_order=0,
                )
            ],
            text_blocks=[
                TextBlock(
                    id="title",
                    source="native",
                    bbox=(60, 40, 700, 100),
                    text="Architecture Sketch",
                    confidence=1.0,
                    font_size=28,
                    block_role="title",
                    use_title_placeholder=True,
                    reading_order=1,
                ),
            ],
            quality_score=QualityScore(1, 1, 1, 1),
            fallback_reason=None,
        )
        render_page_to_slide(presentation, blank, page, slide_width_pt=960, slide_height_pt=540)
        slide = presentation.slides[0]
        children = list(slide.shapes._spTree)  # noqa: SLF001
        title_el = slide.shapes.title._element  # noqa: SLF001
        title_idx = children.index(title_el)
        pic_idxs = [i for i, child in enumerate(children) if child.tag == qn("p:pic")]
        # Page-background fills are omitted when a full-slide raster already
        # includes them; only the picture (+ title) should remain.
        sp_idxs = [
            i
            for i, child in enumerate(children)
            if child.tag == qn("p:sp") and child is not title_el
        ]
        self.assertTrue(pic_idxs)
        self.assertEqual(sp_idxs, [])
        self.assertGreater(title_idx, max(pic_idxs))

    def test_page_background_shape_below_foreground_without_raster(self) -> None:
        """Elements mode: page fill stays at the bottom of z-order."""

        from pptx.oxml.ns import qn

        from pdf2ppt.models import PageResult, QualityScore, ShapeElement
        from pdf2ppt.ppt_render import render_page_to_slide

        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        page = PageResult(
            page_number=1,
            page_kind="digital",
            background_mode="elements",
            width_pt=960,
            height_pt=540,
            shape_elements=[
                ShapeElement(
                    id="bg_rect",
                    kind="rect",
                    bbox=(0, 0, 960, 540),
                    fill_color="#FFFFFF",
                    is_page_background=True,
                    z_order=0,
                ),
                ShapeElement(
                    id="fg_rect",
                    kind="rect",
                    bbox=(100, 100, 300, 200),
                    fill_color="#336699",
                    z_order=2,
                ),
            ],
            text_blocks=[],
            quality_score=QualityScore(1, 1, 1, 1),
            fallback_reason=None,
        )
        render_page_to_slide(presentation, blank, page, slide_width_pt=960, slide_height_pt=540)
        slide = presentation.slides[0]
        children = list(slide.shapes._spTree)  # noqa: SLF001
        sp_idxs = [i for i, child in enumerate(children) if child.tag == qn("p:sp")]
        self.assertEqual(len(sp_idxs), 2)
        # First auto-shape is the page background; second is the foreground rect.
        names = []
        for child in children:
            if child.tag != qn("p:sp"):
                continue
            nv = child.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
            names.append(nv.get("name") if nv is not None else "")
        self.assertEqual(len(names), 2)
        self.assertLess(sp_idxs[0], sp_idxs[1])


class GroupedOcrMaskTests(unittest.TestCase):
    def test_grouped_ocr_lines_union_image_bbox(self) -> None:
        blocks = [
            TextBlock(
                id="l1",
                source="ocr",
                bbox=(72, 130, 400, 160),
                text="• First line",
                confidence=0.9,
                font_size=16,
                image_bbox=(72, 130, 400, 160),
                image_polygon=((72, 130), (400, 130), (400, 160), (72, 160)),
            ),
            TextBlock(
                id="l2",
                source="ocr",
                bbox=(96, 170, 400, 200),
                text="◦ Second line",
                confidence=0.9,
                font_size=16,
                image_bbox=(96, 170, 400, 200),
                image_polygon=((96, 170), (400, 170), (400, 200), (96, 200)),
            ),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        list_blocks = [block for block in applied.text_blocks if block.block_role == "list"]
        self.assertTrue(list_blocks)
        image_bbox = list_blocks[0].image_bbox
        self.assertIsNotNone(image_bbox)
        assert image_bbox is not None
        self.assertAlmostEqual(image_bbox[1], 130, places=0)
        self.assertAlmostEqual(image_bbox[3], 200, places=0)
        self.assertIsNone(list_blocks[0].image_polygon)


class FigureSkipNativeTests(unittest.TestCase):
    def test_skips_figure_crop_over_native_shapes(self) -> None:
        from pdf2ppt.models import ShapeElement

        image = Image.new("RGB", (960, 540), color=(255, 255, 255))
        regions = [
            LayoutRegion(
                id="fig1",
                label="figure",
                bbox=(100, 100, 500, 400),
                confidence=0.95,
                reading_order=1,
                engine="ppstructure",
            )
        ]
        shapes = [
            ShapeElement(
                id="s1",
                kind="rect",
                bbox=(120, 120, 300, 280),
                fill_color="#336699",
            )
        ]
        applied = apply_layout_to_page(
            [],
            LayoutPageResult(regions=regions, engine="ppstructure"),
            page_width=960,
            page_height=540,
            page_image=image,
            shape_elements=shapes,
        )
        self.assertEqual(applied.figure_images, [])


class TableRegionNoMergeTests(unittest.TestCase):
    def test_table_region_without_native_keeps_per_line_boxes(self) -> None:
        blocks = [
            _block("a", "Cell A", (100, 100, 200, 130)),
            _block("b", "Cell B", (220, 100, 320, 130)),
            _block("c", "Cell C", (100, 150, 200, 180)),
        ]
        regions = [
            LayoutRegion(
                id="table",
                label="table",
                bbox=(90, 90, 340, 200),
                confidence=0.9,
                reading_order=1,
                engine="ppstructure",
            )
        ]
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="ppstructure"),
            page_width=960,
            page_height=540,
            table_elements=[],
        )
        self.assertEqual(len(applied.text_blocks), 3)


class ParagraphPropertiesSchemaOrderTests(unittest.TestCase):
    def test_spc_bef_precedes_bullet_children(self) -> None:
        from pptx.oxml.ns import qn

        from pdf2ppt.ppt_render import add_text_block, ppr_children_in_schema_order

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        block = TextBlock(
            id="list1",
            source="native",
            bbox=(72, 120, 600, 280),
            text="One\nTwo",
            confidence=1.0,
            font_size=16,
            block_role="list",
            paragraphs=[
                TextParagraph(text="One", list_level=0, space_before_pt=0),
                TextParagraph(text="Two", list_level=1, space_before_pt=6),
            ],
        )
        add_text_block(slide, block, scale_x=1.0, scale_y=1.0)
        checked = 0
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                p_pr = paragraph._p.pPr  # noqa: SLF001
                if p_pr is None:
                    continue
                self.assertTrue(ppr_children_in_schema_order(p_pr))
                tags = [child.tag for child in p_pr]
                if qn("a:spcBef") in tags and qn("a:buChar") in tags:
                    self.assertLess(tags.index(qn("a:spcBef")), tags.index(qn("a:buChar")))
                    checked += 1
                if qn("a:spcBef") in tags and qn("a:buAutoNum") in tags:
                    self.assertLess(tags.index(qn("a:spcBef")), tags.index(qn("a:buAutoNum")))
                    checked += 1
        self.assertGreaterEqual(checked, 1)

    def test_fixture_outputs_have_valid_ppr_child_order(self) -> None:
        from pptx.oxml.ns import qn

        from pdf2ppt.core import ConversionOptions
        from pdf2ppt.pipeline import convert_pdf
        from pdf2ppt.ppt_render import ppr_children_in_schema_order

        for name in ("bullet_list", "flowchart"):
            pdf = FIXTURES / f"{name}.pdf"
            if not pdf.exists():
                self.skipTest(f"missing {name}.pdf")
            with tempfile.TemporaryDirectory() as tmp:
                out = Path(tmp) / f"{name}.pptx"
                convert_pdf(
                    ConversionOptions(
                        input_path=pdf,
                        output_path=out,
                        report_path=Path(tmp) / f"{name}.json",
                        layout_engine="heuristic",
                    )
                )
                presentation = Presentation(str(out))
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        if not getattr(shape, "has_text_frame", False):
                            continue
                        for paragraph in shape.text_frame.paragraphs:
                            p_pr = getattr(paragraph._p, "pPr", None)  # noqa: SLF001
                            if p_pr is None:
                                continue
                            self.assertTrue(
                                ppr_children_in_schema_order(p_pr),
                                msg=f"{name}: pPr children out of schema order: "
                                f"{[child.tag for child in p_pr]}",
                            )
                            tags = [child.tag for child in p_pr]
                            for bullet_tag in (qn("a:buChar"), qn("a:buAutoNum")):
                                if qn("a:spcBef") in tags and bullet_tag in tags:
                                    self.assertLess(tags.index(qn("a:spcBef")), tags.index(bullet_tag))


class BulletListPositionTests(unittest.TestCase):
    """Heuristic bullet boxes must track PDF glyph positions (|dx|<5, |dy|<3)."""

    @unittest.skipUnless((FIXTURES / "bullet_list.pdf").exists(), "fixture missing")
    def test_bullet_list_line_positions_near_pdf(self) -> None:
        import fitz

        from pdf2ppt.core import ConversionOptions
        from pdf2ppt.layout import _LINE_HEIGHT_FACTOR
        from pdf2ppt.pipeline import convert_pdf
        from pdf2ppt.ppt_render import EMU_PER_PT

        pdf_path = FIXTURES / "bullet_list.pdf"
        # PDF ground truth: body lines only (skip title / footer).
        pdf_lines: list[tuple[str, float, float, float]] = []
        with fitz.open(pdf_path) as document:
            page = document[0]
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    text = "".join(span["text"] for span in line.get("spans", [])).strip()
                    if not text.startswith("-"):
                        continue
                    bbox = line["bbox"]
                    size = float(line["spans"][0]["size"]) if line.get("spans") else 16.0
                    pdf_lines.append((text, float(bbox[0]), float(bbox[1]), size))

        self.assertGreaterEqual(len(pdf_lines), 6)

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "bullet_list.pptx"
            report_path = Path(tmp) / "bullet_list.json"
            convert_pdf(
                ConversionOptions(
                    input_path=pdf_path,
                    output_path=out,
                    report_path=report_path,
                    layout_engine="heuristic",
                )
            )
            presentation = Presentation(str(out))
            slide = presentation.slides[0]

        list_shape = None
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            joined = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if "Ship layout analysis" in joined:
                list_shape = shape
                break
        self.assertIsNotNone(list_shape)
        assert list_shape is not None

        box_left_pt = float(list_shape.left) / EMU_PER_PT
        box_top_pt = float(list_shape.top) / EMU_PER_PT
        cursor_y = box_top_pt
        prev_font_pt = 16.0
        bullet_paras = []
        for paragraph in list_shape.text_frame.paragraphs:
            p_pr = getattr(paragraph._p, "pPr", None)  # noqa: SLF001
            if p_pr is None:
                continue
            has_bullet = any(
                child.tag.endswith("}buChar") or child.tag.endswith("}buAutoNum") for child in p_pr
            )
            if not has_bullet:
                continue
            mar_l = int(p_pr.get("marL") or 0) / EMU_PER_PT
            indent = abs(int(p_pr.get("indent") or 0)) / EMU_PER_PT
            spc_bef = 0.0
            for child in p_pr:
                if child.tag.endswith("}spcBef"):
                    pts = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcPts")
                    if pts is not None and pts.get("val"):
                        spc_bef = int(pts.get("val")) / 100.0
            font_pt = prev_font_pt
            for run in paragraph.runs:
                if run.font.size is not None:
                    font_pt = float(run.font.size.pt)
                    break
            if bullet_paras:
                cursor_y += _LINE_HEIGHT_FACTOR * prev_font_pt + spc_bef
            text_x = box_left_pt + mar_l
            bullet_paras.append((paragraph.text, text_x, cursor_y, indent, font_pt))
            prev_font_pt = font_pt

        self.assertEqual(len(bullet_paras), len(pdf_lines))
        for (_pdf_text, pdf_bullet_x, pdf_y0, pdf_size), (_ppt_text, text_x, text_y, indent, _font) in zip(
            pdf_lines, bullet_paras, strict=True
        ):
            # PDF body text starts ~prefix-width after the bullet glyph.
            expected_text_x = pdf_bullet_x + indent
            self.assertLessEqual(
                abs(text_x - expected_text_x),
                5.0,
                msg=f"dx too large for {_ppt_text!r}: ppt={text_x:.1f} pdf≈{expected_text_x:.1f}",
            )
            self.assertLessEqual(
                abs(text_y - pdf_y0),
                3.0,
                msg=f"dy too large for {_ppt_text!r}: ppt={text_y:.1f} pdf={pdf_y0:.1f}",
            )
            # Hanging indent should be the measured prefix (~10pt), not the 27pt fallback.
            self.assertLessEqual(indent, 18.0)
            self.assertGreaterEqual(indent, 4.0)

    def test_layout_paragraphs_use_baseline_spacing_and_prefix_indent(self) -> None:
        blocks = [
            _block("a", "- Ship layout analysis for scanned decks", (72, 130, 360, 152), font_size=16),
            _block("b", "- Detect multi-level bullets", (96, 170, 280, 192), font_size=16),
            _block("c", "- Map titles to placeholders", (96, 205, 290, 227), font_size=16),
        ]
        regions = analyze_layout_heuristic(blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        list_blocks = [block for block in applied.text_blocks if block.block_role == "list"]
        self.assertTrue(list_blocks)
        paragraphs = list_blocks[0].paragraphs or []
        self.assertGreaterEqual(len(paragraphs), 3)
        # Prefix "- " at 16pt measures ~11pt, far below the 27pt fallback.
        self.assertIsNotNone(paragraphs[0].text_indent_pt)
        assert paragraphs[0].text_indent_pt is not None
        self.assertLess(paragraphs[0].text_indent_pt, 20.0)
        self.assertGreater(paragraphs[0].text_indent_pt, 4.0)
        # Nested line keeps its PDF left edge for marL.
        self.assertAlmostEqual(paragraphs[1].left_pt or 0.0, 96.0, places=0)
        # Baseline spacing ≈ (y0 delta) - 1.2*font, not raw bbox gap (~18).
        self.assertIsNotNone(paragraphs[1].space_before_pt)
        assert paragraphs[1].space_before_pt is not None
        self.assertGreater(paragraphs[1].space_before_pt, 18.5)


class MultilineBodySpacingTests(unittest.TestCase):
    """Ordinary multi-line paragraphs must keep PDF baseline spacing (large_mask)."""

    @unittest.skipUnless((FIXTURES / "large_mask.pdf").exists(), "fixture missing")
    def test_large_mask_layout_baselines_within_3pt_of_pdf(self) -> None:
        import fitz

        from pdf2ppt.layout import _LINE_HEIGHT_FACTOR
        from pdf2ppt.native_extraction import extract_native_text_blocks

        pdf_path = FIXTURES / "large_mask.pdf"
        pdf_line_tops: list[float] = []
        with fitz.open(pdf_path) as document:
            page = document[0]
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    text = "".join(span["text"] for span in line.get("spans", [])).strip()
                    if not text:
                        continue
                    pdf_line_tops.append(float(line["bbox"][1]))
            native_blocks, _images = extract_native_text_blocks(page)

        self.assertGreaterEqual(len(pdf_line_tops), 4)
        regions = analyze_layout_heuristic(native_blocks, page_width=960, page_height=540)
        applied = apply_layout_to_page(
            native_blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        body = next(
            (block for block in applied.text_blocks if (block.paragraphs or []) and len(block.paragraphs or []) >= 4),
            None,
        )
        self.assertIsNotNone(body)
        assert body is not None
        paragraphs = body.paragraphs or []
        self.assertEqual(len(paragraphs), len(pdf_line_tops))

        # Reconstruct rendered line tops the same way PPT single-spacing + spcBef does.
        cursor_y = float(body.bbox[1])
        prev_font = paragraphs[0].font_size or 17.0
        for index, paragraph in enumerate(paragraphs):
            font = paragraph.font_size or prev_font
            if index > 0:
                spc = paragraph.space_before_pt or 0.0
                cursor_y += _LINE_HEIGHT_FACTOR * prev_font + spc
            self.assertLessEqual(
                abs(cursor_y - pdf_line_tops[index]),
                3.0,
                msg=(
                    f"line {index}: ppt≈{cursor_y:.2f} pdf={pdf_line_tops[index]:.2f} "
                    f"spcBef={paragraph.space_before_pt}"
                ),
            )
            # Multi-line body must carry positive baseline extras after the first line.
            if index > 0:
                self.assertIsNotNone(paragraph.space_before_pt)
                assert paragraph.space_before_pt is not None
                self.assertGreater(paragraph.space_before_pt, 0.5)
            prev_font = font

    def test_multiline_member_block_gets_sliced_baseline_spacing(self) -> None:
        # Lines longer than 40 chars so split_wide_multiline_blocks keeps one box.
        text = (
            "First wrapped line of the paragraph continues with more words here\n"
            "Second wrapped line continues with still more wording on this row"
        )
        blocks = [_block("body", text, (60, 130, 860, 180), font_size=17)]
        regions = [
            LayoutRegion(
                id="r1",
                label="text",
                bbox=(60, 130, 860, 180),
                confidence=1.0,
                reading_order=1,
            )
        ]
        applied = apply_layout_to_page(
            blocks,
            LayoutPageResult(regions=regions, engine="heuristic"),
            page_width=960,
            page_height=540,
        )
        self.assertEqual(len(applied.text_blocks), 1)
        paragraphs = applied.text_blocks[0].paragraphs or []
        self.assertEqual(len(paragraphs), 2)
        self.assertIsNotNone(paragraphs[1].space_before_pt)
        assert paragraphs[1].space_before_pt is not None
        # Equal slice of a 50pt-tall 2-line box → ~25pt slots; gap above 1.2*17≈20.4.
        self.assertGreater(paragraphs[1].space_before_pt, 2.0)
        self.assertNotEqual(paragraphs[0].baseline_pt, paragraphs[1].baseline_pt)


if __name__ == "__main__":
    unittest.main()
