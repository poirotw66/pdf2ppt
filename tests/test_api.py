from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import fitz
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from pdf2ppt.api import app
from pdf2ppt.core import InputValidationError, OcrInitializationError, PageConversionError
from pdf2ppt.ir_io import default_assets_dir_name, write_slide_ir
from pdf2ppt.job_store import JobStore
from pdf2ppt.models import ImagePlacement, PageResult, QualityScore, ShapeElement, TextBlock
from pdf2ppt.paths import get_repo_root
from pdf2ppt.safe_zip import ZipSafetyLimits


class ApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.job_store = JobStore(Path(self.temp_dir.name) / "jobs")
        self.client = TestClient(app)
        self.job_store_patcher = patch("pdf2ppt.api.job_store", self.job_store)
        self.job_store_patcher.start()
        self.addCleanup(self.job_store_patcher.stop)

    def test_create_job_and_fetch_metadata(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()

        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(create_response.status_code, 200)
        payload = create_response.json()
        self.assertEqual(payload["status"], "uploaded")
        self.assertEqual(payload["page_count"], 1)

        get_response = self.client.get(f"/jobs/{payload['job_id']}")
        self.assertEqual(get_response.status_code, 200)
        fetched = get_response.json()
        self.assertEqual(fetched["job_id"], payload["job_id"])
        self.assertEqual(fetched["original_filename"], "sample.pdf")

    def test_delete_job_removes_job_artifacts(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()

        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        delete_response = self.client.delete(f"/jobs/{job_id}")

        self.assertEqual(delete_response.status_code, 204)
        self.assertFalse(self.job_store.job_dir(job_id).exists())
        get_response = self.client.get(f"/jobs/{job_id}")
        self.assertEqual(get_response.status_code, 404)
        self.assertEqual(get_response.json()["detail"]["code"], "not-found")

    @patch("pdf2ppt.api._clean_preview_watermark", side_effect=lambda image, _page_rect, **_: image)
    @patch("pdf2ppt.api.OcrEngine")
    def test_detect_generates_preview_payload(
        self,
        ocr_engine_class_mock: unittest.mock.Mock,
        clean_preview_watermark_mock: unittest.mock.Mock,
    ) -> None:
        ocr_engine = unittest.mock.Mock()
        ocr_engine.extract_text_blocks_batch.return_value = [
            SimpleNamespace(
                blocks=[
                    SimpleNamespace(
                        id="ocr_1_1",
                        bbox=(10.0, 20.0, 50.0, 60.0),
                        image_polygon=((10.0, 20.0), (50.0, 20.0), (50.0, 60.0), (10.0, 60.0)),
                        text="demo",
                        confidence=0.95,
                    )
                ]
            )
        ]
        ocr_engine_class_mock.return_value = ocr_engine
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        detect_response = self.client.post(f"/jobs/{job_id}/detect", json={"dpi": 110})
        self.assertEqual(detect_response.status_code, 200)
        ocr_engine.extract_text_blocks_batch.assert_called_once()
        self.assertEqual(ocr_engine.extract_text_blocks_batch.call_args.args[1], [1])
        payload = detect_response.json()
        self.assertEqual(payload["status"], "detected")
        self.assertEqual(len(payload["pages"]), 1)
        first_page = payload["pages"][0]
        self.assertEqual(first_page["page"], 1)
        self.assertTrue(first_page["image_url"].endswith(".jpg"))
        self.assertEqual(first_page["boxes"][0]["source"], "ocr-auto")
        self.assertEqual(first_page["boxes"][0]["text"], "demo")

        preview_response = self.client.get(first_page["image_url"])
        self.assertEqual(preview_response.status_code, 200)
        self.assertEqual(preview_response.headers["content-type"], "image/jpeg")
        self.assertEqual(clean_preview_watermark_mock.call_count, 2)
        self.assertFalse((self.job_store.job_dir(job_id) / "previews").exists())

    def test_preview_cleanup_helper_uses_inpainting_and_preserves_image_size(self) -> None:
        from pdf2ppt.api import _clean_preview_watermark

        image = Image.new("RGB", (320, 240), (250, 250, 250))

        cleaned = _clean_preview_watermark(image, fitz.Rect(0, 0, 320, 240))

        self.assertEqual(cleaned.size, image.size)

    def test_preview_cleanup_helper_skips_inpainting_for_raster_uploads(self) -> None:
        from pdf2ppt.api import _clean_preview_watermark

        image = Image.new("RGB", (320, 240), (250, 250, 250))
        cleaned = _clean_preview_watermark(
            image,
            fitz.Rect(0, 0, 320, 240),
            apply_notebooklm_fallback=False,
        )

        self.assertIs(cleaned, image)

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_disables_notebooklm_fallback_for_raster_uploads(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        image = Image.new("RGB", (320, 240), color=(240, 240, 240))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        create_response = self.client.post(
            "/jobs",
            files={"file": ("slide.png", buffer.getvalue(), "image/png")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={
                "pages": [
                    {
                        "page": 1,
                        "width": 320,
                        "height": 240,
                        "boxes": [
                            {
                                "id": "box_1",
                                "source": "ocr-auto",
                                "bbox": [10.0, 12.0, 80.0, 48.0],
                                "text": "alpha",
                                "confidence": 0.91,
                            }
                        ],
                    }
                ]
            },
        )

        response = self.client.post(f"/jobs/{job_id}/convert", json={})
        self.assertEqual(response.status_code, 200)
        options = convert_pdf_mock.call_args.args[0]
        self.assertFalse(options.apply_notebooklm_watermark_fallback)

    def test_preview_endpoint_returns_not_found_for_missing_page(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        preview_response = self.client.get(f"/jobs/{job_id}/pages/2.jpg")

        self.assertEqual(preview_response.status_code, 404)
        self.assertEqual(preview_response.json()["detail"]["code"], "not-found")

    def test_put_boxes_persists_approved_boxes(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        payload = {
            "pages": [
                {
                    "page": 1,
                    "width": 320,
                    "height": 240,
                    "boxes": [
                        {
                            "id": "box_1",
                            "source": "ocr-user",
                            "bbox": [10.0, 20.0, 80.0, 60.0],
                            "polygon": [[10.0, 20.0], [80.0, 20.0], [80.0, 60.0], [10.0, 60.0]],
                            "text": "approved",
                            "confidence": 1.0,
                        }
                    ],
                }
            ]
        }

        response = self.client.put(f"/jobs/{job_id}/boxes", json=payload)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["status"], "boxes-approved")
        self.assertEqual(body["pages"][0]["boxes"][0]["text"], "approved")

        job_response = self.client.get(f"/jobs/{job_id}")
        self.assertEqual(job_response.status_code, 200)
        job_payload = job_response.json()
        self.assertIsNotNone(job_payload["approved_boxes_path"])
        approved_boxes_path = Path(job_payload["approved_boxes_path"])
        self.assertTrue(approved_boxes_path.exists())

    def test_create_job_returns_input_error_for_invalid_pdf(self) -> None:
        response = self.client.post(
            "/jobs",
            files={"file": ("broken.pdf", b"not-a-pdf", "application/pdf")},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"]["code"], "input-error")

    def test_create_job_rejects_unsupported_upload_with_structured_error(self) -> None:
        response = self.client.post(
            "/jobs",
            files={"file": ("broken.txt", b"plain-text", "text/plain")},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"]["code"], "input-error")
        self.assertEqual(response.json()["detail"]["message"], "Only PDF, PNG, and JPG uploads are supported.")

    def test_create_job_accepts_single_page_png_upload(self) -> None:
        image = Image.new("RGB", (320, 240), color=(240, 240, 240))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        response = self.client.post(
            "/jobs",
            files={"file": ("slide.png", buffer.getvalue(), "image/png")},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(payload["original_filename"], "slide.png")

    def test_create_job_accepts_single_page_jpg_upload(self) -> None:
        image = Image.new("RGB", (400, 300), color=(200, 210, 220))
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG")
        response = self.client.post(
            "/jobs",
            files={"file": ("slide.jpg", buffer.getvalue(), "image/jpeg")},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(payload["original_filename"], "slide.jpg")

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_uses_approved_boxes_and_updates_job(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        boxes_payload = {
            "pages": [
                {
                    "page": 1,
                    "width": 320,
                    "height": 240,
                    "boxes": [
                        {
                            "id": "box_1",
                            "source": "ocr-user",
                            "bbox": [10.0, 20.0, 80.0, 60.0],
                            "text": "",
                            "confidence": 1.0,
                        }
                    ],
                }
            ]
        }
        self.client.put(f"/jobs/{job_id}/boxes", json=boxes_payload)

        response = self.client.post(f"/jobs/{job_id}/convert", json={"write_debug_artifacts": False})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["status"], "converted")
        self.assertEqual(body["page_count"], 1)

        options = convert_pdf_mock.call_args.args[0]
        self.assertIn(1, options.approved_ocr_blocks_by_page)
        self.assertEqual(options.approved_ocr_image_size_by_page[1], (320, 240))
        self.assertEqual(options.ocr_batch_size, 3)

        job_response = self.client.get(f"/jobs/{job_id}")
        job_payload = job_response.json()
        self.assertEqual(job_payload["status"], "converted")
        self.assertIsNotNone(job_payload["output_pptx_path"])
        self.assertIsNotNone(job_payload["report_path"])

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_accepts_custom_ocr_batch_size(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={
                "pages": [
                    {
                        "page": 1,
                        "width": 320,
                        "height": 240,
                        "boxes": [],
                    }
                ]
            },
        )

        response = self.client.post(
            f"/jobs/{job_id}/convert",
            json={"write_debug_artifacts": False, "ocr_batch_size": 6},
        )

        self.assertEqual(response.status_code, 200)
        options = convert_pdf_mock.call_args.args[0]
        self.assertEqual(options.ocr_batch_size, 6)

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_passes_lama_inpaint_options(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={
                "pages": [
                    {
                        "page": 1,
                        "width": 320,
                        "height": 240,
                        "boxes": [],
                    }
                ]
            },
        )

        response = self.client.post(
            f"/jobs/{job_id}/convert",
            json={
                "write_debug_artifacts": False,
                "inpaint_engine": "lama-onnx-cuda",
                "inpaint_model_root": "custom-lama",
                "inpaint_onnx_cuda_provider": "CUDAExecutionProvider",
                "inpaint_onnx_execution_mode": "parallel",
                "inpaint_max_side_px": 1024,
            },
        )

        self.assertEqual(response.status_code, 200)
        options = convert_pdf_mock.call_args.args[0]
        # "lama-onnx-cuda" is a backward-compatible alias that normalizes to the canonical
        # "lama-onnx" engine identifier (Phase 1.1); the request itself keeps working unchanged.
        self.assertEqual(options.inpaint_engine, "lama-onnx")
        self.assertEqual(options.inpaint_model_root, get_repo_root() / "custom-lama")
        self.assertEqual(options.inpaint_onnx_cuda_provider, "CUDAExecutionProvider")
        self.assertEqual(options.inpaint_onnx_execution_mode, "parallel")
        self.assertEqual(options.inpaint_max_side_px, 1024)

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_passes_lama_pytorch_inpaint_options(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={
                "pages": [
                    {
                        "page": 1,
                        "width": 320,
                        "height": 240,
                        "boxes": [],
                    }
                ]
            },
        )

        response = self.client.post(
            f"/jobs/{job_id}/convert",
            json={
                "write_debug_artifacts": False,
                "inpaint_engine": "lama-pytorch",
                "inpaint_model_root": "lama/big-lama",
                "inpaint_lama_repo_root": "lama",
                "inpaint_lama_device": "cuda",
            },
        )

        self.assertEqual(response.status_code, 200)
        options = convert_pdf_mock.call_args.args[0]
        self.assertEqual(options.inpaint_engine, "lama-pytorch")
        self.assertEqual(options.inpaint_model_root, get_repo_root() / "lama" / "big-lama")
        self.assertEqual(options.inpaint_lama_repo_root, get_repo_root() / "lama")
        self.assertEqual(options.inpaint_lama_device, "cuda")

    @patch("pdf2ppt.api.OcrEngine.extract_text_blocks_batch")
    def test_detect_returns_ocr_initialization_error(
        self,
        extract_text_blocks_batch_mock: unittest.mock.Mock,
    ) -> None:
        extract_text_blocks_batch_mock.side_effect = OcrInitializationError("missing OCR runtime")
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        response = self.client.post(f"/jobs/{job_id}/detect", json={"dpi": 110})

        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["detail"]["code"], "ocr-initialization-error")

    @patch("pdf2ppt.api.OcrEngine.extract_text_blocks_batch")
    def test_detect_filters_boxes_below_confidence_threshold(
        self,
        extract_text_blocks_batch_mock: unittest.mock.Mock,
    ) -> None:
        extract_text_blocks_batch_mock.return_value = [
            SimpleNamespace(
                blocks=[
                    SimpleNamespace(
                        id="ocr_1_keep",
                        bbox=(10.0, 20.0, 50.0, 60.0),
                        image_polygon=((10.0, 20.0), (50.0, 20.0), (50.0, 60.0), (10.0, 60.0)),
                        text="keep",
                        confidence=0.92,
                    ),
                    SimpleNamespace(
                        id="ocr_1_drop",
                        bbox=(60.0, 70.0, 110.0, 120.0),
                        image_polygon=((60.0, 70.0), (110.0, 70.0), (110.0, 120.0), (60.0, 120.0)),
                        text="drop",
                        confidence=0.62,
                    ),
                ]
            )
        ]
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        response = self.client.post(f"/jobs/{job_id}/detect", json={"dpi": 110})

        self.assertEqual(response.status_code, 200)
        boxes = response.json()["pages"][0]["boxes"]
        self.assertEqual(len(boxes), 1)
        self.assertEqual(boxes[0]["id"], "ocr_1_keep")

    @patch("pdf2ppt.api.OcrEngine.extract_text_blocks_batch")
    def test_detect_accepts_custom_confidence_threshold(
        self,
        extract_text_blocks_batch_mock: unittest.mock.Mock,
    ) -> None:
        extract_text_blocks_batch_mock.return_value = [
            SimpleNamespace(
                blocks=[
                    SimpleNamespace(
                        id="ocr_1_keep",
                        bbox=(10.0, 20.0, 50.0, 60.0),
                        image_polygon=((10.0, 20.0), (50.0, 20.0), (50.0, 60.0), (10.0, 60.0)),
                        text="keep",
                        confidence=0.52,
                    )
                ]
            )
        ]
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        response = self.client.post(
            f"/jobs/{job_id}/detect",
            json={"dpi": 110, "confidence_threshold": 0.5},
        )

        self.assertEqual(response.status_code, 200)
        boxes = response.json()["pages"][0]["boxes"]
        self.assertEqual(len(boxes), 1)
        self.assertEqual(boxes[0]["id"], "ocr_1_keep")

    def test_download_output_returns_structured_not_found_error(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        response = self.client.get(f"/jobs/{job_id}/output.pptx")

        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["detail"]["code"], "not-found")
        self.assertEqual(response.json()["detail"]["message"], "Converted PPTX not found.")

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_returns_page_conversion_error_detail(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.side_effect = PageConversionError(2, "failed to process page 2")
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={
                "pages": [
                    {
                        "page": 1,
                        "width": 320,
                        "height": 240,
                        "boxes": [],
                    }
                ]
            },
        )

        response = self.client.post(f"/jobs/{job_id}/convert", json={"write_debug_artifacts": False})

        self.assertEqual(response.status_code, 500)
        detail = response.json()["detail"]
        self.assertEqual(detail["code"], "page-conversion-error")
        self.assertEqual(detail["page"], 2)

    def test_download_endpoints_serve_existing_outputs(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        output_path = job_dir / "output.pptx"
        report_path = job_dir / "output.report.json"
        output_path.write_bytes(b"pptx-bytes")
        report_path.write_text('{"ok": true}', encoding="utf-8")
        self.job_store.update_job(
            job_id,
            status="converted",
            output_pptx_path=str(output_path),
            report_path=str(report_path),
        )

        output_response = self.client.get(f"/jobs/{job_id}/output.pptx")
        self.assertEqual(output_response.status_code, 200)
        self.assertEqual(output_response.content, b"pptx-bytes")
        self.assertIn('filename="sample.pptx"', output_response.headers["content-disposition"])

        report_response = self.client.get(f"/jobs/{job_id}/report.json")
        self.assertEqual(report_response.status_code, 200)
        self.assertEqual(report_response.json(), {"ok": True})

    @patch("pdf2ppt.api.convert_pdf")
    def test_emit_ir_endpoint_writes_ir_only(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object(), object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]

        response = self.client.post(f"/jobs/{job_id}/emit-ir", json={"layout_engine": "heuristic"})
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "ir-emitted")
        self.assertEqual(payload["page_count"], 2)
        self.assertTrue(payload["slideir_path"].endswith("output.slideir.json"))
        kwargs = convert_pdf_mock.call_args.kwargs
        self.assertFalse(kwargs["write_pptx"])
        self.assertIsNotNone(kwargs["emit_ir_path"])

    @patch("pdf2ppt.api.convert_pdf")
    def test_convert_emit_ir_flag_is_opt_in(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        self.client.put(
            f"/jobs/{job_id}/boxes",
            json={"pages": [{"page": 1, "width": 320, "height": 240, "boxes": []}]},
        )

        default_response = self.client.post(f"/jobs/{job_id}/convert", json={})
        self.assertEqual(default_response.status_code, 200)
        self.assertIsNone(default_response.json().get("slideir_path"))
        self.assertTrue(convert_pdf_mock.call_args.kwargs.get("write_pptx", True))
        self.assertIsNone(convert_pdf_mock.call_args.kwargs.get("emit_ir_path"))

        convert_pdf_mock.reset_mock()
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        emit_response = self.client.post(
            f"/jobs/{job_id}/convert",
            json={"emit_ir": True, "ir_only": True},
        )
        self.assertEqual(emit_response.status_code, 200)
        self.assertIsNotNone(emit_response.json().get("slideir_path"))
        self.assertFalse(convert_pdf_mock.call_args.kwargs["write_pptx"])

    def test_from_ir_with_assets_zip_returns_pptx_and_cleans_work_dir(self) -> None:
        ir_bytes, zip_bytes = build_slideir_upload_pair(self.temp_dir.name)
        before = {path.name for path in self.job_store.root.glob("pdf2ppt-from-ir-*")}

        response = self.client.post(
            "/from-ir",
            files={
                "ir_file": ("deck.slideir.json", ir_bytes, "application/json"),
                "assets_zip": ("assets.zip", zip_bytes, "application/zip"),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "presentationml.presentation",
            response.headers.get("content-type", ""),
        )
        pptx_bytes = response.content
        self.assertGreater(len(pptx_bytes), 1000)
        presentation = Presentation(io.BytesIO(pptx_bytes))
        self.assertEqual(len(presentation.slides), 1)

        after = {path.name for path in self.job_store.root.glob("pdf2ppt-from-ir-*")}
        self.assertEqual(after, before)

    def test_from_ir_without_assets_zip_returns_400(self) -> None:
        ir_bytes, _zip_bytes = build_slideir_upload_pair(self.temp_dir.name)
        response = self.client.post(
            "/from-ir",
            files={"ir_file": ("deck.slideir.json", ir_bytes, "application/json")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"]["code"], "slideir-error")
        leftover = list(self.job_store.root.glob("pdf2ppt-from-ir-*"))
        self.assertEqual(leftover, [])

    def test_from_ir_zip_path_traversal_rejected(self) -> None:
        ir_bytes, _zip_bytes = build_slideir_upload_pair(self.temp_dir.name)
        evil_outside = Path(self.temp_dir.name) / "evil-outside.txt"
        self.assertFalse(evil_outside.exists())

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w") as archive:
            archive.writestr("../evil-outside.txt", b"pwned")
            archive.writestr("../../evil-outside.txt", b"pwned")
        zip_bytes = zip_buffer.getvalue()

        response = self.client.post(
            "/from-ir",
            files={
                "ir_file": ("deck.slideir.json", ir_bytes, "application/json"),
                "assets_zip": ("assets.zip", zip_bytes, "application/zip"),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"]["code"], "input-error")
        self.assertFalse(evil_outside.exists())
        # Also ensure nothing escaped next to the job store root.
        sibling_evil = self.job_store.root.parent / "evil-outside.txt"
        self.assertFalse(sibling_evil.exists())
        leftover = list(self.job_store.root.glob("pdf2ppt-from-ir-*"))
        self.assertEqual(leftover, [])

    def test_from_ir_rejects_oversized_before_full_read(self) -> None:
        """UploadFile.size must be checked before reading the body (Z1)."""
        huge = SimpleNamespace(
            filename="assets.zip",
            size=400 * 1024 * 1024,
            read=MagicMock(side_effect=AssertionError("must not read oversized body")),
        )
        # Drive the helper directly — TestClient always buffers multipart bodies.
        import asyncio

        from pdf2ppt.api import _read_upload_capped
        from pdf2ppt.safe_zip import reject_oversized_compressed

        limits = ZipSafetyLimits(max_compressed_bytes=50 * 1024 * 1024)

        async def _run() -> None:
            await _read_upload_capped(
                huge,  # type: ignore[arg-type]
                max_bytes=limits.max_compressed_bytes,
                label="assets zip",
                size_checker=lambda size: reject_oversized_compressed(
                    size, limits=limits, label="assets zip"
                ),
            )

        with self.assertRaises(InputValidationError) as ctx:
            asyncio.run(_run())
        self.assertIn("compressed", str(ctx.exception).lower())
        huge.read.assert_not_called()

    def test_from_ir_zip_bomb_rejected(self) -> None:
        ir_bytes, _zip_bytes = build_slideir_upload_pair(self.temp_dir.name)
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("bomb.bin", b"\x00" * 400_000)
        zip_bytes = zip_buffer.getvalue()
        # Force a tight ratio so the highly compressible payload trips the bomb check.
        with patch(
            "pdf2ppt.api.ZipSafetyLimits.from_environ",
            return_value=ZipSafetyLimits(
                max_compressed_bytes=50 * 1024 * 1024,
                max_uncompressed_bytes=50 * 1024 * 1024,
                max_entries=10_000,
                max_entry_bytes=50 * 1024 * 1024,
                max_compression_ratio=5.0,
            ),
        ):
            response = self.client.post(
                "/from-ir",
                files={
                    "ir_file": ("deck.slideir.json", ir_bytes, "application/json"),
                    "assets_zip": ("assets.zip", zip_bytes, "application/zip"),
                },
            )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"]["code"], "input-error")
        self.assertIn("compression", response.json()["detail"]["message"].lower())
        leftover = list(self.job_store.root.glob("pdf2ppt-from-ir-*"))
        self.assertEqual(leftover, [])

    def test_slideir_download_endpoints(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        ir_path = job_dir / "output.slideir.json"
        page = _api_slideir_page()
        write_slide_ir([page], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
        self.job_store.update_job(job_id, status="ir-emitted", report_path=str(job_dir / "output.slideir.report.json"))

        json_response = self.client.get(f"/jobs/{job_id}/slideir.json")
        self.assertEqual(json_response.status_code, 200)
        self.assertEqual(json_response.json()["schema_version"], "1.0")

        zip_response = self.client.get(f"/jobs/{job_id}/slideir-assets.zip")
        self.assertEqual(zip_response.status_code, 200)
        self.assertEqual(zip_response.headers.get("content-type"), "application/zip")
        with zipfile.ZipFile(io.BytesIO(zip_response.content)) as archive:
            names = archive.namelist()
        self.assertTrue(any(name.startswith("pages/") for name in names))

    @patch("pdf2ppt.api.convert_pdf")
    def test_preview_endpoint_builds_ir_and_returns_overlays(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        convert_pdf_mock.return_value = SimpleNamespace(pages=[object()])
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        ir_path = job_dir / "output.slideir.json"

        def _write_ir(*_args: object, **kwargs: object) -> SimpleNamespace:
            write_slide_ir([_api_slideir_page()], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
            return SimpleNamespace(pages=[object()])

        convert_pdf_mock.side_effect = _write_ir

        response = self.client.post(f"/jobs/{job_id}/preview", json={})
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["job_id"], job_id)
        self.assertEqual(payload["page_count"], 1)
        page = payload["pages"][0]
        self.assertEqual(page["page"], 1)
        self.assertTrue(page["image_url"].endswith("/pages/1.jpg"))
        self.assertGreaterEqual(len(page["text_blocks"]), 1)
        self.assertGreaterEqual(len(page["shapes"]), 1)
        self.assertGreaterEqual(len(page["images"]), 1)
        # Planner must stay off for preview.
        options = convert_pdf_mock.call_args.args[0]
        self.assertIsNone(getattr(options, "layout_planner", None))

        get_response = self.client.get(f"/jobs/{job_id}/preview")
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual(get_response.json()["page_count"], 1)

    def test_preview_get_without_ir_returns_404(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        response = self.client.get(f"/jobs/{job_id}/preview")
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["detail"]["code"], "not-found")

    def test_preview_edits_validation_and_from_ir_convert(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        ir_path = job_dir / "output.slideir.json"
        write_slide_ir([_api_slideir_page()], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
        self.job_store.update_job(job_id, status="ir-emitted")

        bad_role = self.client.put(
            f"/jobs/{job_id}/preview-edits",
            json={"pages": [{"page": 1, "role_overrides": {"t1": "subtitle"}}]},
        )
        self.assertEqual(bad_role.status_code, 422)

        bad_id = self.client.put(
            f"/jobs/{job_id}/preview-edits",
            json={"pages": [{"page": 1, "excluded_ids": ["missing"], "role_overrides": {}}]},
        )
        self.assertEqual(bad_id.status_code, 400)
        self.assertEqual(bad_id.json()["detail"]["code"], "input-error")

        save_response = self.client.put(
            f"/jobs/{job_id}/preview-edits",
            json={
                "pages": [
                    {
                        "page": 1,
                        "excluded_ids": ["s1"],
                        "role_overrides": {"t1": "title"},
                        "raster_fallback": False,
                    }
                ]
            },
        )
        self.assertEqual(save_response.status_code, 200)
        self.assertEqual(save_response.json()["status"], "preview-edited")
        self.assertEqual(save_response.json()["edits"]["pages"][0]["role_overrides"]["t1"], "title")

        preview = self.client.get(f"/jobs/{job_id}/preview").json()
        text = next(block for block in preview["pages"][0]["text_blocks"] if block["id"] == "t1")
        self.assertEqual(text["block_role"], "body")
        self.assertEqual(preview["edits"]["pages"][0]["role_overrides"]["t1"], "title")
        shape = next(item for item in preview["pages"][0]["shapes"] if item["id"] == "s1")
        self.assertTrue(shape["excluded"])

        convert_response = self.client.post(f"/jobs/{job_id}/from-ir")
        self.assertEqual(convert_response.status_code, 200)
        payload = convert_response.json()
        self.assertEqual(payload["status"], "converted")
        self.assertEqual(payload["page_count"], 1)
        self.assertTrue(Path(payload["output_pptx_path"]).exists())
        self.assertTrue(Path(payload["edited_slideir_path"]).exists())
        edited = Path(payload["edited_slideir_path"]).read_text(encoding="utf-8")
        self.assertIn('"block_role": "title"', edited)
        self.assertNotIn('"id": "s1"', edited)

        pptx_response = self.client.get(f"/jobs/{job_id}/output.pptx")
        self.assertEqual(pptx_response.status_code, 200)
        presentation = Presentation(io.BytesIO(pptx_response.content))
        self.assertEqual(len(presentation.slides), 1)

    @patch("pdf2ppt.api.convert_pdf")
    def test_rebuild_preview_clears_stale_edits(
        self,
        convert_pdf_mock: unittest.mock.Mock,
    ) -> None:
        """Re-analysis must drop preview_edits.json so positional ids cannot 400 Convert."""
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        ir_path = job_dir / "output.slideir.json"

        def _write_ir(*_args: object, **_kwargs: object) -> SimpleNamespace:
            write_slide_ir([_api_slideir_page()], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
            return SimpleNamespace(pages=[object()])

        convert_pdf_mock.side_effect = _write_ir
        first = self.client.post(f"/jobs/{job_id}/preview", json={})
        self.assertEqual(first.status_code, 200)

        save_response = self.client.put(
            f"/jobs/{job_id}/preview-edits",
            json={
                "pages": [
                    {
                        "page": 1,
                        "excluded_ids": ["s1"],
                        "role_overrides": {"t1": "title"},
                        "raster_fallback": False,
                    }
                ]
            },
        )
        self.assertEqual(save_response.status_code, 200)
        edits_path = job_dir / "preview_edits.json"
        self.assertTrue(edits_path.exists())

        rebuild = self.client.post(f"/jobs/{job_id}/preview", json={})
        self.assertEqual(rebuild.status_code, 200)
        self.assertFalse(edits_path.exists())
        self.assertEqual(rebuild.json()["edits"], {"pages": []})

        convert_response = self.client.post(f"/jobs/{job_id}/from-ir")
        self.assertEqual(convert_response.status_code, 200)

    def test_preview_edits_rejects_oversized_body(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        write_slide_ir(
            [_api_slideir_page()],
            job_dir / "output.slideir.json",
            slide_width_pt=320.0,
            slide_height_pt=240.0,
        )
        self.job_store.update_job(job_id, status="ir-emitted")

        huge = {"pages": [{"page": 1, "excluded_ids": ["x"] * 50_000, "role_overrides": {}}]}
        response = self.client.put(f"/jobs/{job_id}/preview-edits", json=huge)
        self.assertIn(response.status_code, {413, 422})

    def test_page_preview_thumbnail_width(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        full = self.client.get(f"/jobs/{job_id}/pages/1.jpg")
        thumb = self.client.get(f"/jobs/{job_id}/pages/1.jpg?width=120")
        self.assertEqual(full.status_code, 200)
        self.assertEqual(thumb.status_code, 200)
        self.assertLess(len(thumb.content), len(full.content))
        thumb_image = Image.open(io.BytesIO(thumb.content))
        self.assertEqual(thumb_image.width, 120)

    def test_from_ir_job_with_raster_fallback(self) -> None:
        pdf_bytes = build_sample_pdf_bytes()
        create_response = self.client.post(
            "/jobs",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
        job_id = create_response.json()["job_id"]
        job_dir = self.job_store.job_dir(job_id)
        ir_path = job_dir / "output.slideir.json"
        write_slide_ir([_api_slideir_page()], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
        self.job_store.update_job(job_id, status="ir-emitted")

        save_response = self.client.put(
            f"/jobs/{job_id}/preview-edits",
            json={"pages": [{"page": 1, "excluded_ids": [], "role_overrides": {}, "raster_fallback": True}]},
        )
        self.assertEqual(save_response.status_code, 200)

        convert_response = self.client.post(f"/jobs/{job_id}/from-ir")
        self.assertEqual(convert_response.status_code, 200)
        edited_path = Path(convert_response.json()["edited_slideir_path"])
        edited = __import__("json").loads(edited_path.read_text(encoding="utf-8"))
        self.assertEqual(edited["pages"][0]["background_mode"], "full-page")
        self.assertEqual(edited["pages"][0]["text_blocks"], [])
        bg = edited["pages"][0]["background"]
        self.assertIsNotNone(bg)
        assets_dir = job_dir / default_assets_dir_name(ir_path)
        self.assertTrue((assets_dir / bg["path"]).exists())


def _tiny_png() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
        b"\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def _api_slideir_page() -> PageResult:
    return PageResult(
        page_number=1,
        page_kind="digital",
        background_mode="elements",
        width_pt=320.0,
        height_pt=240.0,
        text_blocks=[
            TextBlock(
                id="t1",
                source="native",
                bbox=(40.0, 40.0, 200.0, 80.0),
                text="Hello API IR",
                confidence=1.0,
                font_size=18.0,
                reading_order=0,
            )
        ],
        quality_score=QualityScore(1.0, 1.0, 1.0, 1.0),
        fallback_reason=None,
        image_elements=[ImagePlacement(bbox=(10.0, 100.0, 60.0, 150.0), png_bytes=_tiny_png())],
        shape_elements=[
            ShapeElement(id="s1", kind="rect", bbox=(20.0, 20.0, 80.0, 60.0), fill_color="#336699")
        ],
    )


def build_slideir_upload_pair(temp_root: str) -> tuple[bytes, bytes]:
    """Build SlideIR JSON bytes + zip of its assets_dir contents (pages/...)."""
    work = Path(temp_root) / "slideir-build"
    work.mkdir(parents=True, exist_ok=True)
    ir_path = work / "deck.slideir.json"
    write_slide_ir([_api_slideir_page()], ir_path, slide_width_pt=320.0, slide_height_pt=240.0)
    assets_dir = work / default_assets_dir_name(ir_path)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        for path in sorted(assets_dir.rglob("*")):
            if path.is_file():
                archive.write(path, arcname=str(path.relative_to(assets_dir)))
    return ir_path.read_bytes(), zip_buffer.getvalue()


def build_sample_pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page(width=320, height=240)
    page.insert_text((50, 80), "Hello API")
    return document.tobytes()
