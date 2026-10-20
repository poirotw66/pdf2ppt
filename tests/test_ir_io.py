"""P5 SlideIR serialize / deserialize / CLI round-trip tests."""

from __future__ import annotations

import json
import tempfile
import unittest
import unittest.mock
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from pdf2ppt.cli import build_parser, main
from pdf2ppt.core import ConversionOptions
from pdf2ppt.ir_io import (
    SLIDEIR_SCHEMA_VERSION,
    SlideIrError,
    default_assets_dir_name,
    load_slideir_schema,
    read_slide_ir,
    validate_slide_ir_payload,
    write_slide_ir,
)
from pdf2ppt.models import (
    ImagePlacement,
    PageResult,
    QualityScore,
    ShapeElement,
    TextBlock,
)
from pdf2ppt.pipeline import convert_from_ir, convert_pdf

REPO_ROOT = Path(__file__).resolve().parent.parent
FIXTURES_DIR = REPO_ROOT / "tests" / "fixtures"


def _minimal_page(*, with_image: bool = False, with_background: bool = False) -> PageResult:
    images: list[ImagePlacement] = []
    if with_image:
        # 1x1 PNG
        png = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
            b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
            b"\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        images.append(ImagePlacement(bbox=(10.0, 10.0, 50.0, 50.0), png_bytes=png))
    background = None
    if with_background:
        background = b"\xff\xd8\xff\xd9"  # minimal JPEG-ish bytes
    return PageResult(
        page_number=1,
        page_kind="digital",
        background_mode="elements" if not with_background else "overlay",
        width_pt=612.0,
        height_pt=792.0,
        text_blocks=[
            TextBlock(
                id="t1",
                source="native",
                bbox=(72.0, 72.0, 300.0, 100.0),
                text="Hello",
                confidence=1.0,
                font_size=18.0,
                font_color="#000000",
                reading_order=0,
                block_role="title",
            )
        ],
        quality_score=QualityScore(1.0, 1.0, 1.0, 1.0),
        fallback_reason=None,
        background_image_bytes=background,
        image_elements=images,
        shape_elements=[
            ShapeElement(
                id="s1",
                kind="rect",
                bbox=(20.0, 20.0, 100.0, 80.0),
                fill_color="#ff0000",
            )
        ],
    )


class SlideIrSchemaTests(unittest.TestCase):
    def test_bundled_schema_loads(self) -> None:
        schema = load_slideir_schema()
        self.assertEqual(schema["title"], "pdf2ppt SlideIR")
        self.assertEqual(schema["properties"]["schema_version"]["const"], SLIDEIR_SCHEMA_VERSION)

    def test_default_assets_dir_uses_full_filename(self) -> None:
        self.assertEqual(
            default_assets_dir_name(Path("deck.slideir.json")),
            "deck.slideir.json.assets",
        )

    def test_validate_payload_reports_missing_pages(self) -> None:
        issues = validate_slide_ir_payload({"schema_version": "1.0"})
        self.assertTrue(any("pages" in issue for issue in issues))


class SlideIrRoundTripUnitTests(unittest.TestCase):
    def test_write_read_preserves_text_and_shapes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            ir_path = Path(temp_dir) / "deck.slideir.json"
            page = _minimal_page(with_image=True, with_background=True)
            write_slide_ir([page], ir_path, slide_width_pt=612.0, slide_height_pt=792.0)
            self.assertTrue(ir_path.exists())
            assets = ir_path.parent / default_assets_dir_name(ir_path)
            self.assertTrue(assets.is_dir())
            payload = json.loads(ir_path.read_text(encoding="utf-8"))
            self.assertEqual(validate_slide_ir_payload(payload), [])
            deck = read_slide_ir(ir_path)
            self.assertEqual(len(deck.pages), 1)
            restored = deck.pages[0]
            self.assertEqual(restored.text_blocks[0].text, "Hello")
            self.assertEqual(restored.shape_elements[0].kind, "rect")
            self.assertEqual(len(restored.image_elements), 1)
            self.assertIsNotNone(restored.background_image_bytes)
            self.assertTrue(restored.background_image_bytes.startswith(b"\xff\xd8"))

    def test_rejects_path_traversal_in_assets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            ir_path = Path(temp_dir) / "deck.slideir.json"
            write_slide_ir([_minimal_page()], ir_path)
            payload = json.loads(ir_path.read_text(encoding="utf-8"))
            payload["pages"][0]["images"] = [
                {
                    "id": "evil",
                    "bbox": [0, 0, 1, 1],
                    "asset": {"path": "../secret.png"},
                }
            ]
            ir_path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaises(SlideIrError):
                read_slide_ir(ir_path)

    def test_rejects_unknown_schema_version(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            ir_path = Path(temp_dir) / "deck.slideir.json"
            write_slide_ir([_minimal_page()], ir_path)
            payload = json.loads(ir_path.read_text(encoding="utf-8"))
            payload["schema_version"] = "9.9"
            ir_path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaises(SlideIrError):
                read_slide_ir(ir_path)

    def test_from_ir_renders_pptx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            work = Path(temp_dir)
            ir_path = work / "deck.slideir.json"
            pptx_path = work / "out.pptx"
            write_slide_ir([_minimal_page()], ir_path, slide_width_pt=612.0, slide_height_pt=792.0)
            report = convert_from_ir(ir_path, pptx_path)
            self.assertTrue(pptx_path.exists())
            self.assertEqual(len(report.pages), 1)
            presentation = Presentation(str(pptx_path))
            self.assertEqual(len(presentation.slides), 1)


@unittest.skipUnless((FIXTURES_DIR / "solid_background.pdf").exists(), "fixture missing")
class SlideIrFixtureRoundTripTests(unittest.TestCase):
    def test_fixture_emit_ir_from_ir_pptx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            work = Path(temp_dir)
            pdf_path = FIXTURES_DIR / "solid_background.pdf"
            ir_path = work / "solid.slideir.json"
            pptx_path = work / "solid.pptx"
            report_path = work / "solid.report.json"
            options = ConversionOptions(
                input_path=pdf_path,
                output_path=work / "unused.pptx",
                report_path=report_path,
                vector_mode="auto",
                ocr_model_root=None,
            )
            emit_report = convert_pdf(options, write_pptx=False, emit_ir_path=ir_path)
            self.assertGreaterEqual(len(emit_report.pages), 1)
            self.assertTrue(ir_path.exists())
            payload = json.loads(ir_path.read_text(encoding="utf-8"))
            self.assertEqual(validate_slide_ir_payload(payload), [])
            self.assertEqual(payload["schema_version"], SLIDEIR_SCHEMA_VERSION)

            from_report = convert_from_ir(ir_path, pptx_path)
            self.assertTrue(pptx_path.exists())
            self.assertEqual(len(from_report.pages), len(emit_report.pages))
            presentation = Presentation(str(pptx_path))
            self.assertEqual(len(presentation.slides), len(emit_report.pages))

    def test_vector_shapes_round_trip_keeps_shapes(self) -> None:
        fixture = FIXTURES_DIR / "vector_shapes.pdf"
        if not fixture.exists():
            self.skipTest("vector_shapes fixture missing")
        with tempfile.TemporaryDirectory() as temp_dir:
            work = Path(temp_dir)
            ir_path = work / "shapes.slideir.json"
            pptx_path = work / "shapes.pptx"
            options = ConversionOptions(
                input_path=fixture,
                output_path=work / "unused.pptx",
                report_path=work / "report.json",
                vector_mode="auto",
                ocr_model_root=None,
            )
            convert_pdf(options, write_pptx=False, emit_ir_path=ir_path)
            payload = json.loads(ir_path.read_text(encoding="utf-8"))
            shape_count = sum(len(page.get("shapes") or []) for page in payload["pages"])
            self.assertGreater(shape_count, 0)
            convert_from_ir(ir_path, pptx_path)
            self.assertTrue(pptx_path.exists())

    @unittest.skipUnless((FIXTURES_DIR / "figure_caption.pdf").exists(), "fixture missing")
    def test_figure_caption_from_ir_pptx_byte_matches_convert(self) -> None:
        """SlideIR round-trip must produce the same PPTX bytes as direct convert."""

        with tempfile.TemporaryDirectory() as temp_dir:
            work = Path(temp_dir)
            pdf_path = FIXTURES_DIR / "figure_caption.pdf"
            direct = work / "direct.pptx"
            via_ir = work / "via_ir.pptx"
            ir_path = work / "figure.slideir.json"
            options = ConversionOptions(
                input_path=pdf_path,
                output_path=direct,
                report_path=work / "direct.report.json",
                vector_mode="auto",
                ocr_model_root=None,
                layout_engine="heuristic",
            )
            convert_pdf(options)
            convert_pdf(
                ConversionOptions(
                    input_path=pdf_path,
                    output_path=work / "unused.pptx",
                    report_path=work / "ir.report.json",
                    vector_mode="auto",
                    ocr_model_root=None,
                    layout_engine="heuristic",
                ),
                write_pptx=False,
                emit_ir_path=ir_path,
            )
            convert_from_ir(ir_path, via_ir)
            self.assertEqual(direct.read_bytes(), via_ir.read_bytes())


class SlideIrCliTests(unittest.TestCase):
    def test_parser_accepts_emit_and_from_flags(self) -> None:
        parser = build_parser()
        args = parser.parse_args(
            ["input.pdf", "--emit-ir", "deck.json", "--ir-only"]
        )
        self.assertEqual(args.emit_ir, Path("deck.json"))
        self.assertTrue(args.ir_only)
        args2 = parser.parse_args(["--from-ir", "deck.json", "-o", "out.pptx"])
        self.assertEqual(args2.from_ir, Path("deck.json"))
        self.assertEqual(args2.output_override, Path("out.pptx"))

    def test_legacy_convert_args_still_parse(self) -> None:
        parser = build_parser()
        args = parser.parse_args(["input.pdf", "output.pptx"])
        self.assertEqual(args.input_pdf, Path("input.pdf"))
        self.assertEqual(args.output_pptx, Path("output.pptx"))
        self.assertIsNone(args.emit_ir)
        self.assertFalse(args.ir_only)

    @patch("pdf2ppt.cli.convert_pdf")
    def test_main_emit_ir_only_skips_pptx(self, convert_pdf_mock: unittest.mock.Mock) -> None:
        convert_pdf_mock.return_value = type(
            "R",
            (),
            {"input_path": "a.pdf", "output_path": "", "pages": [object()]},
        )()
        code = main(["emit-ir", "a.pdf", "deck.slideir.json"])
        self.assertEqual(code, 0)
        kwargs = convert_pdf_mock.call_args.kwargs
        self.assertFalse(kwargs["write_pptx"])
        self.assertEqual(kwargs["emit_ir_path"], Path("deck.slideir.json"))

    @patch("pdf2ppt.cli.convert_from_ir")
    def test_main_from_ir_subcommand(self, convert_from_ir_mock: unittest.mock.Mock) -> None:
        convert_from_ir_mock.return_value = type(
            "R",
            (),
            {"pages": [object()]},
        )()
        code = main(["from-ir", "deck.slideir.json", "out.pptx"])
        self.assertEqual(code, 0)
        convert_from_ir_mock.assert_called_once()
        call_args = convert_from_ir_mock.call_args
        self.assertEqual(call_args.args[0], Path("deck.slideir.json"))
        self.assertEqual(call_args.args[1], Path("out.pptx"))

    @patch("pdf2ppt.cli.convert_pdf")
    def test_default_convert_still_writes_pptx(self, convert_pdf_mock: unittest.mock.Mock) -> None:
        convert_pdf_mock.return_value = type(
            "R",
            (),
            {"input_path": "a.pdf", "output_path": "b.pptx", "pages": [object()]},
        )()
        code = main(["a.pdf", "b.pptx"])
        self.assertEqual(code, 0)
        kwargs = convert_pdf_mock.call_args.kwargs
        self.assertTrue(kwargs.get("write_pptx", True))
        self.assertIsNone(kwargs.get("emit_ir_path"))


@unittest.skipUnless((FIXTURES_DIR / "solid_background.pdf").exists(), "fixture missing")
class SlideIrCliSmokeTests(unittest.TestCase):
    def test_cli_emit_then_from_ir(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            work = Path(temp_dir)
            ir_path = work / "deck.slideir.json"
            pptx_path = work / "deck.pptx"
            emit_code = main(
                [
                    "emit-ir",
                    str(FIXTURES_DIR / "solid_background.pdf"),
                    str(ir_path),
                    "--ocr-model-root",
                    str(work / "no-ocr"),
                ]
            )
            self.assertEqual(emit_code, 0)
            self.assertTrue(ir_path.exists())
            from_code = main(["from-ir", str(ir_path), str(pptx_path)])
            self.assertEqual(from_code, 0)
            self.assertTrue(pptx_path.exists())


if __name__ == "__main__":
    unittest.main()
