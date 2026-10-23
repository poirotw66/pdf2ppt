"""Unit tests for the opt-in VLM layout planner (Phase 3)."""

from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image

from pdf2ppt.core import ConversionOptions
from pdf2ppt.ir_io import read_slide_ir, write_slide_ir
from pdf2ppt.layout_planner import (
    GEMINI_LAYOUT_DEFAULT_MODEL,
    LayoutPlannerBudget,
    LayoutPlannerService,
    apply_layout_plan,
    build_layout_planner_service,
    is_known_font_family,
    parse_and_validate_layout_plan,
    resolve_font_family,
    text_has_cjk,
)
from pdf2ppt.layout_planner.gemini import GeminiLayoutPlanner
from pdf2ppt.layout_planner.validate import LayoutPlanValidationError
from pdf2ppt.models import GroupElement, PageResult, QualityScore, TextBlock
from pdf2ppt.pipeline import convert_pdf

REPO_ROOT = Path(__file__).resolve().parent.parent
FIXTURES_DIR = REPO_ROOT / "tests" / "fixtures"


def _blocks() -> list[TextBlock]:
    return [
        TextBlock(
            id="b1",
            source="ocr",
            bbox=(40.0, 30.0, 400.0, 70.0),
            text="Sprint Goals",
            confidence=0.95,
            font_size=28.0,
            reading_order=0,
        ),
        TextBlock(
            id="b2",
            source="ocr",
            bbox=(50.0, 120.0, 500.0, 150.0),
            text="Ship layout analysis",
            confidence=0.9,
            font_size=16.0,
            reading_order=1,
        ),
        TextBlock(
            id="b3",
            source="ocr",
            bbox=(50.0, 155.0, 500.0, 185.0),
            text="Keep heuristic fallback",
            confidence=0.9,
            font_size=16.0,
            reading_order=2,
        ),
        TextBlock(
            id="b4",
            source="ocr",
            bbox=(40.0, 500.0, 200.0, 530.0),
            text="Confidential",
            confidence=0.8,
            font_size=10.0,
            reading_order=3,
        ),
    ]


def _valid_plan_payload() -> dict:
    return {
        "blocks": [
            {
                "id": "b1",
                "role": "title",
                "reading_order": 0,
                "font_family_class": "sans",
                "confidence": 0.95,
            },
            {
                "id": "b2",
                "role": "list_item",
                "reading_order": 1,
                "font_family_class": "sans",
                "confidence": 0.9,
            },
            {
                "id": "b3",
                "role": "list_item",
                "reading_order": 2,
                "font_family_class": "sans",
                "confidence": 0.9,
            },
            {
                "id": "b4",
                "role": "footer",
                "reading_order": 3,
                "font_family_class": "serif",
                "confidence": 0.8,
            },
        ],
        "groups": [
            {
                "id": "g_title",
                "member_ids": ["b1"],
                "role": "title",
                "reading_order": 0,
                "confidence": 0.95,
                "kind": "text_frame",
            },
            {
                "id": "g_list",
                "member_ids": ["b2", "b3"],
                "role": "list_item",
                "reading_order": 1,
                "confidence": 0.9,
                "kind": "paragraph",
            },
        ],
    }


def _pptx_member_digests(path: Path) -> dict[str, bytes]:
    """Stable PPTX comparison: zip member contents, ignoring local header mtimes."""

    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in sorted(archive.namelist())}


class LayoutPlanValidationTests(unittest.TestCase):
    def test_valid_plan_parses(self) -> None:
        plan = parse_and_validate_layout_plan(
            _valid_plan_payload(),
            known_block_ids={"b1", "b2", "b3", "b4"},
            model=GEMINI_LAYOUT_DEFAULT_MODEL,
        )
        self.assertEqual(len(plan.blocks), 4)
        self.assertEqual(len(plan.groups), 2)

    def test_rejects_unknown_ids(self) -> None:
        payload = _valid_plan_payload()
        payload["blocks"][0]["id"] = "missing"
        with self.assertRaises(LayoutPlanValidationError):
            parse_and_validate_layout_plan(
                payload,
                known_block_ids={"b1", "b2", "b3", "b4"},
                model="test",
            )

    def test_rejects_malformed_json_shape(self) -> None:
        with self.assertRaises(LayoutPlanValidationError):
            parse_and_validate_layout_plan(
                ["not", "an", "object"],
                known_block_ids={"b1"},
                model="test",
            )

    def test_rejects_contradictory_groups(self) -> None:
        payload = _valid_plan_payload()
        payload["groups"].append(
            {
                "id": "g_dup",
                "member_ids": ["b2"],
                "role": "body",
                "reading_order": 9,
                "confidence": 0.9,
            }
        )
        with self.assertRaises(LayoutPlanValidationError) as ctx:
            parse_and_validate_layout_plan(
                payload,
                known_block_ids={"b1", "b2", "b3", "b4"},
                model="test",
            )
        self.assertIn("Contradictory", str(ctx.exception))

    def test_rejects_low_confidence(self) -> None:
        payload = _valid_plan_payload()
        payload["blocks"][0]["confidence"] = 0.1
        with self.assertRaises(LayoutPlanValidationError):
            parse_and_validate_layout_plan(
                payload,
                known_block_ids={"b1", "b2", "b3", "b4"},
                model="test",
                min_confidence=0.55,
            )


class LayoutPlanApplyTests(unittest.TestCase):
    def test_default_apply_defers_grouping_to_heuristic(self) -> None:
        plan = parse_and_validate_layout_plan(
            _valid_plan_payload(),
            known_block_ids={"b1", "b2", "b3", "b4"},
            model="test",
        )
        source = _blocks()
        result = apply_layout_plan(
            source,
            plan,
            page_number=1,
            page_height=540.0,
            skip_header_footer=True,
            apply_grouping=False,
            apply_fonts=False,
        )
        self.assertTrue(result.used_planner)
        self.assertFalse(result.frames_applied)
        # Default mode does not rewrite frames — heuristic grouping stays intact.
        self.assertEqual([block.id for block in result.text_blocks], [block.id for block in source])
        self.assertEqual(result.group_elements, [])
        self.assertIsNotNone(result.plan)

    def test_apply_grouping_merges_adjacent_list_lines(self) -> None:
        plan = parse_and_validate_layout_plan(
            _valid_plan_payload(),
            known_block_ids={"b1", "b2", "b3", "b4"},
            model="test",
        )
        result = apply_layout_plan(
            _blocks(),
            plan,
            page_number=1,
            page_height=540.0,
            page_width=960.0,
            skip_header_footer=True,
            apply_grouping=True,
            apply_fonts=False,
        )
        self.assertEqual(len(result.text_blocks), 2)
        merged = result.text_blocks[1]
        self.assertIn("Ship layout analysis", merged.text)
        self.assertIn("Keep heuristic fallback", merged.text)
        self.assertEqual(merged.line_bboxes, [(50.0, 120.0, 500.0, 150.0), (50.0, 155.0, 500.0, 185.0)])
        # Merged frames must not pin font size (disables shrink-to-fit).
        self.assertFalse(merged.pin_font_size)
        title = result.text_blocks[0]
        self.assertFalse(title.pin_font_size)

    def test_rejects_table_cell_and_wide_horizontal_merge(self) -> None:
        blocks = [
            TextBlock(id="c1", source="ocr", bbox=(40, 100, 120, 130), text="A", confidence=0.9, font_size=12),
            TextBlock(id="c2", source="ocr", bbox=(200, 100, 280, 130), text="B", confidence=0.9, font_size=12),
            TextBlock(id="c3", source="ocr", bbox=(360, 100, 440, 130), text="C", confidence=0.9, font_size=12),
        ]
        payload = {
            "blocks": [
                {"id": "c1", "role": "table_cell", "reading_order": 0, "font_family_class": "sans", "confidence": 0.9},
                {"id": "c2", "role": "table_cell", "reading_order": 1, "font_family_class": "sans", "confidence": 0.9},
                {"id": "c3", "role": "table_cell", "reading_order": 2, "font_family_class": "sans", "confidence": 0.9},
            ],
            "groups": [
                {
                    "id": "g_all",
                    "member_ids": ["c1", "c2", "c3"],
                    "role": "table_cell",
                    "reading_order": 0,
                    "confidence": 0.9,
                }
            ],
        }
        plan = parse_and_validate_layout_plan(payload, known_block_ids={"c1", "c2", "c3"}, model="t")
        result = apply_layout_plan(
            blocks,
            plan,
            apply_grouping=True,
            page_width=960.0,
            table_bboxes=[(30.0, 90.0, 450.0, 140.0)],
        )
        self.assertEqual(len(result.text_blocks), 3)
        self.assertEqual(result.group_elements, [])

    def test_keeps_known_pdf_font_even_when_fonts_enabled(self) -> None:
        blocks = _blocks()
        blocks[0].font_family = "Helvetica"
        blocks[0].source = "native"
        plan = parse_and_validate_layout_plan(
            _valid_plan_payload(),
            known_block_ids={"b1", "b2", "b3", "b4"},
            model="test",
        )
        result = apply_layout_plan(blocks, plan, apply_fonts=True, apply_grouping=False)
        title = next(block for block in result.text_blocks if block.id == "b1" or block.block_role == "title")
        self.assertEqual(title.font_family, "Helvetica")

    def test_font_family_class_mapping(self) -> None:
        self.assertEqual(
            resolve_font_family("serif", text="Hello", existing_font_family=None, apply_font_class=True),
            "Times New Roman",
        )
        self.assertTrue(text_has_cjk("中文標題"))
        self.assertFalse(text_has_cjk("Q3 2026 - Finance Team"))
        # Latin+digits must NOT get a CJK font.
        self.assertEqual(
            resolve_font_family(
                "sans",
                text="Q3 2026 - Finance Team",
                existing_font_family=None,
                apply_font_class=True,
            ),
            "Arial",
        )
        cjk_font = resolve_font_family(
            "sans", text="財務摘要", existing_font_family=None, apply_font_class=True
        ) or ""
        self.assertTrue(
            any(token in cjk_font for token in ("Noto Sans CJK", "PingFang", "JhengHei")),
            cjk_font,
        )
        self.assertTrue(is_known_font_family("Helvetica"))
        self.assertFalse(is_known_font_family(None))


class GeminiPlannerMockTests(unittest.TestCase):
    def test_valid_transport_response(self) -> None:
        def transport(_body: dict) -> dict:
            return {
                "candidates": [
                    {
                        "content": {
                            "parts": [{"text": json.dumps(_valid_plan_payload())}],
                        }
                    }
                ]
            }

        planner = GeminiLayoutPlanner(transport=transport, api_key=None)
        image = Image.new("RGB", (64, 64), color=(255, 255, 255))
        from pdf2ppt.layout_planner import blocks_to_evidence

        plan = planner.plan_page(page_image=image, blocks=blocks_to_evidence(_blocks()), page_number=1)
        self.assertEqual(plan.blocks[0].role, "title")

    def test_malformed_json_raises(self) -> None:
        def transport(_body: dict) -> dict:
            return {"candidates": [{"content": {"parts": [{"text": "{not-json"}]}}]}

        planner = GeminiLayoutPlanner(transport=transport)
        image = Image.new("RGB", (32, 32), color=(0, 0, 0))
        from pdf2ppt.layout_planner import LayoutPlannerError, blocks_to_evidence

        with self.assertRaises(LayoutPlannerError):
            planner.plan_page(page_image=image, blocks=blocks_to_evidence(_blocks())[:1])

    def test_timeout_counts_against_budget(self) -> None:
        def transport(_body: dict) -> dict:
            raise TimeoutError("timed out")

        budget = LayoutPlannerBudget(max_calls=5, max_consecutive_failures=10)
        service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=transport),
            budget=budget,
            cache_dir=None,
        )
        result = service.try_plan_page(
            page_image=Image.new("RGB", (32, 32), color=(128, 128, 128)),
            text_blocks=_blocks(),
            page_number=1,
            page_height=540.0,
        )
        self.assertIsNone(result)
        self.assertEqual(budget.fallbacks, 1)
        self.assertEqual(budget.calls, 1)
        self.assertEqual(budget.http_attempts, 1)
        self.assertEqual(budget.records[0].fallback_kind, "timeout")
        self.assertTrue(budget.records[0].http_attempted)

    def test_failed_requests_count_toward_cap(self) -> None:
        calls = {"n": 0}

        def transport(_body: dict) -> dict:
            calls["n"] += 1
            return {"candidates": [{"content": {"parts": [{"text": "{bad"}]}}]}

        budget = LayoutPlannerBudget(max_calls=2, max_consecutive_failures=99)
        service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=transport),
            budget=budget,
            cache_dir=None,
        )
        image = Image.new("RGB", (16, 16))
        for page in range(1, 6):
            service.try_plan_page(
                page_image=image,
                text_blocks=_blocks(),
                page_number=page,
                page_height=540.0,
            )
        self.assertEqual(calls["n"], 2)
        self.assertEqual(budget.calls, 2)
        self.assertEqual(budget.http_attempts, 2)

    def test_circuit_breaker_stops_after_consecutive_failures(self) -> None:
        calls = {"n": 0}

        def transport(_body: dict) -> dict:
            calls["n"] += 1
            raise TimeoutError("down")

        budget = LayoutPlannerBudget(max_calls=50, max_consecutive_failures=2)
        service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=transport),
            budget=budget,
            cache_dir=None,
        )
        image = Image.new("RGB", (16, 16))
        for page in range(1, 6):
            service.try_plan_page(
                page_image=image,
                text_blocks=_blocks(),
                page_number=page,
                page_height=540.0,
            )
        self.assertEqual(calls["n"], 2)
        self.assertTrue(budget.circuit_open)

    def test_circuit_cancels_queued_reserved_calls(self) -> None:
        """Once the breaker trips, reserved-but-not-started calls must not dispatch."""

        import threading
        import time

        calls_lock = threading.Lock()
        http_calls: list[int] = []
        release_first = threading.Event()

        def transport(_body: dict) -> dict:
            with calls_lock:
                http_calls.append(1)
                ordinal = len(http_calls)
            if ordinal == 1:
                # Hold the concurrency slot until siblings have reserved + queued.
                release_first.wait(timeout=3.0)
                raise TimeoutError("fail-open-circuit")
            raise AssertionError("queued call was dispatched after circuit open")

        budget = LayoutPlannerBudget(
            max_calls=50,
            concurrency=1,
            max_consecutive_failures=1,
            max_cost_usd=10.0,
        )
        service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=transport, timeout_s=5.0),
            budget=budget,
            cache_dir=None,
        )
        image = Image.new("RGB", (16, 16))
        barrier = threading.Barrier(3)
        errors: list[BaseException] = []

        def worker(page_number: int) -> None:
            try:
                barrier.wait(timeout=3.0)
                service.try_plan_page(
                    page_image=image,
                    text_blocks=_blocks(),
                    page_number=page_number,
                    page_height=540.0,
                )
            except BaseException as error:  # noqa: BLE001 — collect for assertion
                errors.append(error)

        threads = [threading.Thread(target=worker, args=(index + 1,)) for index in range(3)]
        for thread in threads:
            thread.start()
        for _ in range(80):
            with calls_lock:
                if http_calls:
                    break
            time.sleep(0.025)
        # Give the other two workers time to reserve and block on the slot.
        time.sleep(0.35)
        release_first.set()
        for thread in threads:
            thread.join(timeout=5.0)
        self.assertEqual(errors, [])
        self.assertEqual(len(http_calls), 1)
        self.assertTrue(budget.circuit_open)
        # Only the in-flight HTTP attempt remains counted; queued reserves refunded.
        self.assertEqual(budget.calls, 1)
        self.assertEqual(budget.http_attempts, 1)
        self.assertAlmostEqual(budget.estimated_cost_usd, 0.0, places=6)
        cancelled = [record for record in budget.records if record.fallback_kind == "circuit"]
        self.assertGreaterEqual(len(cancelled), 2)
        self.assertTrue(all(not record.http_attempted for record in cancelled))

    def test_cancel_reservation_refunds_call_and_cost(self) -> None:
        from pdf2ppt.layout_planner.budget import (
            DEFAULT_WORST_CASE_CALL_COST_USD,
            LayoutPlannerCallRecord,
        )

        budget = LayoutPlannerBudget(
            max_calls=5,
            max_cost_usd=1.0,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        self.assertTrue(budget.try_reserve())
        self.assertEqual(budget.calls, 1)
        self.assertAlmostEqual(budget.estimated_cost_usd, DEFAULT_WORST_CASE_CALL_COST_USD, places=6)
        budget.cancel_reservation(
            LayoutPlannerCallRecord(
                model="t",
                page_number=1,
                cache_hit=False,
                success=False,
                latency_ms=1.0,
                estimated_cost_usd=0.0,
                fallback_kind="circuit",
                note="test cancel",
            )
        )
        self.assertEqual(budget.calls, 0)
        self.assertEqual(budget.http_attempts, 0)
        self.assertAlmostEqual(budget.estimated_cost_usd, 0.0, places=6)
        self.assertFalse(budget.records[0].http_attempted)

    def test_atomic_reserve_prevents_overshoot(self) -> None:
        budget = LayoutPlannerBudget(max_calls=3, concurrency=8)
        reserved = sum(1 for _ in range(8) if budget.try_reserve())
        self.assertEqual(reserved, 3)
        self.assertEqual(budget.calls, 3)

    def test_cost_reconciles_usage_and_worst_case_cap(self) -> None:
        from pdf2ppt.layout_planner.budget import (
            DEFAULT_WORST_CASE_CALL_COST_USD,
            LayoutPlannerCallRecord,
            worst_case_call_cost_usd,
        )

        # Reserve from max_output_tokens worst case (~8k in + 8192 out).
        self.assertAlmostEqual(
            worst_case_call_cost_usd(), DEFAULT_WORST_CASE_CALL_COST_USD, places=6
        )
        budget = LayoutPlannerBudget(
            max_calls=10,
            max_cost_usd=0.020,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        self.assertTrue(budget.try_reserve())
        budget.commit_success(
            LayoutPlannerCallRecord(
                model="t",
                page_number=1,
                cache_hit=False,
                success=True,
                latency_ms=1.0,
                estimated_cost_usd=0.00327,
                prompt_tokens=1793,
                output_tokens=672,
            )
        )
        self.assertAlmostEqual(budget.estimated_cost_usd, 0.00327, places=5)
        # Next worst-case 0.014 fits under 0.020.
        self.assertTrue(budget.try_reserve())
        # After a second 0.00327, remaining room is ~0.013 < 0.014 → refuse.
        budget.commit_success(
            LayoutPlannerCallRecord(
                model="t",
                page_number=2,
                cache_hit=False,
                success=True,
                latency_ms=1.0,
                estimated_cost_usd=0.00327,
            )
        )
        self.assertFalse(budget.try_reserve())
        self.assertEqual(budget.calls, 2)

    def test_commit_failure_refunds_unused_worst_case_reserve(self) -> None:
        from pdf2ppt.layout_planner.budget import (
            DEFAULT_WORST_CASE_CALL_COST_USD,
            LayoutPlannerCallRecord,
        )

        budget = LayoutPlannerBudget(
            max_calls=5,
            max_cost_usd=1.0,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        self.assertTrue(budget.try_reserve())
        budget.commit_failure(
            LayoutPlannerCallRecord(
                model="t",
                page_number=1,
                cache_hit=False,
                success=False,
                latency_ms=1.0,
                estimated_cost_usd=0.0,
                fallback_kind="timeout",
            )
        )
        self.assertAlmostEqual(budget.estimated_cost_usd, 0.0, places=6)
        self.assertAlmostEqual(
            budget.records[0].reserved_cost_usd, DEFAULT_WORST_CASE_CALL_COST_USD, places=6
        )

    def test_blocks_to_evidence_caps_count_and_chars(self) -> None:
        from pdf2ppt.layout_planner.base import blocks_to_evidence

        blocks = [
            TextBlock(
                id=f"b{i}",
                source="ocr",
                bbox=(0.0, float(i), 10.0, float(i) + 1.0),
                text="x" * 500,
                confidence=0.9,
            )
            for i in range(100)
        ]
        evidence = blocks_to_evidence(blocks, max_blocks=5, max_chars_per_block=40)
        self.assertEqual(len(evidence), 5)
        self.assertLessEqual(len(evidence[0].text), 40)


class LayoutPlannerFailureRefundServiceTests(unittest.TestCase):
    """End-to-end service-path refunds: failures must not book the worst-case reserve."""

    def _service(
        self,
        transport,
        *,
        max_calls: int = 10,
        max_cost_usd: float = 1.0,
        max_consecutive_failures: int = 99,
        api_key: str | None = "test-key",
    ) -> tuple[LayoutPlannerService, LayoutPlannerBudget]:
        from pdf2ppt.layout_planner.budget import DEFAULT_WORST_CASE_CALL_COST_USD

        budget = LayoutPlannerBudget(
            max_calls=max_calls,
            max_cost_usd=max_cost_usd,
            max_consecutive_failures=max_consecutive_failures,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=transport, api_key=api_key),
            budget=budget,
            cache_dir=None,
        )
        return service, budget

    def _plan_once(self, service: LayoutPlannerService, page_number: int = 1):
        return service.try_plan_page(
            page_image=Image.new("RGB", (32, 32), color=(200, 200, 200)),
            text_blocks=_blocks(),
            page_number=page_number,
            page_height=540.0,
        )

    def _assert_failure_refund(
        self,
        budget: LayoutPlannerBudget,
        *,
        expected_cost: float,
        expected_calls: int,
        fallback_kind: str | None = None,
    ) -> None:
        from pdf2ppt.layout_planner.budget import DEFAULT_WORST_CASE_CALL_COST_USD

        self.assertEqual(budget.calls, expected_calls)
        self.assertEqual(budget.http_attempts, expected_calls)
        self.assertAlmostEqual(budget.estimated_cost_usd, expected_cost, places=6)
        # Full worst-case reserve must not stick on the report after reconcile.
        self.assertLess(
            budget.estimated_cost_usd + 1e-9,
            expected_calls * DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        for record in budget.records:
            if not record.http_attempted:
                continue
            self.assertAlmostEqual(
                record.reserved_cost_usd, DEFAULT_WORST_CASE_CALL_COST_USD, places=6
            )
            if fallback_kind is not None and not record.success:
                self.assertEqual(record.fallback_kind, fallback_kind)

    def test_http_error_refunds_reserve(self) -> None:
        from pdf2ppt.layout_planner.base import LayoutPlannerError

        def transport(_body: dict) -> dict:
            raise LayoutPlannerError("Gemini layout HTTP 500: upstream boom", kind="api_error")

        service, budget = self._service(transport)
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="api_error"
        )
        self.assertAlmostEqual(budget.records[0].estimated_cost_usd, 0.0, places=6)

    def test_timeout_refunds_reserve(self) -> None:
        def transport(_body: dict) -> dict:
            raise TimeoutError("timed out")

        service, budget = self._service(transport)
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="timeout"
        )

    def test_invalid_key_401_refunds_and_opens_circuit(self) -> None:
        import io
        import urllib.error

        from pdf2ppt.layout_planner.budget import DEFAULT_WORST_CASE_CALL_COST_USD

        def fake_urlopen(_request, timeout=None):  # noqa: ARG001
            raise urllib.error.HTTPError(
                url="https://example.invalid/generateContent",
                code=401,
                msg="Unauthorized",
                hdrs=None,  # type: ignore[arg-type]
                fp=io.BytesIO(b'{"error":{"message":"API key not valid"}}'),
            )

        budget = LayoutPlannerBudget(
            max_calls=5,
            max_cost_usd=1.0,
            max_consecutive_failures=99,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        service = LayoutPlannerService(
            GeminiLayoutPlanner(api_key="bad-key", transport=None),
            budget=budget,
            cache_dir=None,
        )
        with patch("urllib.request.urlopen", side_effect=fake_urlopen):
            self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="auth_error"
        )
        self.assertTrue(budget.circuit_open)

    def test_invalid_key_403_refunds_and_opens_circuit(self) -> None:
        import io
        import urllib.error

        from pdf2ppt.layout_planner.budget import DEFAULT_WORST_CASE_CALL_COST_USD

        def fake_urlopen(_request, timeout=None):  # noqa: ARG001
            raise urllib.error.HTTPError(
                url="https://example.invalid/generateContent",
                code=403,
                msg="Forbidden",
                hdrs=None,  # type: ignore[arg-type]
                fp=io.BytesIO(b'{"error":{"message":"Permission denied for API key"}}'),
            )

        budget = LayoutPlannerBudget(
            max_calls=5,
            max_cost_usd=1.0,
            max_consecutive_failures=99,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )
        service = LayoutPlannerService(
            GeminiLayoutPlanner(api_key="bad-key", transport=None),
            budget=budget,
            cache_dir=None,
        )
        with patch("urllib.request.urlopen", side_effect=fake_urlopen):
            self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="auth_error"
        )
        self.assertTrue(budget.circuit_open)

    def test_malformed_json_with_usage_bills_usage_cost(self) -> None:
        from pdf2ppt.layout_planner.budget import cost_from_usage_metadata

        usage = {"promptTokenCount": 2000, "candidatesTokenCount": 400}
        expected = cost_from_usage_metadata(usage)
        assert expected is not None

        def transport(_body: dict) -> dict:
            return {
                "candidates": [{"content": {"parts": [{"text": "{not-json"}]}}],
                "usageMetadata": usage,
            }

        service, budget = self._service(transport)
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=expected, expected_calls=1, fallback_kind="validation"
        )
        self.assertAlmostEqual(budget.records[0].estimated_cost_usd, expected, places=6)

    def test_malformed_json_without_usage_refunds_to_zero(self) -> None:
        def transport(_body: dict) -> dict:
            return {"candidates": [{"content": {"parts": [{"text": "{not-json"}]}}]}

        service, budget = self._service(transport)
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="validation"
        )

    def _validation_failure_payload_unknown_id(self) -> dict:
        payload = _valid_plan_payload()
        payload["blocks"][0]["id"] = "missing"
        return payload

    def _validation_failure_payload_overlap(self) -> dict:
        payload = _valid_plan_payload()
        payload["groups"].append(
            {
                "id": "g_dup",
                "member_ids": ["b2"],
                "role": "body",
                "reading_order": 9,
                "confidence": 0.9,
            }
        )
        return payload

    def _transport_plan(self, plan: dict, usage: dict | None = None):
        def transport(_body: dict) -> dict:
            body: dict = {
                "candidates": [{"content": {"parts": [{"text": json.dumps(plan)}]}}],
            }
            if usage is not None:
                body["usageMetadata"] = usage
            return body

        return transport

    def test_unknown_id_with_usage_bills_usage_and_records_tokens(self) -> None:
        from pdf2ppt.layout_planner.budget import cost_from_usage_metadata

        usage = {"promptTokenCount": 1800, "candidatesTokenCount": 350}
        expected = cost_from_usage_metadata(usage)
        assert expected is not None

        service, budget = self._service(
            self._transport_plan(self._validation_failure_payload_unknown_id(), usage)
        )
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=expected, expected_calls=1, fallback_kind="validation"
        )
        record = budget.records[0]
        self.assertEqual(record.prompt_tokens, 1800)
        self.assertEqual(record.output_tokens, 350)
        self.assertIn("unknown", (record.note or "").lower())

    def test_unknown_id_without_usage_refunds_to_zero(self) -> None:
        service, budget = self._service(
            self._transport_plan(self._validation_failure_payload_unknown_id())
        )
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="validation"
        )
        record = budget.records[0]
        self.assertIsNone(record.prompt_tokens)
        self.assertIsNone(record.output_tokens)
        self.assertEqual(budget.to_dict()["estimated_cost_usd"], 0.0)

    def test_overlap_with_usage_bills_usage_and_records_tokens(self) -> None:
        from pdf2ppt.layout_planner.budget import cost_from_usage_metadata

        usage = {"promptTokenCount": 2200, "candidatesTokenCount": 480}
        expected = cost_from_usage_metadata(usage)
        assert expected is not None

        service, budget = self._service(
            self._transport_plan(self._validation_failure_payload_overlap(), usage)
        )
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=expected, expected_calls=1, fallback_kind="validation"
        )
        record = budget.records[0]
        self.assertEqual(record.prompt_tokens, 2200)
        self.assertEqual(record.output_tokens, 480)
        self.assertIn("Contradictory", record.note or "")

    def test_overlap_without_usage_refunds_to_zero(self) -> None:
        service, budget = self._service(
            self._transport_plan(self._validation_failure_payload_overlap())
        )
        self.assertIsNone(self._plan_once(service))
        self._assert_failure_refund(
            budget, expected_cost=0.0, expected_calls=1, fallback_kind="validation"
        )
        record = budget.records[0]
        self.assertIsNone(record.prompt_tokens)
        self.assertIsNone(record.output_tokens)
        printed = budget.to_dict()["estimated_cost_usd"]
        self.assertEqual(printed, 0.0)
        self.assertEqual(str(printed), "0.0")

    def test_retries_each_count_as_calls_and_usd_cap_still_blocks(self) -> None:
        """Each failed attempt counts; refunds free the USD reserve so the cap still gates."""
        from pdf2ppt.layout_planner.budget import DEFAULT_WORST_CASE_CALL_COST_USD

        calls = {"n": 0}

        def failing_transport(_body: dict) -> dict:
            calls["n"] += 1
            return {"candidates": [{"content": {"parts": [{"text": "{bad"}]}}]}

        # Without refund, two failures would book ~2× reserve and exceed 1.5×.
        # With refund, both attempts proceed and reported cost stays 0.
        service, budget = self._service(
            failing_transport,
            max_calls=10,
            max_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD * 1.5,
            max_consecutive_failures=99,
        )
        self.assertIsNone(self._plan_once(service, page_number=1))
        self.assertIsNone(self._plan_once(service, page_number=2))
        self.assertEqual(calls["n"], 2)
        self.assertEqual(budget.calls, 2)
        self.assertEqual(budget.http_attempts, 2)
        self.assertAlmostEqual(budget.estimated_cost_usd, 0.0, places=6)

        # Tight cap: exactly one worst-case reserve. A real success leaves too
        # little room for another reserve → service records a budget fallback.
        tight = LayoutPlannerBudget(
            max_calls=10,
            max_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
            max_consecutive_failures=99,
            worst_case_call_cost_usd=DEFAULT_WORST_CASE_CALL_COST_USD,
        )

        def success_transport(_body: dict) -> dict:
            calls["n"] += 1
            return {
                "candidates": [
                    {
                        "content": {
                            "parts": [{"text": json.dumps(_valid_plan_payload())}],
                        }
                    }
                ],
                "usageMetadata": {
                    "promptTokenCount": 2000,
                    "candidatesTokenCount": 500,
                },
            }

        tight_service = LayoutPlannerService(
            GeminiLayoutPlanner(transport=success_transport, api_key="test-key"),
            budget=tight,
            cache_dir=None,
        )
        self.assertIsNotNone(self._plan_once(tight_service, page_number=1))
        self.assertEqual(tight.calls, 1)
        self.assertGreater(tight.estimated_cost_usd, 0.0)
        self.assertFalse(tight.try_reserve())
        blocked = self._plan_once(tight_service, page_number=2)
        self.assertIsNone(blocked)
        self.assertEqual(tight.calls, 1)
        self.assertEqual(tight.records[-1].fallback_kind, "budget")


class LayoutPlannerServiceConfigTests(unittest.TestCase):
    def test_unconfigured_returns_none_and_records_skip_reason(self) -> None:
        env = {key: value for key, value in __import__("os").environ.items() if "API_KEY" not in key}
        with patch.dict("os.environ", env, clear=True):
            options = ConversionOptions(
                input_path=Path("in.pdf"),
                output_path=Path("out.pptx"),
                report_path=Path("out.report.json"),
                layout_planner="gemini",
            )
            self.assertIsNone(build_layout_planner_service(options))
            self.assertIsNotNone(options.layout_planner_skip_reason)
            self.assertIsInstance(options.layout_planner_budget, LayoutPlannerBudget)


class SlideIrGroupsRoundTripTests(unittest.TestCase):
    def test_groups_and_plan_metadata_round_trip(self) -> None:
        page = PageResult(
            page_number=1,
            page_kind="digital",
            background_mode="elements",
            width_pt=612.0,
            height_pt=792.0,
            text_blocks=_blocks()[:1],
            quality_score=QualityScore(1.0, 1.0, 1.0, 1.0),
            fallback_reason=None,
            group_elements=[
                GroupElement(
                    id="g1",
                    member_ids=["b1"],
                    bbox=(40.0, 30.0, 400.0, 70.0),
                    provenance="layout_planner",
                    confidence=0.9,
                )
            ],
            plan_metadata={"engine": "layout_planner", "model": "test"},
            layout_engine="layout_planner",
        )
        with tempfile.TemporaryDirectory() as tmp:
            ir_path = Path(tmp) / "deck.slideir.json"
            write_slide_ir(
                [page],
                ir_path,
                slide_width_pt=612.0,
                slide_height_pt=792.0,
                plan_metadata={"engine": "layout_planner", "stats": {"calls": 1}},
            )
            deck = read_slide_ir(ir_path)
            self.assertEqual(deck.plan_metadata.get("engine"), "layout_planner")
            self.assertEqual(len(deck.pages[0].group_elements), 1)
            self.assertEqual(deck.pages[0].group_elements[0].member_ids, ["b1"])
            self.assertEqual(deck.pages[0].plan_metadata.get("engine"), "layout_planner")


@unittest.skipUnless((FIXTURES_DIR / "bullet_list.pdf").is_file(), "fixture missing")
class LayoutPlannerOffByteIdenticalTests(unittest.TestCase):
    def _convert(self, fixture: Path, out: Path, report: Path, **kwargs: object) -> None:
        options = ConversionOptions(
            input_path=fixture,
            output_path=out,
            report_path=report,
            mode="editable",
            ocr_model_root=None,
            layout_engine="heuristic",
            vector_mode="off",
            inpaint_engine="white-box",
            **kwargs,  # type: ignore[arg-type]
        )
        with patch("pdf2ppt.pipeline.OcrEngine") as ocr_cls:
            ocr_cls.return_value.extract_text_blocks_batch.return_value = []
            convert_pdf(options)

    def test_planner_off_matches_default_pipeline_parts(self) -> None:
        fixture = FIXTURES_DIR / "bullet_list.pdf"
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            out_default = tmp_path / "default.pptx"
            out_explicit = tmp_path / "explicit.pptx"
            self._convert(fixture, out_default, tmp_path / "d.report.json", layout_planner=None)
            self._convert(fixture, out_explicit, tmp_path / "e.report.json", layout_planner=None)
            self.assertEqual(_pptx_member_digests(out_default), _pptx_member_digests(out_explicit))

    def test_planner_without_api_key_matches_default_parts(self) -> None:
        fixture = FIXTURES_DIR / "bullet_list.pdf"
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            out_default = tmp_path / "default.pptx"
            out_flag = tmp_path / "flag.pptx"
            env = {
                key: value
                for key, value in __import__("os").environ.items()
                if key not in {"GEMINI_API_KEY", "GOOGLE_API_KEY"}
            }
            with patch.dict("os.environ", env, clear=True):
                self._convert(fixture, out_default, tmp_path / "d.report.json", layout_planner=None)
                self._convert(
                    fixture,
                    out_flag,
                    tmp_path / "f.report.json",
                    layout_planner="gemini",
                    layout_planner_api_key=None,
                )
            self.assertEqual(_pptx_member_digests(out_default), _pptx_member_digests(out_flag))


def _pptx_text_frame_count(path: Path) -> int:
    from pptx import Presentation

    presentation = Presentation(str(path))
    count = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                if (shape.text_frame.text or "").strip():
                    count += 1
    return count


def _pptx_list_autonumber_starts(path: Path) -> list[int | None]:
    """Return startAt (or None=default 1) for each buAutoNum paragraph on slide 0."""

    from pptx import Presentation
    from pptx.oxml.ns import qn

    presentation = Presentation(str(path))
    starts: list[int | None] = []
    if not presentation.slides:
        return starts
    slide = presentation.slides[0]
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            try:
                p_pr = paragraph._p.pPr  # noqa: SLF001
            except Exception:  # noqa: BLE001
                continue
            if p_pr is None:
                continue
            for child in p_pr:
                if child.tag == qn("a:buAutoNum"):
                    raw = child.get("startAt")
                    starts.append(int(raw) if raw is not None else None)
    return starts


@unittest.skipUnless((FIXTURES_DIR / "agenda_numbered_list.pdf").is_file(), "fixture missing")
class LayoutPlannerDefaultKeepsHeuristicGroupingTests(unittest.TestCase):
    """Default planner mode must not split lines vs planner-off (e02/e05 regressions)."""

    def _convert_with_transport(self, fixture: Path, out: Path, report: Path, transport) -> None:
        options = ConversionOptions(
            input_path=fixture,
            output_path=out,
            report_path=report,
            mode="editable",
            ocr_model_root=None,
            layout_engine="heuristic",
            vector_mode="off",
            inpaint_engine="white-box",
            layout_planner="gemini",
            layout_planner_transport=transport,
            layout_planner_apply_grouping=False,
            layout_planner_apply_fonts=False,
            layout_planner_budget=LayoutPlannerBudget(max_calls=5),
        )
        convert_pdf(options)

    def _off_convert(self, fixture: Path, out: Path, report: Path) -> None:
        options = ConversionOptions(
            input_path=fixture,
            output_path=out,
            report_path=report,
            mode="editable",
            ocr_model_root=None,
            layout_engine="heuristic",
            vector_mode="off",
            inpaint_engine="white-box",
            layout_planner=None,
        )
        convert_pdf(options)

    @staticmethod
    def _role_transport(role_by_substr: dict[str, str]):
        def transport(body: dict) -> dict:
            # Extract live block ids from the request and assign roles by text.
            text = ""
            try:
                parts = body["contents"][0]["parts"]
                text = next(part["text"] for part in parts if isinstance(part, dict) and "text" in part)
            except Exception:  # noqa: BLE001
                text = ""
            blocks_json = text.split("BLOCKS:\n", 1)[-1] if "BLOCKS:\n" in text else "[]"
            try:
                req_blocks = json.loads(blocks_json)
            except Exception:  # noqa: BLE001
                req_blocks = []
            out_blocks = []
            out_groups = []
            for index, item in enumerate(req_blocks if isinstance(req_blocks, list) else []):
                bid = str(item.get("id"))
                item_text = str(item.get("text") or "")
                role = "body"
                for needle, mapped in role_by_substr.items():
                    if needle in item_text:
                        role = mapped
                        break
                out_blocks.append(
                    {
                        "id": bid,
                        "role": role,
                        "reading_order": index,
                        "font_family_class": "sans",
                        "confidence": 0.9,
                    }
                )
                out_groups.append(
                    {
                        "id": f"g_{bid}",
                        "member_ids": [bid],
                        "role": role,
                        "reading_order": index,
                        "confidence": 0.9,
                        "kind": "text_frame",
                    }
                )
            payload = {"blocks": out_blocks, "groups": out_groups}
            return {"candidates": [{"content": {"parts": [{"text": json.dumps(payload)}]}}]}

        return transport

    def test_agenda_frame_count_and_numbering_match_planner_off(self) -> None:
        fixture = FIXTURES_DIR / "agenda_numbered_list.pdf"
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            off = tmp_path / "off.pptx"
            on = tmp_path / "on.pptx"
            self._off_convert(fixture, off, tmp_path / "off.report.json")
            self._convert_with_transport(
                fixture,
                on,
                tmp_path / "on.report.json",
                self._role_transport({"Agenda": "title", "Financial": "list_item", "Regional": "list_item"}),
            )
            self.assertEqual(_pptx_text_frame_count(off), _pptx_text_frame_count(on))
            off_starts = _pptx_list_autonumber_starts(off)
            on_starts = _pptx_list_autonumber_starts(on)
            # Split boxes restart at 1 each (None); a single list frame has one
            # sequence. Counts must match planner-off.
            self.assertEqual(len(off_starts), len(on_starts))
            self.assertEqual(off_starts, on_starts)

    def test_cjk_paragraph_frame_count_matches_planner_off(self) -> None:
        fixture = FIXTURES_DIR / "cjk_paragraph.pdf"
        if not fixture.is_file():
            self.skipTest("cjk_paragraph.pdf missing")
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            off = tmp_path / "off.pptx"
            on = tmp_path / "on.pptx"
            self._off_convert(fixture, off, tmp_path / "off.report.json")
            self._convert_with_transport(
                fixture,
                on,
                tmp_path / "on.report.json",
                self._role_transport({"品牌": "title", "我們": "body", "創辦人": "body"}),
            )
            self.assertEqual(_pptx_text_frame_count(off), _pptx_text_frame_count(on))


def _pptx_slide_text(path: Path) -> str:
    from pdf2ppt.refine import extract_pptx_slide_text

    return extract_pptx_slide_text(path)


def _normalize_pptx_chars(text: str) -> str:
    return "".join(ch for ch in text if not ch.isspace())


class AnnotateGroupedBlocksTitlePlaceholderTests(unittest.TestCase):
    """Mixed title+body frames must not use the title-only placeholder."""

    def test_mixed_title_body_frame_keeps_all_text_and_no_placeholder(self) -> None:
        from pdf2ppt.layout_planner.apply import (
            annotate_grouped_blocks_with_plan,
            planner_text_regression,
        )
        from pdf2ppt.layout_planner.types import LayoutPlan, PlannedBlock

        merged = TextBlock(
            id="merged",
            source="native",
            bbox=(80.0, 67.0, 512.0, 209.0),
            text=(
                "Card at 93% of the page\n"
                "This body paragraph sits on a large colored card. It wraps\n"
                "across several lines in the PDF so we can check line spacing\n"
                "and z-order together on one slide."
            ),
            confidence=0.9,
            font_size=28.0,
            reading_order=0,
            block_role="body",
            layout_label="text",
        )
        sources = [
            TextBlock(
                id="native_1_1_0",
                source="native",
                bbox=(80.0, 67.0, 405.0, 109.0),
                text="Card at 93% of the page",
                confidence=0.9,
                font_size=30.0,
            ),
            TextBlock(
                id="native_1_2_0",
                source="native",
                bbox=(80.0, 140.0, 512.0, 209.0),
                text=(
                    "This body paragraph sits on a large colored card. It wraps\n"
                    "across several lines in the PDF so we can check line spacing\n"
                    "and z-order together on one slide."
                ),
                confidence=0.9,
                font_size=16.0,
            ),
        ]
        plan = LayoutPlan(
            blocks=[
                PlannedBlock(
                    id="native_1_1_0",
                    role="title",
                    reading_order=0,
                    font_family_class="sans",
                    confidence=0.95,
                ),
                PlannedBlock(
                    id="native_1_2_0",
                    role="body",
                    reading_order=1,
                    font_family_class="sans",
                    confidence=0.9,
                ),
            ],
            groups=[],
            model="test",
        )
        out = annotate_grouped_blocks_with_plan(
            [merged], plan, source_blocks=sources, page_height=540.0
        )
        self.assertEqual(len(out), 1)
        self.assertFalse(out[0].use_title_placeholder)
        self.assertNotEqual(out[0].block_role, "title")
        self.assertEqual(
            _normalize_pptx_chars(out[0].text), _normalize_pptx_chars(merged.text)
        )
        self.assertIsNone(planner_text_regression([merged], out, plan))

    def test_planner_text_regression_detects_dropped_body(self) -> None:
        from pdf2ppt.layout_planner.apply import planner_text_regression
        from pdf2ppt.layout_planner.types import LayoutPlan, PlannedBlock

        before = [
            TextBlock(
                id="merged",
                source="native",
                bbox=(0, 0, 100, 100),
                text="Title\nBody paragraph that must stay",
                confidence=0.9,
            )
        ]
        after = [
            TextBlock(
                id="merged",
                source="native",
                bbox=(0, 0, 100, 100),
                text="Title",
                confidence=0.9,
            )
        ]
        plan = LayoutPlan(
            blocks=[
                PlannedBlock(
                    id="t",
                    role="title",
                    reading_order=0,
                    font_family_class="sans",
                    confidence=0.9,
                )
            ],
            groups=[],
            model="test",
        )
        reason = planner_text_regression(before, after, plan)
        self.assertIsNotNone(reason)
        self.assertIn("dropped", reason or "")

    def test_planner_text_regression_allows_matched_ocr_header_footer(self) -> None:
        """Heuristic frame ids differ from plan ids; HF removal must still be allowed."""

        from pdf2ppt.layout_planner.apply import planner_text_regression
        from pdf2ppt.layout_planner.types import LayoutPlan, PlannedBlock

        header_src = TextBlock(
            id="ocr_1_0",
            source="ocr",
            bbox=(40.0, 10.0, 400.0, 30.0),
            text="ACME Corp | Q3 Review",
            confidence=0.9,
            font_size=10.0,
        )
        body_src = TextBlock(
            id="ocr_1_1",
            source="ocr",
            bbox=(40.0, 80.0, 500.0, 120.0),
            text="Sprint Goals",
            confidence=0.95,
            font_size=28.0,
        )
        # Heuristic frames use layout_* ids — never present in the Gemini plan.
        header_frame = TextBlock(
            id="layout_text_1_0",
            source="ocr",
            bbox=(40.0, 10.0, 400.0, 30.0),
            text="ACME Corp | Q3 Review",
            confidence=0.9,
            font_size=10.0,
        )
        body_frame = TextBlock(
            id="layout_text_1_1",
            source="ocr",
            bbox=(40.0, 80.0, 500.0, 120.0),
            text="Sprint Goals",
            confidence=0.95,
            font_size=28.0,
        )
        plan = LayoutPlan(
            blocks=[
                PlannedBlock(
                    id="ocr_1_0",
                    role="header",
                    reading_order=0,
                    font_family_class="sans",
                    confidence=0.9,
                ),
                PlannedBlock(
                    id="ocr_1_1",
                    role="title",
                    reading_order=1,
                    font_family_class="sans",
                    confidence=0.95,
                ),
            ],
            groups=[],
            model="test",
        )
        # After annotate: header frame dropped, body kept.
        reason = planner_text_regression(
            [header_frame, body_frame],
            [body_frame],
            plan,
            source_blocks=[header_src, body_src],
        )
        self.assertIsNone(reason)

    def test_planner_text_regression_still_blocks_body_loss_with_sources(self) -> None:
        from pdf2ppt.layout_planner.apply import planner_text_regression
        from pdf2ppt.layout_planner.types import LayoutPlan, PlannedBlock

        title_src = TextBlock(
            id="native_1_0",
            source="native",
            bbox=(80.0, 60.0, 400.0, 100.0),
            text="Card title",
            confidence=0.9,
            font_size=28.0,
        )
        body_src = TextBlock(
            id="native_1_1",
            source="native",
            bbox=(80.0, 140.0, 500.0, 200.0),
            text="Body paragraph that must stay",
            confidence=0.9,
            font_size=14.0,
        )
        merged = TextBlock(
            id="layout_text_1_0",
            source="native",
            bbox=(80.0, 60.0, 500.0, 200.0),
            text="Card title\nBody paragraph that must stay",
            confidence=0.9,
            font_size=28.0,
        )
        dropped = TextBlock(
            id="layout_text_1_0",
            source="native",
            bbox=(80.0, 60.0, 500.0, 200.0),
            text="Card title",
            confidence=0.9,
            font_size=28.0,
        )
        plan = LayoutPlan(
            blocks=[
                PlannedBlock(
                    id="native_1_0",
                    role="title",
                    reading_order=0,
                    font_family_class="sans",
                    confidence=0.95,
                ),
                PlannedBlock(
                    id="native_1_1",
                    role="body",
                    reading_order=1,
                    font_family_class="sans",
                    confidence=0.9,
                ),
            ],
            groups=[],
            model="test",
        )
        reason = planner_text_regression(
            [merged],
            [dropped],
            plan,
            source_blocks=[title_src, body_src],
        )
        self.assertIsNotNone(reason)
        self.assertIn("dropped", reason or "")


@unittest.skipUnless((FIXTURES_DIR / "card_93pct.pdf").is_file(), "fixture missing")
class LayoutPlannerTitleBodyNoDropTests(unittest.TestCase):
    """card_93pct / dark_97pct: default + fonts PPTX text must equal planner-off."""

    def _role_transport(self, title_needles: tuple[str, ...]):
        def transport(body: dict) -> dict:
            text = ""
            try:
                parts = body["contents"][0]["parts"]
                text = next(
                    part["text"] for part in parts if isinstance(part, dict) and "text" in part
                )
            except Exception:  # noqa: BLE001
                text = ""
            blocks_json = text.split("BLOCKS:\n", 1)[-1] if "BLOCKS:\n" in text else "[]"
            try:
                req_blocks = json.loads(blocks_json)
            except Exception:  # noqa: BLE001
                req_blocks = []
            out_blocks = []
            out_groups = []
            for index, item in enumerate(req_blocks if isinstance(req_blocks, list) else []):
                bid = str(item.get("id"))
                item_text = str(item.get("text") or "")
                role = (
                    "title"
                    if any(needle in item_text for needle in title_needles)
                    else "body"
                )
                out_blocks.append(
                    {
                        "id": bid,
                        "role": role,
                        "reading_order": index,
                        "font_family_class": "sans",
                        "confidence": 0.95,
                    }
                )
                out_groups.append(
                    {
                        "id": f"g_{bid}",
                        "member_ids": [bid],
                        "role": role,
                        "reading_order": index,
                        "confidence": 0.95,
                        "kind": "text_frame",
                    }
                )
            payload = {"blocks": out_blocks, "groups": out_groups}
            return {"candidates": [{"content": {"parts": [{"text": json.dumps(payload)}]}}]}

        return transport

    def _convert(
        self,
        fixture: Path,
        out: Path,
        report: Path,
        *,
        planner: bool,
        apply_fonts: bool = False,
        transport=None,
    ) -> None:
        kwargs: dict = {
            "input_path": fixture,
            "output_path": out,
            "report_path": report,
            "mode": "editable",
            "ocr_model_root": None,
            "layout_engine": "heuristic",
            "vector_mode": "off",
            "inpaint_engine": "white-box",
        }
        if planner:
            kwargs.update(
                {
                    "layout_planner": "gemini",
                    "layout_planner_transport": transport,
                    "layout_planner_apply_grouping": False,
                    "layout_planner_apply_fonts": apply_fonts,
                    "layout_planner_budget": LayoutPlannerBudget(max_calls=5),
                }
            )
        else:
            kwargs["layout_planner"] = None
        convert_pdf(ConversionOptions(**kwargs))  # type: ignore[arg-type]

    def _assert_text_equals_off(
        self, fixture: Path, title_needles: tuple[str, ...], *, apply_fonts: bool
    ) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            off = tmp_path / "off.pptx"
            on = tmp_path / "on.pptx"
            self._convert(fixture, off, tmp_path / "off.json", planner=False)
            self._convert(
                fixture,
                on,
                tmp_path / "on.json",
                planner=True,
                apply_fonts=apply_fonts,
                transport=self._role_transport(title_needles),
            )
            off_text = _pptx_slide_text(off)
            on_text = _pptx_slide_text(on)
            self.assertEqual(
                _normalize_pptx_chars(off_text),
                _normalize_pptx_chars(on_text),
                msg=f"fonts={apply_fonts}\nOFF:\n{off_text}\nON:\n{on_text}",
            )
            # Body content must survive (not only the title line).
            self.assertGreater(len(_normalize_pptx_chars(on_text)), 40)

    def test_card_93pct_default_text_equals_planner_off(self) -> None:
        self._assert_text_equals_off(
            FIXTURES_DIR / "card_93pct.pdf", ("Card at 93%",), apply_fonts=False
        )

    def test_card_93pct_fonts_text_equals_planner_off(self) -> None:
        self._assert_text_equals_off(
            FIXTURES_DIR / "card_93pct.pdf", ("Card at 93%",), apply_fonts=True
        )

    def test_dark_97pct_default_text_equals_planner_off(self) -> None:
        fixture = FIXTURES_DIR / "dark_97pct.pdf"
        if not fixture.is_file():
            self.skipTest("dark_97pct.pdf missing")
        self._assert_text_equals_off(fixture, ("Dark 97%",), apply_fonts=False)

    def test_dark_97pct_fonts_text_equals_planner_off(self) -> None:
        fixture = FIXTURES_DIR / "dark_97pct.pdf"
        if not fixture.is_file():
            self.skipTest("dark_97pct.pdf missing")
        self._assert_text_equals_off(fixture, ("Dark 97%",), apply_fonts=True)


class RecordedOracleDetectionTests(unittest.TestCase):
    def _oracle_fn(self):
        import importlib.util

        path = REPO_ROOT / "scripts" / "bench_layout_planner.py"
        spec = importlib.util.spec_from_file_location("bench_layout_planner_mod", path)
        assert spec is not None and spec.loader is not None
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod._looks_like_ground_truth_oracle

    def test_rejects_missing_live_metadata(self) -> None:
        fn = self._oracle_fn()
        self.assertTrue(fn({"blocks": [{"id": "native_1_0_0", "role": "title"}]}, None))

    def test_rejects_gt_copy_plus_extra_id(self) -> None:
        fn = self._oracle_fn()
        gt = {
            "paragraphs": [
                {"id": "title", "role": "title", "confidence": 0.92},
                {"id": "body1", "role": "body", "confidence": 0.92},
            ]
        }
        recorded = {
            "recorded_from": "live",
            "model": "gemini-3.1-flash-lite",
            "blocks": [
                {"id": "title", "role": "title", "confidence": 0.92},
                {"id": "body1", "role": "body", "confidence": 0.92},
                {"id": "decoy_extra", "role": "other", "confidence": 0.5},
            ],
        }
        self.assertTrue(fn(recorded, gt))

    def test_accepts_real_native_ids(self) -> None:
        fn = self._oracle_fn()
        gt = {
            "paragraphs": [
                {"id": "title", "role": "title"},
                {"id": "body1", "role": "body"},
            ]
        }
        recorded = {
            "recorded_from": "live",
            "model": "gemini-3.1-flash-lite",
            "blocks": [
                {"id": "native_1_0_0", "role": "title", "confidence": 1.0},
                {"id": "native_1_1_0", "role": "body", "confidence": 1.0},
            ],
        }
        self.assertFalse(fn(recorded, gt))


@unittest.skipUnless(
    __import__("os").environ.get("PDF2PPT_RUN_LAYOUT_PLANNER") == "1",
    "set PDF2PPT_RUN_LAYOUT_PLANNER=1 with GEMINI_API_KEY for live tests",
)
class LayoutPlannerLiveTests(unittest.TestCase):
    def test_live_gemini_call_cheap(self) -> None:
        import os

        self.assertTrue(os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY"))
        planner = GeminiLayoutPlanner(timeout_s=60.0)
        image = Image.new("RGB", (320, 180), color=(255, 255, 255))
        from pdf2ppt.layout_planner import blocks_to_evidence

        plan = planner.plan_page(
            page_image=image,
            blocks=blocks_to_evidence(_blocks()),
            page_number=1,
        )
        self.assertGreaterEqual(len(plan.blocks), 1)


if __name__ == "__main__":
    unittest.main()
